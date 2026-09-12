import docx, os, hashlib
p=r"D:\Document\数学建模\2026CUMCM\tmp\paper-review\问题分析_去第一人称专业版.docx"
d=docx.Document(p)
out=[]
for i,par in enumerate(d.paragraphs):
    t=par.text.strip()
    if t:
        out.append(f"[{i}][{par.style.name}] {t}")
out.append(f"\n--- tables: {len(d.tables)} ---")
for ti,tb in enumerate(d.tables):
    out.append(f"== table {ti} ({len(tb.rows)}x{len(tb.columns)}) ==")
    for r in tb.rows:
        out.append(" | ".join(c.text.strip().replace("\n"," ") for c in r.cells))
txt="\n".join(out)
open(r"D:\Document\数学建模\2026CUMCM\tmp\paper-review\docx_text.txt","w",encoding="utf-8").write(txt)
print("paragraphs:",len(d.paragraphs),"chars:",len(txt))
print("inline shapes:",len(d.inline_shapes))
for s in d.inline_shapes:
    print("shape",s.type,s.width,s.height)
