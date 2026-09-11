"""Read-only original PDF/XLSX audit. CLI only; no solver, no VS verification."""
from pathlib import Path
import json, hashlib, math
import pymupdf as fitz
import openpyxl

root=Path('D:/Document/数学建模/2026CUMCM')
src=root/'problem_files/CUMCM2026Problems/A题'
out=root/'notes/A-restart/2026-09-11'
out.mkdir(parents=True,exist_ok=True)
doc=fitz.open(src/'A题.pdf')
(out/'fresh_problem_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n'+p.get_text() for i,p in enumerate(doc)),encoding='utf-8')
for i,p in enumerate(doc):
    p.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(out/f'fresh_page_{i+1}.png')
audit={'scope':'fresh original-only CLI audit; no PDE, no Visual Studio GUI or human code review','sources':[],'workbooks':[]}
for p in sorted(src.rglob('*')):
    if p.is_file(): audit['sources'].append({'path':str(p.relative_to(root)), 'size':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
for p in sorted(src.rglob('*.xlsx')):
    wb=openpyxl.load_workbook(p,data_only=False)
    b={'file':str(p.relative_to(src)),'sheets':[]}
    for ws in wb:
        rows=list(ws.values)
        nonempty=[(i+1,list(r)) for i,r in enumerate(rows) if any(v is not None for v in r)]
        cols=[]
        for j in range(ws.max_column):
            vals=[r[j] for r in rows[1:]]
            nums=[v for v in vals if isinstance(v,(int,float)) and not isinstance(v,bool)]
            cols.append({'column':j+1,'header':rows[0][j] if rows else None,'nonempty_below_first_row':sum(v is not None for v in vals),'numeric_count':len(nums),'min':min(nums) if nums else None,'max':max(nums) if nums else None,'missing_below_first_row':sum(v is None for v in vals)})
        s={'name':ws.title,'max_row':ws.max_row,'max_column':ws.max_column,'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'nonempty_rows':len(nonempty),'first_rows':nonempty[:8],'last_rows':nonempty[-5:],'columns':cols,'formula_cells':[(c.coordinate,c.value) for row in ws for c in row if c.data_type=='f'],'error_cells':[(c.coordinate,c.value) for row in ws for c in row if c.data_type=='e']}
        s['all_nonempty_cells']={c.coordinate:c.value for row in ws for c in row if c.value is not None}
        times=[r[0] for r in rows[1:] if isinstance(r[0],(int,float))]
        s['first_column_numeric_sequence']={'count':len(times),'unique':len(set(times)),'steps':sorted(set(round(b-a,10) for a,b in zip(times,times[1:])))}
        b['sheets'].append(s)
    audit['workbooks'].append(b)
env=list(openpyxl.load_workbook(src/'附件/附件1.xlsx',data_only=True).active.values)[1:]
rad=list(openpyxl.load_workbook(src/'附件/附件2.xlsx',data_only=True).active.values)[1:]
audit['derived_checks']={'environment_duplicate_rows':len(env)-len(set(env)),'radius_duplicate_rows':len(rad)-len(set(rad)),'radius_increases':[(a,b) for a,b in zip(rad,rad[1:]) if b[1]>a[1]],'radius_final_constant_run_start_s':next(r[0] for i,r in enumerate(rad) if all(x[1]==rad[-1][1] for x in rad[i:])),'last_hour_environment':{str(j):{'min':min(r[j] for r in env if r[0]>=10800),'max':max(r[j] for r in env if r[0]>=10800),'mean':sum(r[j] for r in env if r[0]>=10800)/sum(r[0]>=10800 for r in env)} for j in [1,2]}}
audit['problem_constants']={'length_cm':25,'initial_radius_cm':2,'initial_T_C':28,'initial_C_kg_per_kg':2.55,'h_W_m2_K':25,'hm_m_s':8e-7,'q1':{'rho_kg_m3':820,'cp_J_kg_K':2600,'k_W_m_K':0.36,'D_m2_s':'7e-9*exp(-0.89/C)'},'q2_q3':{'rho':'650+128*C','cp':'1450+2736*C/(C+1)','k':'0.21+0.38*C/(C+1)','D':'2.4e-3*exp(-0.45/C)*exp(-3850/T_K)'},'q4':{'rho':'760+90*C','cp':'1850+2150*C/(C+1)','k':'0.12+0.20*C/(C+1)','D':'4.2e-4*exp(-0.30/C)*exp(-3850/T_K)'},'dryness_criterion':'everywhere C < 0.15 kg/kg','rounding_decimals':4}
audit['evidence_limits']=['Environmental data end at 14400 s; later behavior requires explicit assumption.','Radius data end at 259200 s; later radius requires explicit extrapolation if needed.','No internal measured T or C exists for calibration/validation.','No axial size evolution, latent heat, or sorption isotherm is supplied.','Q2 complete result time horizon not explicitly repeated after 3h table requirement.','Q4 moving surface is explicit in template; out-of-domain values must not be filled as physical concentrations.','No PDE computation, GUI reproduction, or human review performed.']
rho_d0=(760+90*2.55)/(1+2.55)
r_min_cm=2*math.sqrt(rho_d0/760)
first_violation=next(r for r in rad if r[1]<r_min_cm)
audit['q4_conditional_global_dry_mass_feasibility']={'assumptions':['rho(C)=760+90C is actual total wet mass per current volume','C is water mass/dry mass and everywhere nonnegative','cylindrical length stays fixed at 25 cm','initial C=2.55 is uniform','dry mass conserved'],'rho_d_formula_kg_m3':'(760+90*C)/(1+C)=90+670/(1+C)','rho_d_upper_bound_kg_m3':760,'rho_d_initial_kg_m3':rho_d0,'necessary_radius_min_cm':r_min_cm,'first_record_below_radius_bound':{'time_s':first_violation[0],'time_h':first_violation[0]/3600,'radius_cm':first_violation[1]},'final_time_s':rad[-1][0],'final_radius_cm':rad[-1][1],'final_max_possible_dry_mass_ratio':760/rho_d0*(rad[-1][1]/2)**2,'initial_dry_mass_kg':rho_d0*math.pi*0.02**2*0.25,'final_max_possible_dry_mass_kg':760*math.pi*(rad[-1][1]/100)**2*0.25,'conclusion':'At 72h even C=0 everywhere cannot conserve initial dry mass under these joint assumptions. This is conditional inconsistency, not a standalone rejection of the empirical formula.','precision_note':'First crossing uses published rounded radius records, without interpolation; nominal crossing has a 0.0002029565 cm margin, smaller than 0.001 cm reporting precision.'}
(out/'data_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print((out/'fresh_problem_text.txt').read_text(encoding='utf-8'))
for b in audit['workbooks']:
    print(json.dumps({**b,'sheets':[{k:v for k,v in s.items() if k!='all_nonempty_cells'} for s in b['sheets']]},ensure_ascii=False,default=str))
