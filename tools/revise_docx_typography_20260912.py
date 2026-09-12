"""Edit the user DOCX in place structurally, while delivering a separate copy."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from copy import deepcopy
import re,json,hashlib,sys
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path.cwd();assert ROOT.as_posix()=='D:/Document/数学建模/2026CUMCM'
QA=ROOT/'paper_output/qa/manuscript_revision_20260912'
sys.path.insert(0,str(ROOT/'tools'))
import build_manuscript_20260912 as layout
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_COLOR_INDEX

SOURCE=QA/'user_edited_source.docx'
OUT=ROOT/'paper_output/paper/A题_论文_公式规范与精简附录版.docx'
NS={'w':qn('w:p').split('}')[0][1:],'m':qn('m:oMath').split('}')[0][1:]}
visible=lambda node:''.join(node.xpath('.//w:t/text()|.//m:t/text()',namespaces=NS)) if type(node) is etree._Element else ''.join(t.text or '' for t in node.iter() if t.tag in (qn('w:t'),qn('m:t')))
def child(parent,name):
    item=parent.find(qn(name))
    if item is None:item=OxmlElement(name);parent.append(item)
    return item
def wprop(rpr,name,value):child(rpr,name).set(qn('w:val'),value)
def font(rpr):
    f=child(rpr,'w:rFonts')
    for key in ('ascii','hAnsi','cs'):f.set(qn('w:'+key),'Times New Roman')
    f.set(qn('w:eastAsia'),'宋体')
    for key in list(f.attrib):
        if 'Theme' in key:del f.attrib[key]

def replace_text(p,old,new):
    """Replace exact visible ordinary text, preserving all surrounding nodes."""
    texts=[x for x in p.iter(qn('w:t'))]
    total=''.join(x.text or '' for x in texts)
    assert total.count(old)==1,(old,total)
    start=total.index(old);end=start+len(old);offset=0;written=False
    for t in texts:
        s=t.text or '';a=offset;b=offset+len(s);offset=b
        if b<=start or a>=end:continue
        left=s[:max(0,start-a)];right=s[max(0,end-a):] if end<b else ''
        t.text=left+(new if not written else '')+right;written=True

def ordinary_style(p,symbol_cell=False,unit_cell=False):
    runs=[r for r in p.iter(qn('w:r')) if r.find(qn('w:t')) is not None and all(c.tag in (qn('w:rPr'),qn('w:t')) for c in r)]
    total=''.join(''.join(t.text or '' for t in r.findall(qn('w:t'))) for r in runs)
    marks=[{} for _ in total]
    def mark(a,b,**values):
        for i in range(a,b):marks[i].update(values)
    # Units are upright. Single-symbol prose recognition avoids Latin words.
    variables=set('TCDRBNAxrtuwpabkhmgMIzqJy')
    if not unit_cell:
        for match in re.finditer(r'(?<![A-Za-z])([TCDRBNAxrtuwpabkhmgMIzqJy])(?![A-Za-z])|[α-ωΑ-Ωℓ]',total):
            token=match.group();before=total[max(0,match.start()-12):match.start()]
            is_unit=token in 'mh' and bool(re.search(r'[0-9⁰¹²³⁴⁵⁶⁷⁸⁹]\s*$',before))
            if not is_unit:mark(*match.span(),italic=True,vert='baseline')
        patterns=[(r'ρ(eff|d0|d)',False),(r'(?<![A-Za-z])C(ref|eq|s|0)',False),
                  (r'(?<![A-Za-z])T(∞|s|0)',False),(r'(?<![A-Za-z])R(0)',False),
                  (r'(?<![A-Za-z])Y(∞)',False),(r'(?<![A-Za-z])D([0134])',False),
                  (r'(?<![A-Za-z])c(p)',False),(r'(?<![A-Za-z])q(r)',False),
                  (r'(?<![A-Za-z])j(w)',False),(r'(?<![A-Za-z])t(c|r)',False),
                  (r'(?<![A-Za-z])[xw](i)',True),(r'(?<![A-Za-z])K(p)',True),
                  (r'(?<![A-Za-z])[TC](r|x)',True)]
        for pattern,it in patterns:
            for m in re.finditer(pattern,total):
                if m.end()<len(total) and total[m.end()].islower() and total[m.end()].isascii():continue
                mark(m.start(),m.start()+1,italic=True,vert='baseline')
                mark(m.start(1),m.end(1),italic=it,vert='subscript')
        for m in re.finditer(r'(?<![A-Za-z])N(?=\d)',total):mark(*m.span(),italic=True,vert='baseline')
        for m in re.finditer(r'10(-14)',total):mark(m.start(1),m.end(1),italic=False,vert='superscript')
        for m in re.finditer(r'(?<![A-Za-z])t(\*)',total):
            mark(m.start(),m.start()+1,italic=True,vert='baseline');mark(m.start(1),m.end(1),italic=False,vert='superscript')
    # Scientific unit strings can contain letters which are variables elsewhere.
    for m in re.finditer(r'(?:kg/(?:kg|m[²³]?)|kg|cm|m[²³]?/s|[JW]/\([^）)]*\)|W/m²|°C|℃)',total):
        mark(*m.span(),italic=False,vert='baseline')
    for m in re.finditer(r'(?<=/)h\b',total):
        mark(*m.span(),italic=False,vert='baseline')
    if unit_cell:
        for i in range(len(total)):marks[i]={'italic':False}
    pos=0
    for r in runs:
        text=''.join(t.text or '' for t in r.findall(qn('w:t')));parent=r.getparent();at=parent.index(r)
        chunks=[]
        for j,ch in enumerate(text):
            flags=marks[pos+j]
            if not chunks or chunks[-1][1]!=flags:chunks.append([ch,dict(flags)])
            else:chunks[-1][0]+=ch
        pos+=len(text)
        for chunk,flags in chunks:
            fresh=deepcopy(r)
            for t in list(fresh.findall(qn('w:t'))):fresh.remove(t)
            pr=child(fresh,'w:rPr');font(pr)
            if 'italic' in flags:
                wprop(pr,'w:i','1' if flags['italic'] else '0');wprop(pr,'w:iCs','1' if flags['italic'] else '0')
            if 'vert' in flags:wprop(pr,'w:vertAlign',flags['vert'])
            t=OxmlElement('w:t');t.text=chunk;t.set(qn('xml:space'),'preserve');fresh.append(t)
            parent.insert(at,fresh);at+=1
        parent.remove(r)

FUNCTIONS={'exp','Ei','max','min','lim','ln','log','sin','cos','ceil','Bi','rtol','atol'}
def mathematical_style(doc):
    count=0
    for math in doc.element.body.iter(qn('m:oMath')):
        # Keep multi-letter descriptive suffixes and named operators as one
        # normal-text math run; LibreOffice otherwise spaces their letters.
        for sub in math.iter(qn('m:sub')):
            suffix=visible(sub)
            if suffix in {'eff','eq','ref','rep','rep,h','d0'} and all(x.tag==qn('m:r') for x in sub):
                replacement=deepcopy(sub[0])
                for t in list(replacement.findall(qn('m:t'))):replacement.remove(t)
                t=OxmlElement('m:t');t.text=suffix;replacement.append(t)
                for x in list(sub):sub.remove(x)
                sub.append(replacement)
        for parent in list(math.iter()):
            nodes=list(parent)
            for i in range(len(nodes)-1):
                a,b=nodes[i],nodes[i+1]
                if a.getparent() is parent and b.getparent() is parent and a.tag==b.tag==qn('m:r') and visible(a)+visible(b) in {'Ei','Bi'}:
                    value=visible(a)+visible(b)
                    for t in list(a.findall(qn('m:t'))):a.remove(t)
                    t=OxmlElement('m:t');t.text=value;a.append(t);parent.remove(b)
        for run in list(math.iter(qn('m:r'))):
            text=''.join(t.text or '' for t in run.findall(qn('m:t')))
            if not text:continue
            original_sty=run.find(qn('m:rPr'))
            sty=original_sty.find(qn('m:sty')) if original_sty is not None else None
            was_plain=sty is not None and sty.get(qn('m:val')) in ('p','b')
            was_bold=sty is not None and sty.get(qn('m:val')) in ('bi','b')
            desc=False
            for ancestor in run.iterancestors():
                if ancestor.tag==qn('m:sub'):
                    subtext=visible(ancestor)
                    base=visible(ancestor.getparent().find(qn('m:e')))
                    desc=subtext in {'eff','eq','ref','rep','rep,h','d0','d','w','s','c','H','h'} or (subtext=='p' and base=='c') or (subtext=='r' and base in {'t','q'})
                    break
                if ancestor.tag==qn('m:oMath'):break
            tokens=re.findall(r'exp|Ei|max|min|lim|ceil|log|ln|sin|cos|rtol|atol|Bi|[A-Za-zα-ωΑ-Ωℓ]|[^A-Za-zα-ωΑ-Ωℓ]+',text)
            if desc:tokens=[text]
            parent=run.getparent();at=parent.index(run)
            for token in tokens:
                italic=not desc and not was_plain and token not in FUNCTIONS and bool(re.fullmatch(r'[A-Za-zα-ωΑ-Ωℓ]',token))
                # Upright physical units inside display formulas, based on their original plain runs or explicit trailing unit position.
                if token in {'K','s'} and (was_plain or text.strip() in {'K','s','s.'} or re.search(r'\d\s*'+token+r'\b',text)):italic=False
                if token=='m' and text.strip()=='m':
                    # m as the base of m²/s in Q1 diffusion units.
                    par=run.getparent()
                    if par.tag==qn('m:e') and par.getparent().tag==qn('m:sSup') and visible(par.getparent().find(qn('m:sup')))=='2':italic=False
                if token=='K' and run.getparent().tag==qn('m:e'):
                    expression=run.getparent().getparent()
                    if expression.tag==qn('m:sSub') and visible(expression.find(qn('m:sub')))=='p':
                        italic=True
                fresh=deepcopy(run)
                for t in list(fresh.findall(qn('m:t'))):fresh.remove(t)
                mp=child(fresh,'m:rPr');ms=child(mp,'m:sty');ms.set(qn('m:val'),('bi' if italic else 'b') if was_bold else ('i' if italic else 'p'))
                if not italic and (desc or bool(re.search(r'[A-Za-z]',token)) or token in {'⌈','⌉','⌊','⌋'}):
                    child(mp,'m:nor').set(qn('m:val'),'1')
                elif italic:
                    for normal in list(mp.findall(qn('m:nor'))):mp.remove(normal)
                wp=child(fresh,'w:rPr');font(wp);wprop(wp,'w:i','1' if italic else '0');wprop(wp,'w:iCs','1' if italic else '0')
                if was_bold:wprop(wp,'w:b','1')
                t=OxmlElement('m:t');t.text=token;t.set(qn('xml:space'),'preserve');fresh.append(t)
                parent.insert(at,fresh);at+=1;count+=1
            parent.remove(run)
    return count

def heading(d,text,level):
    p=d.add_paragraph(style=f'Heading {min(level,3)}');p.add_run(text)
    p.paragraph_format.first_line_indent=Cm(0)
    if text=='附录':p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True

def appendix_body(d,text,**kwargs):
    result=layout.body(d,text,**kwargs)
    if text.startswith('来源：'):
        p=d.paragraphs[-1]
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1.1
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.font.size=Pt(10.5)
    return result

code_lines=0
def code(d,text):
    global code_lines
    lines=[line for line in text.splitlines() if line.strip()]
    for i,line in enumerate(lines):
        p=d.add_paragraph();p._p.set(qn('w:rsidR'),'00000002')
        pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.space_before=pf.space_after=Pt(0);pf.line_spacing=Pt(11.5)
        pf.keep_with_next=i<len(lines)-1;pf.widow_control=True
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9.5)
        code_lines+=1

def main():
    d=Document(SOURCE);body=d.element.body
    boundary=next(i for i,p in enumerate(body) if visible(p).strip()=='附录')
    original_blocks=[visible(p) for p in body[:boundary]]
    for p in list(body)[boundary:]:
        if p.tag!=qn('w:sectPr'):body.remove(p)
    changes=[];comments_to_add=[]
    # Latest explicit reply: retain the deleted bibliography entries and ONLY
    # comment on unmatched in-text citations, preserving the user's body text.
    for p in d.paragraphs:
        for n in (3,4,5,6):
            if f'[{n}]' in p.text:
                changes.append({'kind':'unmatchedCitationComment','citation':f'[{n}]','bodyTextPreserved':True})
                comments_to_add.append((p,f'正文引用[{n}]在当前参考文献表中没有对应条目。按你的最新要求维持已删除条目的状态，正文此处仅加批注，未恢复或补写文献。请团队在定稿时核查处理。'))
    cell=d.tables[0].cell(15,0)
    if '··' in cell.text:
        replace_text(cell.paragraphs[0]._p,'··','·');changes.append({'kind':'typo','before':'B=ρeff··cp','after':'B=ρeff·cp'})
    figs=[]
    current_section=''
    for p in d.paragraphs:
        if re.match(r'^5\.\d+\.\d+\s',p.text):current_section=p.text
        m=re.match(r'图像主题([1-6])：',p.text)
        if m:
            n=int(m.group(1));old=m.group(0);new=f'【绘图位置{n}】图{n}：'
            replace_text(p._p,old,new)
            p.paragraph_format.keep_together=True;p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(7)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
            if p.runs:p.runs[0].bold=True
            figs.append({'number':n,'section':current_section,'text':p.text})
            changes.append({'kind':'figureMarker','before':old,'after':new})
    assert len(figs)==6
    # Mark descriptive suffixes and variables in ordinary prose and table cells.
    for p in d.paragraphs:
        if not p.style.name.startswith(('Heading','Title')):ordinary_style(p._p)
    for ti,tb in enumerate(d.tables):
        for ri,row in enumerate(tb.rows):
            for ci,c in enumerate(row.cells):
                for p in c.paragraphs:ordinary_style(p._p,symbol_cell=ti==0 and ci==0,unit_cell=ti==0 and ci==2)
    # Retire comments already resolved by the user's edits, and update the
    # remaining partial resolutions without touching the user's prose.
    retired=[]
    for comment in list(d.comments):
        note=comment.text
        cid=comment._element.get(qn('w:id'))
        replacement=None
        if '此处“长2cm”应改为“长25cm”' in note or '原文共有8名作者，前三人之后宜补“等”' in note:
            retired.append(cid);comment._element.getparent().remove(comment._element)
        elif '“BD”应为“BDF”' in note:
            replacement='你已将BD改为BDF，本次保留。此段“精确处理”仍应结合后文理解：Kirchhoff势仅解析处理浓度方向积分，整个数值通量及方程求解仍有离散误差。分钟量级为已记录环境表现。'
        elif 'Q4材料骨架随径向收缩运动' in note:
            replacement='你已补充干物质没有生成、损失或相对骨架迁移的限定，本次保留。句首“骨架不发生迁移”仍需与Q4材料整体收缩运动区别理解；后文按材料随动建立方程。'
        if replacement is not None:
            texts=list(comment._element.iter(qn('w:t')))
            for i,t in enumerate(texts):t.text=replacement if i==0 else ''
    for node in list(d.element.body.iter()):
        if node.tag in (qn('w:commentRangeStart'),qn('w:commentRangeEnd'),qn('w:commentReference')) and node.get(qn('w:id')) in retired:
            node.getparent().remove(node)
    for p,note in comments_to_add:d.add_comment(p.runs,text=note,author='审阅',initials='')
    appendix=QA/'compact_appendix_for_paper.md'
    if '--body-only' not in sys.argv:
        assert appendix.exists(),'Compact appendix is not ready.'
        layout.equation_index=19;layout.fmt.add_heading=heading;layout.fmt.add_code_block=code
        layout.fmt.add_body_paragraph=appendix_body
        stats=layout.fmt.render_markdown(d,appendix.read_text(encoding='utf-8'),{}, {})
        assert not stats['formula_fallbacks'],stats['formula_fallbacks']
    math_runs=mathematical_style(d)
    for p in d.paragraphs:
        child(p._p.get_or_add_pPr(),'w:snapToGrid').set(qn('w:val'),'0')
    # Preserve all original opaque package parts; only XML actually edited is transplanted.
    temp=QA/'edited_package.docx';d.save(temp)
    target=QA/'styled_body.docx' if '--body-only' in sys.argv else OUT
    with ZipFile(SOURCE) as original,ZipFile(temp) as generated,ZipFile(target,'w') as result:
        for info in original.infolist():
            data=generated.read(info.filename) if info.filename in {'word/document.xml','word/comments.xml'} else original.read(info.filename)
            if info.filename=='docProps/core.xml':
                core=etree.fromstring(data)
                for x in core:
                    if etree.QName(x).localname in ('creator','lastModifiedBy'):x.text=''
                data=etree.tostring(core,xml_declaration=True,encoding='UTF-8',standalone=True)
            result.writestr(info,data)
    report={'source':str(SOURCE.relative_to(ROOT)),'sourceSha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
       'output':str(target.relative_to(ROOT)),'outputSha256':hashlib.sha256(target.read_bytes()).hexdigest(),
       'changes':changes,'figureBriefs':figs,'formattedMathRuns':math_runs,'codeLines':code_lines,
       'preservedOpaqueParts':'All source ZIP parts except document.xml, comments.xml and anonymized core.xml copied byte-for-byte.'}
    (QA/('body_style_build.json' if '--body-only' in sys.argv else 'revision_build.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':report['output'],'codeLines':code_lines,'mathRuns':math_runs,'sha256':report['outputSha256']},ensure_ascii=False))

if __name__=='__main__':main()
