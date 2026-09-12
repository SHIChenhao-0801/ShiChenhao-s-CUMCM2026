"""Independent read-only verification of the exported workbook and CSV values."""
from __future__ import annotations
import csv
import hashlib
import json
import math
from pathlib import Path
import zipfile
import openpyxl

ROOT = Path.cwd().resolve()
assert ROOT.as_posix().endswith('/数学建模/2026CUMCM')
QA = ROOT / 'paper_output/qa/figure_package_20260912'
PACKAGE = ROOT / 'paper_output/handoff/six_figures_20260912'
INPUT = ROOT / 'tmp/cache/figure_package_20260912/workbook_data.json'
XLSX = PACKAGE / '六图绘图数据.xlsx'
data = json.loads(INPUT.read_text(encoding='utf-8'))
report = {'status': 'CHECKING', 'numeric_relative_tolerance': 2e-14, 'sheets': [], 'errors': []}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()

def compare(expected, actual, context):
    if isinstance(expected, (float, int)) and not isinstance(expected, bool):
        if not isinstance(actual, (float, int)) or isinstance(actual, bool):
            report['errors'].append(f'{context}: numeric type lost')
            return
        if not math.isfinite(actual) or abs(actual - expected) > 2e-14 * abs(expected):
            report['errors'].append(f'{context}: number differs: {expected!r}, {actual!r}')
    elif expected in (None, '') and actual in (None, ''):
        return
    elif expected != actual:
        report['errors'].append(f'{context}: value differs')

with zipfile.ZipFile(XLSX) as archive:
    assert archive.testzip() is None
    report['zipCRC'] = 'PASS'
wb = openpyxl.load_workbook(XLSX, data_only=False, read_only=False)
assert wb.sheetnames == [s['name'] for s in data['sheets']]
for spec in data['sheets']:
    ws = wb[spec['name']]
    cols = spec['columns']
    assert [ws.cell(6, i + 1).value for i in range(len(cols))] == cols, spec['name']
    cells_checked = 0
    for ri, row in enumerate(spec['rows'], 7):
        for ci, expected in enumerate(row, 1):
            cell = ws.cell(ri, ci)
            compare(expected, cell.value, f'{spec["name"]}!{cell.coordinate}')
            assert cell.data_type != 'f', f'Unexpected formula {spec["name"]}!{cell.coordinate}'
            if isinstance(expected, (float, int)):
                assert cell.data_type == 'n', f'Not numeric: {spec["name"]}!{cell.coordinate}'
            cells_checked += 1
    if len(spec['rows']) > 20:
        assert ws.freeze_panes == 'A7', (spec['name'], ws.freeze_panes)
    assert not ws.merged_cells.ranges, (spec['name'], 'unexpected merges')
    assert ws.sheet_view.showGridLines is False, spec['name']
    references = spec.get('csv', [])
    references = [references] if isinstance(references, str) else references
    csv_reports = []
    matched = False
    for relative in references:
        csv_path = PACKAGE / relative
        assert csv_path.is_file(), csv_path
        with csv_path.open(encoding='utf-8-sig', newline='') as f:
            records = list(csv.reader(f))
        if records and records[0] == cols and len(records) - 1 == len(spec['rows']):
            for ri, (actualrow, expectedrow) in enumerate(zip(records[1:], spec['rows']), 2):
                assert len(actualrow) == len(cols), (csv_path, ri)
                for ci, (actual, expected) in enumerate(zip(actualrow, expectedrow), 1):
                    value = float(actual) if isinstance(expected, (float, int)) and not isinstance(expected, bool) else (None if expected is None and actual == '' else actual)
                    compare(expected, value, f'{relative}:{ri}:{ci}')
            matched = True
            csv_reports.append({'file': relative, 'rows': len(records) - 1, 'SHA256': sha(csv_path), 'wholeTableCompared': True})
        else:
            csv_reports.append({'file': relative, 'rows': max(0, len(records) - 1), 'SHA256': sha(csv_path), 'wholeTableCompared': False})
    assert matched, f'No corresponding full CSV table for {spec["name"]}'
    report['sheets'].append({'name': spec['name'], 'rows': len(spec['rows']), 'columns': len(cols), 'cellsCompared': cells_checked, 'CSV': csv_reports, 'freezePanes': ws.freeze_panes})
report['cellCount'] = sum(s['cellsCompared'] for s in report['sheets'])
report['sheetCount'] = len(report['sheets'])
report['rowCount'] = sum(s['rows'] for s in report['sheets'])
report['xlsxSHA256'] = sha(XLSX)
report['status'] = 'PASS' if not report['errors'] else 'FAIL'
report['visualQA'] = 'PENDING_VIEW_IMAGE'
(QA / 'workbook_independent_audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: report[k] for k in ['status', 'sheetCount', 'rowCount', 'cellCount', 'xlsxSHA256']}, ensure_ascii=False))
if report['errors']:
    raise SystemExit(1)
