from pathlib import Path
import copy
import csv
import hashlib
import json
import re
import sys
import unicodedata

ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(ROOT / '.agents/skills/paper-formal-writer/scripts'))
sys.path.insert(0, str(ROOT / 'tmp/cache/manuscript_20260912/python_deps'))
from formula_omml import latex_to_omml, inline_formula_tokens
import formula_omml
from latex2mathml.converter import convert as latex_mathml
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

QA = ROOT / 'paper_output/qa/appendix_full_20260913'
OUT = ROOT / 'paper_output/paper/药材热湿耦合模型与干燥时间计算_完整附录.docx'
MANIFEST = QA / 'source_inventory/source_manifest.json'
SOURCE = ROOT / 'paper_output/paper/药材热湿耦合模型与干燥时间计算_格式与语言修订版.docx'
source_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
assert source_sha == '2fc45b6bbc7aff10db0de198e47f5595b83f8990969cd7258353fa31ddfbe44a'
manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
entries = manifest['entries']
d = Document()
bookmarks = []
code_index = []
headings = []
equations = []
eq_map = {}
toc_entries = []
page_map_path = QA / 'toc_page_map.json'
page_map = json.loads(page_map_path.read_text(encoding='utf-8')) if page_map_path.exists() else {}
bookmark_id = 0

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def font(run, size=None, bold=None, mono=False):
    run.font.name = 'Consolas' if mono else 'Times New Roman'
    rpr = run._r.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), '宋体')
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return run

def bookmark(p, name):
    global bookmark_id
    bookmark_id += 1
    a, b = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    a.set(qn('w:id'), str(bookmark_id)); a.set(qn('w:name'), name)
    b.set(qn('w:id'), str(bookmark_id))
    p._p.append(a); p._p.append(b)
    bookmarks.append(name)

def field(p, instruction, cached, bookmark_name=None):
    global bookmark_id
    if bookmark_name:
        bookmark_id += 1
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(bookmark_id)); start.set(qn('w:name'), bookmark_name)
        p._p.append(start)
        bid = bookmark_id
        bookmarks.append(bookmark_name)
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), instruction)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = str(cached)
    r.append(t); f.append(r); p._p.append(f)
    if bookmark_name:
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bid)); p._p.append(end)

def math(latex):
    latex = latex.replace(r'\lim_', r'\mathrm{lim}_')
    if r'\begin{aligned}' in latex:
        arr = OxmlElement('m:eqArr')
        for row in latex.replace(r'\begin{aligned}', '').replace(r'\end{aligned}', '').split(r'\\'):
            row = row.replace('&', '').strip()
            if row:
                e = OxmlElement('m:e')
                for child in list(convert_math(row)):
                    e.append(child)
                arr.append(e)
        om = OxmlElement('m:oMath'); om.append(arr)
    else:
        om = convert_math(latex)
    for sub in om.iter(qn('m:sub')):
        text = ''.join(sub.itertext())
        if text in {'eff', 'eq', 'ref', 'end', 'out', 'obs', 'RMS'} and all(e.tag == qn('m:r') for e in sub):
            for e in list(sub):
                sub.remove(e)
            sub.append(formula_omml._text_run(text, plain=True))
    for r in om.iter(qn('m:r')):
        rp = r.find(qn('m:rPr'))
        if ''.join(r.itertext()) in {'⌈', '⌉', '⌊', '⌋'}:
            if rp is None:
                rp = OxmlElement('m:rPr'); r.insert(0, rp)
            upright = rp.find(qn('m:sty'))
            if upright is None:
                upright = OxmlElement('m:sty'); rp.append(upright)
            upright.set(qn('m:val'), 'p')
        sty = rp.find(qn('m:sty')) if rp is not None else None
        if sty is not None and sty.get(qn('m:val')) in ('p', 'b'):
            normal = OxmlElement('m:nor'); normal.set(qn('m:val'), '1'); rp.append(normal)
            wp = OxmlElement('w:rPr')
            fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Times New Roman'); fonts.set(qn('w:hAnsi'),'Times New Roman')
            wp.append(fonts)
            italic = OxmlElement('w:i'); italic.set(qn('w:val'),'0'); wp.append(italic)
            r.insert(1, wp)
    return om

def convert_math(latex):
    root = etree.fromstring(latex_mathml(formula_omml.normalize_latex(latex)).encode())
    def propagate(node, inherited=None):
        variant = node.get('mathvariant', inherited)
        if etree.QName(node).localname in ('mi', 'mn', 'mo', 'mtext') and variant:
            node.set('mathvariant', variant)
            if variant == 'script' and node.text and len(node.text) == 1 and node.text.isascii() and node.text.isalpha():
                names = {'B':'ℬ','E':'ℰ','F':'ℱ','H':'ℋ','I':'ℐ','L':'ℒ','M':'ℳ','R':'ℛ','e':'ℯ','g':'ℊ','o':'ℴ'}
                c = node.text
                try:
                    node.text = names[c] if c in names else unicodedata.lookup('MATHEMATICAL SCRIPT '+('CAPITAL ' if c.isupper() else 'SMALL ')+c.upper())
                except KeyError:
                    pass
        for child in node:
            propagate(child, variant)
    propagate(root)
    out = OxmlElement('m:oMath')
    for child in formula_omml._convert_node(root):
        out.append(child)
    return out

def inline_plain(p, text, size=12):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1（\2）', text)
    pieces = re.split(r'(`[^`]*`)', text)
    for piece in pieces:
        if piece.startswith('`') and piece.endswith('`'):
            font(p.add_run(piece[1:-1]), size, mono=True)
        else:
            font(p.add_run(piece), size)

def inline(p, text, size=12):
    parts = re.split(r'(\[EQ_A_[A-Za-z0-9_]+\])', text)
    for part in parts:
        if re.fullmatch(r'\[EQ_A_[A-Za-z0-9_]+\]', part):
            key = part[1:-1]
            assert key in eq_map, key
            field(p, f' REF {key} \\h ', f'（A.{eq_map[key]}）')
            continue
        cursor = 0
        for tok in inline_formula_tokens(part):
            inline_plain(p, part[cursor:tok.start], size)
            p._p.append(math(tok.latex)); cursor = tok.end
        inline_plain(p, part[cursor:], size)

def para(text='', size=12, style=None):
    p = d.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(.85)
    inline(p, text, size)
    return p

def heading(text, level=1, name=None):
    p = d.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.first_line_indent = Cm(0)
    if level == 1:
        p.paragraph_format.page_break_before = True
    font(p.add_run(text), 15 if level == 1 else 12.5, True)
    name = name or f'HEAD_{len(headings)+1:03d}'
    bookmark(p, name)
    headings.append({'text': text, 'bookmark': name, 'level': level, 'paragraph_index': len(d.paragraphs)-1})
    return p

def display(latex, key):
    global bookmark_id
    p = d.add_paragraph(style='Equation')
    p.paragraph_format.tab_stops.add_tab_stop(Cm(7.25), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t'); p._p.append(math(latex)); p.add_run('\t')
    n = eq_map[key]
    bookmark_id += 1
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id)); start.set(qn('w:name'), key)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(bookmark_id))
    p._p.append(start)
    font(p.add_run('（A.'), 12)
    field(p, f' SEQ AppendixA \\r {n} ', str(n))
    font(p.add_run('）'), 12)
    p._p.append(end)
    bookmarks.append(key)
    equations.append({'key': key, 'number': f'A.{n}', 'latex': latex, 'paragraph_index': len(d.paragraphs)-1})

def table(rows, widths=None, size=10.5):
    cols = len(rows[0]); t = d.add_table(rows=0, cols=cols)
    t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths is None and rows[0] == ['编号', '相对文件路径']:
        widths = [1.5, 14.5]
    widths = widths or [16 / cols] * cols
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for ri, values in enumerate(rows):
        cells = t.add_row().cells
        trp = cells[0]._tc.getparent().get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trp.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(cells):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:'+side); item.set(qn('w:val'), 'single')
                item.set(qn('w:sz'), '4'); item.set(qn('w:color'), 'D9D9D9'); borders.append(item)
            cp.append(borders)
            if ri == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'EEEEEE'); cp.append(fill)
            p = cell.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.12
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(values[ci])) < 25 else WD_ALIGN_PARAGRAPH.LEFT
            if rows[0] == ['编号', '相对文件路径'] and ci == 1 and ri > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p, str(values[ci]), size)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def md(text):
    lines = text.splitlines(); i = 0; pending = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        match = re.match(r'<!--\s*(EQ_A_[A-Za-z0-9_]+)\s*-->', line)
        if match:
            pending = match.group(1); i += 1; continue
        if line.startswith('<!--'):
            i += 1; continue
        if line == '$$':
            formula = []; i += 1
            while i < len(lines) and lines[i].strip() != '$$':
                formula.append(lines[i]); i += 1
            assert pending
            display('\n'.join(formula), pending); pending = None; i += 1; continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            heading(line.lstrip('#').strip(), min(level, 3)); i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\-\s]+', c) for c in cells):
                    rows.append(cells)
                i += 1
            table(rows); continue
        if line.startswith('```'):
            i += 1
            code_paragraphs = []
            while i < len(lines) and not lines[i].startswith('```'):
                p = d.add_paragraph(style='Code'); font(p.add_run(lines[i]), 9, mono=True); i += 1
                code_paragraphs.append(p)
            for p in code_paragraphs[:-1]:
                p.paragraph_format.keep_with_next = True
            i += 1; continue
        if line.startswith('- '):
            p = para('• '+line[2:]); p.paragraph_format.first_line_indent = Cm(0)
            i += 1; continue
        parts = [line]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||\$\$|<!--|```|- )', lines[i].strip()):
            parts.append(lines[i].strip()); i += 1
        combined = ''.join(parts)
        p = para(combined)
        if re.search(r'（\d+份）$', combined):
            p.paragraph_format.keep_with_next = True
        if 'REVIEW_REQUIRED' in combined or 'export_six_figure_package_20260912.py' in combined:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def code(rec, number):
    path = Path(rec['appendix_source_path']); raw = path.read_bytes()
    assert sha(path) == rec['sha256'], str(path)
    text = raw.decode(rec.get('source_text_encoding', 'utf-8-sig'))
    lines = text.splitlines()
    hp = heading(f'D.{number} {Path(rec["relative_path"]).name}', 2, 'FILE_'+rec['id'])
    hp.paragraph_format.page_break_before = True
    for label, val in [('路径', rec['display_path']), ('用途', rec['purpose']), ('使用范围', rec['adoption_scope'])]:
        p = para(f'{label}：{val}', 10.5); p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(3)
    p = para(f'语言：{rec["language"]}；源文件共 {len(lines)} 行。', 9)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.space_after = Pt(2)
    p = para(f'SHA256：{rec["sha256"]}', 9)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.keep_with_next = True
    first = len(d.paragraphs)
    last_function_start = len(lines)
    if rec['id'] in ('SRC008', 'SRC017'):
        last_function_start = max(i for i, line in enumerate(lines) if line.startswith('def '))
    for idx, line in enumerate(lines):
        p = d.add_paragraph(style='Code')
        if last_function_start <= idx < len(lines)-1:
            p.paragraph_format.keep_with_next = True
        if idx == 0:
            bookmark(p, 'CODE_'+rec['id']+'_BEGIN')
        font(p.add_run(line), 9, mono=True)
        if idx == len(lines)-1:
            bookmark(p, 'CODE_'+rec['id']+'_END')
    code_index.append({**rec, 'start_paragraph_index': first, 'end_paragraph_index': len(d.paragraphs)-1, 'code_line_count': len(lines), 'section': f'D.{number}'})

s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
s.header_distance, s.footer_distance = Cm(1), Cm(1.25)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Code', 'Equation', 'TOC Entry'):
    if name not in d.styles:
        d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st = d.styles[name]
    st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    for old in list(st._element.iter(qn('w:pBdr'))):
        old.getparent().remove(old)
    pf = st.paragraph_format; pf.line_spacing = 1.25; pf.space_after = Pt(5)
    pf.widow_control = True
    if name.startswith('Heading'):
        pf.keep_with_next = True; pf.space_before = Pt(12); pf.space_after = Pt(7)
    if name == 'Code':
        st.font.name = 'Consolas'; st.font.size = Pt(9)
        pf.line_spacing = Pt(11); pf.space_after = pf.space_before = Pt(0)
        pf.keep_together = pf.keep_with_next = pf.widow_control = False
        pf.first_line_indent = pf.left_indent = pf.right_indent = Cm(0)
        wrap = OxmlElement('w:wordWrap'); wrap.set(qn('w:val'),'0'); st._element.get_or_add_pPr().append(wrap)
    if name == 'Equation':
        pf.line_spacing = 1.1; pf.space_before = Pt(5); pf.space_after = Pt(7)
        pf.keep_together = True; pf.first_line_indent = Cm(0)
    if name == 'TOC Entry':
        st.font.size = Pt(10.5); pf.line_spacing = 1.05; pf.space_after = Pt(3)
        pf.first_line_indent = Cm(0)
foot = s.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(foot, ' PAGE ', '1')
d.core_properties.author = ''; d.core_properties.last_modified_by = ''
d.core_properties.title = '药材热湿耦合模型与干燥时间计算完整附录'
d.core_properties.subject = '公式推导 源程序 支撑材料与计算核查'
d.core_properties.comments = ''

derivations = (QA / 'derivations.md').read_text(encoding='utf-8')
keys = re.findall(r'<!--\s*(EQ_A_[A-Za-z0-9_]+)\s*-->', derivations)
assert len(keys) == len(set(keys))
eq_map = {key: idx for idx, key in enumerate(keys, 1)}
title = d.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(title.add_run('药材热湿耦合模型与干燥时间计算'), 18, True)
title2 = d.add_paragraph(); title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(title2.add_run('完整附录'), 16, True)
para('本附录给出固定半径与径向收缩模型的完整推导、离散与求解方法、支撑材料目录、计算结果及源程序。公式采用独立编号，源程序按文件完整列出；四问的全量径向与逐秒结果保存在所列 Excel 工作簿中。')
para('物性关系使用题目给定经验式，计算环境与复现结论对应具体代码版本。检验与探索程序按实际使用范围列示，其历史试验不均属于最终模型。')
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run('附录目录'), 14, True)
toc_position = len(d.paragraphs)
toc_placeholder = d.add_paragraph('TOC_PLACEHOLDER')
if not derivations.lstrip().startswith('# 附录 A'):
    heading('附录 A 模型与数值方法完整推导')
md(derivations)
md((QA / 'source_inventory/support_content.md').read_text(encoding='utf-8'))
heading('附录 C 结果表与数值精度核查')
heading('C.1 临界事件与严格达标时刻', 2)
para('下表保留事件根和报告时刻的未舍入查询结果。报告时刻对应 0.0001 h 的时间网格，并逐点检查最大含水率严格小于 0.15 kg/kg；显示为 0.1500 的四位浓度单元格不作为等于阈值的判据。全域判定与离散最大值的关系见附录 A。')
table([['量', '问题三', '问题四'], ['事件根 / h', '57.47230195056044', '51.09057478683054'], ['报告烘干时长 / h', '57.4724', '51.0906'], ['报告时刻最大 C / (kg/kg)', '0.14999989718225315', '0.14999995339076624'], ['径向区间数 N', '3200', '6400']], [5.4,5.3,5.3])
heading('C.2 题目要求的径向采样结果', 2)
para('表 C.1 至表 C.7 按已冻结 CSV 原样显示，时间、位置和单位与正文结果表对应。完整结果工作簿另见附录 B 的文件清单。径向含水率均为干基含水率；径向列名中的 0、0.5、1、1.5、2 表示物理半径，单位为 cm。问题四的“药材表面”列查询当时真实表面 r=R(t)，位于当时材料范围之外的固定半径留空。')
tables = [
('q1_paper_temperature.csv','问题一温度（°C）'), ('q1_paper_moisture.csv','问题一干基含水率（kg/kg）'),
('q2_paper_temperature.csv','问题二温度（°C）'), ('q2_paper_moisture.csv','问题二干基含水率（kg/kg）'),
('q3_paper_moisture.csv','问题三干基含水率（kg/kg）'), ('q4_paper_moisture.csv','问题四干基含水率（kg/kg）'),
('q4_paper_radius.csv','问题四半径变化')]
result_tables = []
for idx, (filename, caption) in enumerate(tables, 1):
    path = ROOT / '支撑材料/04_结果表格/Referrence Table Files' / filename
    rows = list(csv.reader(path.open(encoding='utf-8-sig', newline='')))
    p = para(f'表 C.{idx} {caption}', 11); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.keep_with_next = True
    if idx == 7:
        p.paragraph_format.page_break_before = True
    table(rows)
    result_tables.append({'path': str(path.relative_to(ROOT)), 'sha256': sha(path), 'rows': rows})
heading('C.3 全量导出及一致性检查', 2)
para('本轮独立 Windows Sandbox 正式计算使用 Q1、Q23 的 N=3200 和 Q4 的 N=6400，内部求解与导出记录耗时 1810.9893 s。四份工作簿合计核查 9,335,598 个数据单元格，另对正文 297 个结果单元格进行实际轨迹查询核对。工作簿全量回读结合逐行 CSV 对照，正文采样由求解轨迹直接查询；两项核查覆盖方式不同。三条保存轨迹的 21 个数组与冻结参照逐元素相同。')
para('数值复现说明指定代码、输入与环境能够再现冻结结果，并不单独证明物理模型的实验预测精度。与离散方法相关的基准、网格和容差检验应在相应实际执行范围内解释，不能将历史检验脚本的成功退出直接等同于全部检验条件通过。')
heading('附录 D 完整无注释源程序')
para(f'本附录按文件列出 {len(entries)} 份完整源码，共 {sum(e["lines"] for e in entries):,} 行。源码正文不含解释性注释或 Python 文档字符串；文件路径、用途与采用范围在源码块前列明。长行仅由 Word 自动折行，行内字符串和缩进保持原文件内容。')
para('D 中同时保留正式求解程序和实际使用过的检验、对照及绘图程序。历史失败或仅在有限设置下运行的程序不作为已通过的模型验证证据，具体范围见附录 B。源码中的原有路径常量按原文件保留；复现入口及路径适配方式按附录 B 执行。')
rows = [['小节', '程序文件', '语言', '行数']]
for i, rec in enumerate(entries, 1):
    rows.append([f'D.{i}', Path(rec['relative_path']).name, rec['language'], rec['lines']])
table(rows, [1.5, 9.2, 3.6, 1.7], 9)
for i, rec in enumerate(entries, 1):
    code(rec, i)

selected = [h for h in headings if h['level'] == 1 or (h['level'] == 2 and not h['text'].startswith('D.'))]
for item in selected:
    p = d.add_paragraph(style='TOC Entry')
    if item['level'] == 2:
        p.paragraph_format.left_indent = Cm(.45)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), item['bookmark'])
    rr=OxmlElement('w:r'); tt=OxmlElement('w:t'); tt.text=item['text']; rr.append(tt); link.append(rr); p._p.append(link)
    p.add_run('\t')
    page = page_map.get(item['bookmark'], '')
    field(p, f' PAGEREF {item["bookmark"]} \\h ', page)
    toc_placeholder._p.addprevious(p._p)
    toc_entries.append({**item, 'cached_page': page})
toc_placeholder._p.getparent().remove(toc_placeholder._p)
shift = len(selected)-1
for rec in code_index:
    rec['start_paragraph_index'] += shift
    rec['end_paragraph_index'] += shift
for rec in headings + equations:
    rec['paragraph_index'] += shift
for p in d.paragraphs:
    pr=p._p.get_or_add_pPr(); snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pr.append(snap)
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); d.settings.element.append(update)
d.save(OUT)
(QA / 'code_index.json').write_text(json.dumps({'docx':str(OUT),'docx_sha256':sha(OUT),'entries':code_index},ensure_ascii=False,indent=2),encoding='utf-8')
(QA / 'build_manifest.json').write_text(json.dumps({'docx':str(OUT),'docx_sha256':sha(OUT),'source_paper_sha256':source_sha,'source_count':len(entries),'source_lines':sum(e['lines'] for e in entries),'paragraphs':len(d.paragraphs),'equations':equations,'headings':headings,'toc':toc_entries,'result_tables':result_tables,'bookmarks':bookmarks,'omml_count':len(d.element.xpath('//m:oMath'))},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT),'bytes':OUT.stat().st_size,'sources':len(entries),'code_lines':sum(e['lines'] for e in entries),'equations':len(equations),'omml':len(d.element.xpath('//m:oMath'))},ensure_ascii=False))
