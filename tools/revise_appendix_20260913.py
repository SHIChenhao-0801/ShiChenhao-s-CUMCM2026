from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
import hashlib, json, sys
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')
QA = ROOT / 'paper_output/qa/appendix_revision_20260913'
SOURCE = ROOT / '药材热湿耦合模型与干燥时间计算_文献公式修订版(1).docx'
OUT = ROOT / '药材热湿耦合模型与干燥时间计算_附录修订版.docx'
inventory = json.loads((QA / 'source_inventory.json').read_text(encoding='utf-8'))
files = [f for f in inventory['files'] if f['recommendation'] == 'retain']
core_order = ['runDelivery.py', 'q1Model.py', 'q2Model.py', 'q3Model.py', 'q4Model.py', 'dryingCore.py', 'analyticJacobian.py', 'diskDense.py', 'exportOutputs.py', 'reference_data.py', 'runLogged.ps1']
core = sorted([f for f in files if '/03_程序代码/' in f['relative_path']], key=lambda f: core_order.index(f['filename']))
checks = sorted([f for f in files if '/Python检验源码/' in f['relative_path']], key=lambda f: f['relative_path'])
matlab = sorted([f for f in files if '/MATLAB对照证据/' in f['relative_path']], key=lambda f: ['runCrossCheck.m', 'compareMatlab.mjs'].index(f['filename']))
files = core + checks + matlab
assert len(files) == 36
d = Document(SOURCE)
body = d.element.body
old_elements = list(body)
protected = [deepcopy(e) for e in old_elements[:470]]
sect = deepcopy(old_elements[-1])
caption_prototype = deepcopy(old_elements[74])
code_ppr = deepcopy(old_elements[555].find(qn('w:pPr')))
code_rpr = deepcopy(old_elements[555].find(qn('w:r')).find(qn('w:rPr')))
for e in old_elements[470:]:
    body.remove(e)
body.append(sect)

def prop(parent, name, attrs):
    e = OxmlElement(name)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    parent.append(e)
    return e

def font(run, size=12, bold=False, mono=False):
    run.font.name = 'Consolas' if mono else 'Times New Roman'
    rp = run._r.get_or_add_rPr()
    rp.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def paragraph(text, size=12, indent=True, keep=False):
    p = d.add_paragraph()
    f = p.paragraph_format
    f.first_line_indent = Pt(24 if indent else 0)
    f.space_before = Pt(0)
    f.space_after = Pt(5)
    f.line_spacing = 1.25
    f.keep_with_next = keep
    f.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text), size)
    return p

def heading(text, level=2, page_break=False):
    p = d.add_paragraph()
    prop(p._p.get_or_add_pPr(), 'w:pStyle', {'w:val': '11' if level == 2 else '12'})
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    p.add_run(text)
    return p

def caption(text, page_break=False):
    p = d.add_paragraph()
    pp = p._p
    pp.clear()
    pp.append(deepcopy(caption_prototype.find(qn('w:pPr'))))
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p

table_records = []
def source_table(title, group, page_break=False):
    caption(title, page_break)
    t = d.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [650, 4150, 4272]
    for col, width in zip(t.columns, widths):
        col.width = Twips(width)
    for c, text in zip(t.rows[0].cells, ['编号', '程序文件', '主要用途']):
        c.text = text
    for f in group:
        cells = t.add_row().cells
        cells[0].text = str(files.index(f) + 1)
        cells[1].text = 'reference/reference_data.py' if f['filename'] == 'reference_data.py' else f['filename']
        cells[2].text = f['description']
    pr = t._tbl.tblPr
    borders = prop(pr, 'w:tblBorders', {})
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        prop(borders, 'w:' + edge, {'w:val': 'nil'})
    for ri, row in enumerate(t.rows):
        rp = row._tr.get_or_add_trPr()
        prop(rp, 'w:cantSplit', {})
        if ri == 0:
            prop(rp, 'w:tblHeader', {})
        for ci, c in enumerate(row.cells):
            c.width = Twips(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = c._tc.get_or_add_tcPr()
            border = prop(cp, 'w:tcBorders', {})
            for edge in ['top', 'left', 'bottom', 'right']:
                visible = (ri == 0 and edge in ('top', 'bottom')) or (ri == len(t.rows) - 1 and edge == 'bottom')
                prop(border, 'w:' + edge, {'w:val': 'single' if visible else 'nil', **({'w:sz': '8', 'w:color': '000000', 'w:space': '0'} if visible else {})})
            mar = prop(cp, 'w:tcMar', {})
            for edge in ['top', 'left', 'bottom', 'right']:
                prop(mar, 'w:' + edge, {'w:w': '75' if edge in ('top', 'bottom') else '90', 'w:type': 'dxa'})
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or ci == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    font(r, 9 if ri > 0 and ci == 1 else 10.5, ri == 0, ri > 0 and ci == 1)
    table_records.append({'caption': title, 'rows': len(t.rows), 'body_index': list(body).index(t._tbl)})

heading('9.2 程序源码', page_break=True)
paragraph('以下列出核心计算、数值检验及 MATLAB 对照程序。核心计算程序位于支撑材料的“03_程序代码”；Python 检验程序位于“05_数值检验与实验/Python检验源码/paper_output/code”；MATLAB 对照程序位于“05_数值检验与实验/MATLAB对照证据”。相同计算逻辑的核心模块仅列出一份。', size=10.5)
source_table('表9 核心计算程序及用途', core)
source_table('表10 Python 数值检验程序及用途', checks[:12], page_break=True)
source_table('续表10 Python 数值检验程序及用途', checks[12:], page_break=True)
source_table('表11 MATLAB 对照程序及用途', matlab)

code_records = []
for i, f in enumerate(files, 1):
    heading(f"9.2.{i} {f['filename']}", 3, page_break=(i == 1))
    path = Path(f['path'])
    assert hashlib.sha256(path.read_bytes()).hexdigest() == f['sha256']
    lines = path.read_text(encoding='utf-8-sig').splitlines()
    first = len(body) - 1
    for line in lines:
        p = d.add_paragraph()
        p._p.append(deepcopy(code_ppr))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        wrap = p._p.find(qn('w:pPr')).find(qn('w:wordWrap'))
        wrap.set(qn('w:val'), '1')
        r = p.add_run(line)
        r._r.insert(0, deepcopy(code_rpr))
    code_records.append({'number': i, **f, 'first_body_index': first, 'paragraph_count': len(lines)})

heading('9.3 AI 使用报告', page_break=True)
paragraph('本次建模过程中，使用 Deepseek Harness 与 Deepseek v4.1 flash 组合辅助检查源代码语法并进行部分程序调试；使用 Codex 与 GPT6-Astra 组合辅助理解题目、检索参考文献、检查建模结果和整理公式。具体使用情况如下。')
heading('9.3.1 Deepseek Harness 与 Deepseek v4.1 flash', 3)
paragraph('该组合用于检查已编写源代码的语法问题，以及辅助部分程序的调试。语法检查围绕程序的语言规则展开，包括括号与引号的配对、缩进、语句结构和函数调用形式等。对程序运行中出现的报错，结合相关代码片段分析可能原因，提供定位问题和调整语句的思路，以便修改后继续检查。')
paragraph('在部分程序调试中，借助该组合解释报错信息，并辅助梳理变量、参数与函数之间的传递关系。其作用是帮助发现源代码中的具体问题、理解异常发生的位置，为后续修改提供参考。')
heading('9.3.2 Codex 与 GPT6-Astra', 3)
paragraph('该组合用于解释题目，辅助梳理各问的研究目标、已知条件、输入数据和输出要求，理解温度、干基含水率、物性参数及半径变化等量在题目中的含义和联系。参考文献搜索主要围绕药材干燥、热湿耦合传递及相关数值方法展开，用于寻找可查阅的文献线索和理论资料。')
paragraph('在建模结果检查方面，借助该组合分析数值结果与题设条件、模型方程及边界条件之间的对应关系，检查温度、含水率和干燥时间等结果的单位、变化趋势与表达是否一致。公式整理主要包括梳理推导步骤、统一符号写法、检查公式之间的衔接，以及整理公式与变量说明的对应关系，使数学表达更清楚、连贯。')

def xml_signature(e):
    return (e.tag, tuple(sorted(e.attrib.items())), e.text, e.tail, tuple(xml_signature(c) for c in e))

for a, b in zip(protected, list(body)[:470]):
    assert xml_signature(a) == xml_signature(b)
assert xml_signature(body[-1]) == xml_signature(sect)
with ZipFile(SOURCE) as zin, ZipFile(OUT, 'w') as zout:
    for item in zin.infolist():
        content = etree.tostring(d.element, xml_declaration=True, encoding='UTF-8', standalone=True) if item.filename == 'word/document.xml' else zin.read(item.filename)
        zout.writestr(item, content)
check = Document(OUT)
for rec in code_records:
    actual = [e.xpath('string(.)') for e in list(check.element.body)[rec['first_body_index']:rec['first_body_index'] + rec['paragraph_count']]]
    expected = Path(rec['path']).read_text(encoding='utf-8-sig').splitlines()
    assert actual == expected, rec['filename']
audit = {'source': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(), 'output': str(OUT), 'output_sha256': hashlib.sha256(OUT.read_bytes()).hexdigest(), 'protected_body_elements': 232, 'protected_appendix_elements': 238, 'protected_elements_equal': True, 'other_zip_parts_unchanged': True, 'source_count': len(files), 'source_lines': sum(len(Path(f['path']).read_text(encoding='utf-8-sig').splitlines()) for f in files), 'source_extraction_exact': True, 'sources': code_records, 'new_tables': table_records}
(QA / 'build_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: v for k, v in audit.items() if k not in ('sources', 'new_tables')}, ensure_ascii=False))
