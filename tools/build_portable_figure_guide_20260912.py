"""Create the portable Word instructions included in the actual-data ZIP."""
from pathlib import Path
import json, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path.cwd(); assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
OUT = ROOT / 'paper_output/handoff/six_figures_20260912'
data = json.loads((ROOT/'tmp/cache/figure_package_20260912/workbook_data.json').read_text(encoding='utf-8'))
d = Document(); sec = d.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2); sec.right_margin = Cm(2); sec.footer_distance = Cm(.75)
for name, size, bold in [('Normal',11,False),('Title',21,True),('Heading 1',16,True),('Heading 2',11,True)]:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size)
    st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if bold else '宋体')
    pf=st.paragraph_format; pf.line_spacing=1.15; pf.space_after=Pt(6)
    if name.startswith('Heading'): pf.space_before=Pt(10); pf.keep_with_next=True
for border in list(d.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='A题六图绘图说明'

def p(text, style=None, bold=False, size=None):
    par=d.add_paragraph(style=style);run=par.add_run(text);run.bold=bold
    if size: run.font.size=Pt(size)
    par.paragraph_format.widow_control=True
    return par
def h(text): return p(text,'Heading 2')
def table(headers, rows, widths):
    tb=d.add_table(rows=1,cols=len(headers));tb.autofit=False;tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(tb.columns,widths):col.width=Cm(w)
    for c,v in zip(tb.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(tb.add_row().cells,row): c.text=str(v)
    for ri,row in enumerate(tb.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if ri==0:tr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for k,v in [('top',80),('bottom',80),('left',100),('right',100)]:
                e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar); borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF1F4');pr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1.05
                for run in par.runs:run.font.size=Pt(10);run.bold=ri==0
    return tb
def card(n,title,loc):
    par=p(f'图{n} {title}','Heading 1');par.paragraph_format.page_break_before=True;par.paragraph_format.space_before=Pt(0)
    p('论文位置：'+loc,size=10)
def files(n):
    h('打开哪些数据表')
    p(f'Excel：六图绘图数据.xlsx。下列CSV均位于包内 CSV/图{n}/。',size=10)
    table(['Excel工作表','CSV文件名','数据行数'],
          [[s['name'],Path(s['csv']).name,len(s['rows'])] for s in data['sheets'] if s['figure']==n],[3.5,11,2.5])
def source(text):
    h('包内原始文件')
    p(text,size=10)

p('A题六图绘图说明','Title')
p('本资料包已包含六张图的全部作图数据。解压后，直接打开“六图绘图数据.xlsx”，按本说明选择工作表和数据列即可绘图。CSV提供同一套数值，便于导入Origin、MATLAB或其他绘图软件。')
table(['包内文件或目录','用途'],[
    ['六图绘图数据.xlsx','14张数据工作表，包含实际数值。'],
    ['CSV/图1/ 至 CSV/图6/','按图号分开的14份CSV，与Excel逐表对应。'],
    ['原始数据/','15份原始CSV、NPZ、JSON及正文核对表，供追溯。'],
    ['请先阅读.txt、数据清单.json','使用入口、表格行数、数据来源与转换记录。']],[6,11])
h('六张图索引')
table(['图号','绘图内容','本说明页码'],[
    ['图1','指定时刻的径向温度与含水率','2'],
    ['图2','环境观测与4小时后的平台延拓','3'],
    ['图3','中心与表面的热湿响应','4'],
    ['图4','全域达标过程与临界放大','5'],
    ['图5','收缩半径与六时刻含水率剖面','6'],
    ['图6','三组参数下的两类临界时长','7']],[1.5,13,2.5])
h('数据使用约定')
p('Excel第6行为字段名、第7行起为数值；CSV第1行为字段名、第2行起为数值。所有数据都已实际放入资料包，无需访问作者电脑或另外下载数据。')
p('列名中的h、s、cm和°C表示已经转换好的单位，不再重复换算。含水率C采用干基kg水/kg干物质；图2环境水分指标为kg/kg，不能标成相对湿度%。')
p('图1、图3、图4取N3200结果，图5取N6400结果；图6全部为N800情景。Excel用于直接绘图；需完整数值精度时取CSV或原始文件。')
p('数学变量使用Times New Roman斜体，单位及说明性下标用正体。颜色和线型同时区分曲线。以下论文页码对应当前稿，插图后按章节重新核对。',size=10)

card(1,'指定时刻的径向温度与含水率','5.1.4，当前第9页，表1与表2分析之后')
p('绘制左右两个面板：左图为温度，右图为干基含水率。两个面板均画100、300、600、900、1200、1500、1800 s七条曲线，使用相同的颜色与图例次序。')
files(1)
h('如何选取横轴与纵轴')
table(['数据列','取法'],[
    ['r_cm','两张表第一列均为横轴，范围0–2 cm。'],
    ['T_100s 至 T_1800s','温度表后七列分别为七条曲线，纵轴T / °C。'],
    ['C_100s 至 C_1800s','含水率表后七列分别为七条曲线，纵轴C / (kg/kg)。']],[5.5,11.5])
h('绘制要求')
p('每条曲线使用完整121个径向点，中心为r=0，表面为r=2 cm。绘制模型曲线；不要只拿正文五个半径的四位数值拼接作图。')
p('共享图例写“时间 / s”。两张表的同一列序对应同一时刻。温度列已经是摄氏度，不再减273.15；半径已经是厘米，不再乘100。')
h('建议图注')
p('问题一在固定半径2 cm、附录2物性下的径向温度与干基含水率分布，采用N3200生产结果。表面先响应，内部含水率变化存在滞后。')
source('原始数据/fig_q1_profiles_data.npz；正文核对表在 原始数据/正文核对表/q1_paper_temperature.csv 和 q1_paper_moisture.csv。')

card(2,'环境观测与4小时后的平台延拓','5.2.2，当前第11页，边界延拓公式之后')
p('上图绘制0–4 h环境温度观测点及线性插值，下图绘制环境水分指标；另用0–60 h的长时间轴或小插图展示4 h后的平台。t=4 h处画竖线，并注明两侧含义。')
files(2)
h('如何取数')
table(['数据列','取法'],[
    ['time_h','横轴时间 / h。观测表241行，0–4 h。'],
    ['temperature_C','温度纵轴 / °C。'],
    ['air_moisture_kg_per_kg','观测表的环境水分指标，纵轴 / (kg/kg)。'],
    ['boundary_moisture_kg_per_kg','平台表水分指标，作为独立假设系列。'],
    ['endpoint_included','0表示4 h平台左端开放；1表示60 h绘图端点包含。此列不作纵轴。']],[7.3,9.7])
h('平台与观测的边界')
table(['位置','温度 / °C','水分指标 / (kg/kg)'],[
    ['0 h观测','28','0.01963'],['4 h末次观测','50.165','0.04986'],['4 h右侧至60 h假设平台','50','0.05']],[7.3,4,5.7])
p('观测点相邻连直线即为分段线性插值。平台表仅两行，是水平线的端点；应画虚线，无实测标记。4 h平台端点用空心点或图注说明右侧极限，保留观测末点的小跳变。',size=10)
h('建议图注')
p('0–4 h为环境观测与分段线性插值；4 h后延拓为50°C和0.05 kg/kg的假设平台。空气指标映射为材料等效边界量属于模型假设。',size=10)
source('原始数据/A_environment_observed.csv；平台设置见 原始数据/Q23/summary.json 的settings。')

card(3,'中心与表面的热湿响应','5.2.4，当前第12页，表3与表4分析之后')
p('绘制温度与含水率两个时间面板，每图均含中心、表面两条曲线。温度图展示0–4 h；含水率图展示从0至临界达标附近，两图允许使用不同的横轴范围。')
files(3)
h('如何选列')
table(['工作表','横轴与纵轴'],[
    ['图3温度','time_h为X；centre_T_C与surface_T_C为两条Y，单位°C。'],
    ['图3含水率','time_h为X；centre_C_kg_per_kg与surface_C_kg_per_kg为两条Y，单位干基kg/kg。']],[4,13])
p('数据已按时间范围筛选：温度242个保存点，含水率3451个保存点。含水率表止于临界根57.47230195056044 h；若额外标严格报告点，使用“图4事件标记”，并单独标明其含义。')
h('线型与比较要求')
p('两个面板的中心曲线使用同一种颜色与线型，表面同理。中心为r=0，表面为r=2 cm。可以在温度图标出4 h观测结束时刻，在含水率图标出0.15阈值。')
p('采用问题二从均匀初始状态独立计算的完整轨迹。升温接近完成不代表内部水分已达到干燥阈值。')
h('建议图注')
p('问题二采用附录3整组变物性、固定半径和N3200结果；4 h后使用名义环境平台。中心与表面热湿响应表明热平衡与干燥达标具有不同时间尺度。')
source('原始数据/Q23/sampled_solution.npz；事件记录在 原始数据/Q23/summary.json；正文四位核对表在 原始数据/正文核对表/q2_paper_temperature.csv 和 q2_paper_moisture.csv。')

card(4,'全域达标过程与临界放大','5.3.4，当前第15页，达标过程说明处')
p('主图画中心、表面、全域最大含水率，并叠加0.15 kg/kg阈值线。放大图采用相对临界根的秒差为横轴，以浓度减0.15为纵轴，区分临界等号根和严格报告点。')
files(4)
h('主图与放大图各自取数')
p('主曲线：time_h作X，centre_C、surface_C、max_C作三条Y，threshold_C作水平线。max_C来自全体生产节点。中心与最大值重合时直接重合，不能人为错开。',size=10)
p('全域放大：delta_time_s作X，max_C_minus_threshold作Y，仅含两个已有保存点。中心放大：delta_time_s作X，centre_C_minus_threshold作Y，可展示根前120 s内的中心样点；标签必须写“中心样点”。',size=10)
p('事件标记：以delta_time_s和max_C_minus_threshold画两个独立点。kind=critical_equality为临界根；strict_report为严格报告点。负值表示低于阈值，勿先舍入为四位再画。',size=10)
table(['标记','时间 / h','最大干基含水率'],[
    ['临界根','57.47230195056044','0.15'],
    ['严格报告点','57.4724','0.14999989718225315']],[3.5,6.8,6.7])
h('图注与范围')
p('两标记相差约0.353 s。主图可设0–60 h，但数据止于206902 s，后续留空。临界根为等号条件；严格报告点经全域最大值核验低于0.15。放大图连线仅连接已有保存点，不新增预测点。',size=10)
source('原始数据/fig_q3_drying_data.npz；原始数据/Q23/sampled_solution.npz；原始数据/Q23/summary.json。')

card(5,'收缩半径与域内含水率剖面','5.4.4，当前第18页，表6与表7分析之后')
p('左图绘制半径随时间变化，并标出六个选定时刻；右图绘制0、6、12、24、48 h和临界根处的含水率剖面。每条剖面都在当时真实表面终止。')
files(5)
h('如何选取成对坐标')
p('左图：图5半径的time_h作X、radius_cm作Y。图5选定时刻表提供六个标记点及完整事件时刻。',size=10)
p('右图：图5含水率剖面每相邻两列是一组X/Y。依次选r_0h_cm与C_0h、r_6h_cm与C_6h、r_12h_cm与C_12h、r_24h_cm与C_24h、r_48h_cm与C_48h、r_event_cm与C_event。',bold=True,size=10)
p('每组21点。r已经按各自时刻计算为material_x×radius_m×100，单位cm；C为干基kg/kg。不要给六条曲线统一套第一组半径列。',size=10)
table(['时刻 / h','0','6','12','24','48','临界根'],[
    ['半径 / cm','2.0000','1.3740','1.2480','1.2040','1.2000','1.2000']],[3.2,2.2,2.2,2.2,2.2,2.2,2.8])
p('临界根为51.09057478683054 h。横轴可统一显示0–2 cm，但各曲线在自身半径处终止；域外没有药材，应留空，不填零、不补线。使用散点连线图以正确匹配每组X/Y。',size=10)
h('建议图注')
p('附录4物性下的径向同比收缩结果，固定轴向长度，N6400。与问题三的比较同时改变物性和几何，不能将全部干燥时差解释为纯收缩效应。',size=10)
source('原始数据/A_radius_observed.csv；原始数据/Q4/sampled_solution.npz；原始数据/Q4/summary.json；正文表在 原始数据/正文核对表/q4_paper_moisture.csv 和 q4_paper_radius.csv。')

card(6,'经验边界阻力情景对比','5.4.6，当前第19页，补充修正之后')
p('横轴为经验参数p=1、2、4，绘制两组带标记折线或分组柱图：问题二三定半径、附录3物性；问题四收缩、附录4物性。纵轴为临界时长 / h。')
files(6)
h('直接使用的六个数据点')
six=next(s for s in data['sheets'] if s['figure']==6)
table(['p','Q23_event_h','Q4_event_h'],six['rows'],[1.5,7.75,7.75])
p('p列作横轴；Q23_event_h和Q4_event_h为两组纵轴。绘图保留原始数值，文字标注可以显示四位小数。三行表格实际包含两组共六个点。')
h('比较范围与标注')
p('所有点统一采用N800网格、零潜热负荷，取顶层event_h的临界等号时长。p=1也使用同一N800情景族结果，不能用正文N3200或N6400生产时长替换。')
p('仅绘制已计算的p=1、2、4。折线用作阅读辅助，不补入其他p值，不添加未经计算的误差棒或置信带。')
h('建议图注')
p('在相同N800网格及零潜热负荷下比较p=1、2、4的经验边界阻力情景。该参数族尚未由真实药材等温线标定，结果说明所选边界解释对临界时长的影响。')
source('原始数据/isotherm_closure.json；scenarios中的六条iso_p1/2/4_Q23/Q4记录，均为computed_event_only。')

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页').font.size=Pt(9)
d.save(OUT/'六图绘图说明.docx')
print(json.dumps({'docx':(OUT/'六图绘图说明.docx').relative_to(ROOT).as_posix()},ensure_ascii=False))
