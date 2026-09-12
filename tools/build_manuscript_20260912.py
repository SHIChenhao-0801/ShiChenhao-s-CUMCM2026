"""Format the audited paper, preserving supplied front text and full code.

Uses the shared native-equation formatter with competition-specific typography,
requested textual drawing briefs, exact code paragraphs and review comments.
"""
from pathlib import Path
import copy
import hashlib
import json
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
if Path.cwd().resolve() != ROOT:
    raise RuntimeError('Run from the contest workspace.')
sys.path.insert(0, str(ROOT / '.agents/skills/paper-formal-writer/scripts'))
sys.path.insert(0, str(ROOT / 'tmp/cache/manuscript_20260912/python_deps'))
import format_formal_docx as fmt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from formula_omml import display_omml

QA = ROOT / 'paper_output/qa/manuscript_20260912'
SOURCE_DOCX = ROOT / 'reference_materials/peer-models/2026-09-12_front_matter/数学建模大赛a.docx'
code_records = json.loads((QA / 'code_appendix_manifest.json').read_text(encoding='utf-8'))['files']
extra_manifest = QA / 'verification_code_appendix_manifest.json'
if extra_manifest.exists():
    code_records += json.loads(extra_manifest.read_text(encoding='utf-8'))['files']
code_index = 0
equation_index = 0
original_font = fmt.apply_run_font


def run_font(run, font_name='宋体', size=None, bold=None):
    original_font(run,font_name,size,bold)
    if font_name=='宋体':
        run.font.name='Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    run.font.color.rgb=RGBColor(0,0,0)


def new_document():
    d = Document(SOURCE_DOCX)
    for node in list(d.element.body):
        if node.tag != qn('w:sectPr'):
            d.element.body.remove(node)
    return d


def configure(d):
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
    s.header_distance, s.footer_distance = Cm(1.0), Cm(1.25)
    for grid in list(s._sectPr.findall(qn('w:docGrid'))):
        s._sectPr.remove(grid)
    for section in d.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                p.clear()
    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Cm(.85)
    for name, size in [('Title',18),('Heading 1',15),('Heading 2',13),('Heading 3',12)]:
        if name not in d.styles:
            d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st=d.styles[name]
        st.font.name='黑体'; st.font.size=Pt(size); st.font.bold=True
        st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
        st.paragraph_format.space_before=Pt(10)
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.line_spacing=1.25
        st.paragraph_format.keep_with_next=True
    foot=s.footer.paragraphs[0]
    foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.paragraph_format.first_line_indent=Cm(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    foot._p.append(field)
    d.core_properties.author=''
    d.core_properties.last_modified_by=''
    d.core_properties.title='考虑物性变化与径向收缩的药材热湿耦合模型及数值验证'
    d.core_properties.subject=''
    d.core_properties.comments=''


def body(d,text,**kwargs):
    p=d.add_paragraph()
    p.paragraph_format.first_line_indent=Cm(.85)
    p.paragraph_format.line_spacing=1.25
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if text.startswith('图像主题'):
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.space_before=Pt(7)
        p.paragraph_format.space_after=Pt(7)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if re.match(r'^表(?:[A-Z])?\d+\s',text):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.keep_with_next=True
    return fmt.add_inline_content(p,text,font_size=12,**kwargs)


def heading(d,text,level):
    if not d.paragraphs:
        p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    else:
        p=d.add_paragraph(style=f'Heading {min(level,3)}')
    p.paragraph_format.first_line_indent=Cm(0)
    if text in ('摘要','1 问题重述','附录'):
        if text!='摘要':p.paragraph_format.page_break_before=True
        if text=='摘要':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    if re.match(r'^[DE]\.\d+\s',text):
        p.paragraph_format.page_break_before=True


def equation(d,latex,**kwargs):
    global equation_index
    equation_index+=1
    latex=re.sub(r'\\qquad\s*\\mathrm\{\(\d+\)\}', '', latex).strip()
    # Avoid LibreOffice's malformed limLow import while retaining a native
    # upright limit operator and its exact limit condition.
    latex=latex.replace(r'\lim_',r'\mathrm{lim}_')
    p=d.add_paragraph()
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(6)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if r'\begin{aligned}' in latex:
        rows=latex.replace(r'\begin{aligned}','').replace(r'\end{aligned}','').split(r'\\')
        elem=OxmlElement('m:oMathPara')
        math=OxmlElement('m:oMath');elem.append(math)
        array=OxmlElement('m:eqArr');math.append(array)
        for row in rows:
            row=row.replace('&','').strip()
            if not row:continue
            slot=OxmlElement('m:e');array.append(slot)
            converted=display_omml(row).find(qn('m:oMath'))
            for child in list(converted):slot.append(child)
    else:
        elem=display_omml(latex)
    # Equations retain their native editable structure; size follows the text.
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    p._p.append(elem.find(qn('m:oMath')))
    run_font(p.add_run(f'\t({equation_index})'),'宋体',12)
    return 1


def code(d,received):
    global code_index
    if code_index>=len(code_records):
        raise RuntimeError('Unexpected code block; code listings must be indexed.')
    rec=code_records[code_index];path=ROOT/rec['path']
    raw=path.read_bytes()
    assert hashlib.sha256(raw).hexdigest()==rec['sha256'],rec['path']
    original=raw.decode('utf-8-sig')
    normalized=lambda t:'\n'.join(x.rstrip() for x in t.strip('\n').splitlines())
    if normalized(received)!=normalized(original):
        import difflib
        details='\n'.join(difflib.unified_diff(original.splitlines(),received.splitlines()))
        (QA/'source_listing_mismatch.txt').write_text(details,encoding='utf-8')
        raise AssertionError(rec['path'])
    for line in original.splitlines():
        p=d.add_paragraph()
        p.style='Normal'
        pf=p.paragraph_format
        pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0)
        pf.line_spacing=Pt(10.5);pf.space_before=Pt(0);pf.space_after=Pt(0)
        pf.keep_together=False;pf.keep_with_next=False;pf.widow_control=False
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        run=p.add_run(line)
        fmt.apply_run_font(run,'Consolas',8.5)
        run._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
        # Stable listing membership permits exact code-text reconstruction.
        p._p.set(qn('w:rsidR'),'00000001')
    code_index+=1


def finish(path):
    d=Document(path)
    for table in d.tables:
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        for column in table.columns:column.width=Cm(16/len(table.columns))
        # Three-line table; no coloured cells or vertical rules.
        for i,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cell.width=Cm(16/len(table.columns))
                pr=cell._tc.get_or_add_tcPr()
                for old in list(pr):
                    if old.tag in (qn('w:shd'),qn('w:tcBorders')):pr.remove(old)
                borders=OxmlElement('w:tcBorders')
                for side in ('top','bottom','left','right'):
                    edge=OxmlElement('w:'+side)
                    on=(side=='top' and i==0) or (side=='bottom' and (i==0 or i==len(table.rows)-1))
                    edge.set(qn('w:val'),'single' if on else 'nil')
                    edge.set(qn('w:sz'),'8' if i in (0,len(table.rows)-1) else '4')
                    borders.append(edge)
                pr.append(borders)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.1
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.first_line_indent=Cm(0)
                    p.paragraph_format.left_indent=Cm(0)
                    p.paragraph_format.right_indent=Cm(0)
                    for r in p.runs:fmt.apply_run_font(r,'宋体',10.5,i==0)
    comments=[
        ('解析稀疏Jacobian配合BD处理刚性','原文按要求保留。“BD”应为“BDF”。Kirchhoff原函数改善面通量计算，仍使用温度面均值和有限精度求积，不能把“精确处理”理解为没有离散或积分误差。分钟量级是已记录计算环境下的运行表现，不是其他设备的时间保证。'),
        ('长2cm、半径2cm','原文按要求保留。题面及前文1.2给出的长度为25 cm，此处“长2cm”应改为“长25cm”；后续模型按0.25 m计算。'),
        ('Arhenius形式','原文按要求保留。Q1的D1(C)=7e−9 exp(−0.89/C)没有温度项；Arrhenius温度因子属于附录3/4。后文按各问实际公式区分。'),
        ('保留由坐标变换产生的网格速度相关项','原文按要求保留。采用同比材料收缩时，坐标变化项与物理平流相消，不另加网格平流。附录A.3给出完整推导。'),
        ('附录4扩散系数整体小于附录3','原文按要求保留。同T、同C下D4/D3=0.175exp(0.15/C)，只有C>0.0860600 kg/kg时小于1，低含水率表面可能反转。'),
        ('假设药材干物质骨架不发生迁移','原文按要求保留。Q4材料骨架随径向收缩运动，应限定为干物质没有生成、损失或相对骨架迁移，而不是空间静止。'),
        ('王乐毅,李长河,刘明政.中药材干燥技术','原条目按要求保留。刊方确认：王乐意,李长河,刘明政,等.中药材干燥技术与装备研究现状[J].农业工程学报,2024,40(2):1–28. DOI:10.11975/j.issn.1002-6819.202306104。'),
        ('王晓辉,王学成,唐培渝.数值模拟仿真','原条目按要求保留。PubMed PMID37474981确认2023,48(13):3440–3447，原页码正确；原文共有8名作者，前三人之后宜补“等”。')
    ]
    placed=[]
    for snippet,note in comments:
        matches=[p for p in d.paragraphs if snippet in p.text and p.runs]
        if len(matches)!=1:raise RuntimeError(f'Comment anchor ambiguous: {snippet}: {len(matches)}')
        d.add_comment(matches[0].runs,text=note,author='审阅',initials='')
        placed.append(snippet)
    # Assert every nonempty original body paragraph and all symbol cells survive.
    original=Document(SOURCE_DOCX)
    final_paras=[p.text for p in d.paragraphs]
    missing=[]
    skip={'问题重述','问题分析','模型假设','符号说明','模型的建立与求解','参考文献',
          '模型的评价','6.1模型的优点','6.2模型的不足','6.3模型的推广'}
    for p in original.paragraphs:
        if not p.text.strip() or p.text in skip:continue
        if not any(p.text in t for t in final_paras):missing.append(p.text[:80])
    final_cells=[c.text for tb in d.tables for row in tb.rows for c in row.cells]
    for tb in original.tables:
        for row in tb.rows:
            for c in row.cells:
                if c.text not in final_cells:missing.append('symbol cell:'+c.text)
    assert not missing,missing
    all_paragraphs=list(d.paragraphs)+[p for tb in d.tables for row in tb.rows for cell in row.cells for p in cell.paragraphs]
    for p in all_paragraphs:
        pr=p._p.get_or_add_pPr()
        snap=pr.find(qn('w:snapToGrid'))
        if snap is None:snap=OxmlElement('w:snapToGrid');pr.append(snap)
        snap.set(qn('w:val'),'0')
    d.save(path)
    (QA/'docx_structure_build.json').write_text(json.dumps({
        'sourceFrontTextPreserved':True,'commentsAdded':len(placed),
        'codeFiles':code_index,'codeLines':sum(x['lineCount'] for x in code_records),
        'displayEquationsNumbered':equation_index,'tables':len(d.tables),
        'docxSha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'userFigures':'Six textual briefs; no images generated.',
        'output':str(path.relative_to(ROOT))},ensure_ascii=False,indent=2),encoding='utf-8')


fmt.Document=new_document
fmt.configure_document=configure
fmt.add_body_paragraph=body
fmt.add_heading=heading
fmt.add_display_formula=equation
fmt.add_code_block=code
fmt.apply_run_font=run_font

if __name__=='__main__':
    result=fmt.main()
    if result:raise SystemExit(result)
    assert code_index==len(code_records)
    finish(fmt.DOCX_FILE)
