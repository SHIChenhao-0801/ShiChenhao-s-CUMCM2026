"""Create a practical Word handoff for the six agreed paper figures."""
from pathlib import Path
import json,hashlib,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path.cwd();assert ROOT.as_posix()=='D:/Document/数学建模/2026CUMCM'
QA=ROOT/'paper_output/qa/figure_guide_20260912';QA.mkdir(parents=True,exist_ok=True)
OUT=ROOT/'paper_output/paper/A题六张图绘制要求与数据说明.docx'
PROD='paper_output/results/production/final_v6a/'
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7);s.left_margin=Cm(2.0);s.right_margin=Cm(2.0)
s.footer_distance=Cm(.75)
for name,size,bold in [('Normal',11,False),('Title',21,True),('Heading 1',16,True),('Heading 2',11,True)]:
    st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if bold else '宋体')
    pf=st.paragraph_format;pf.line_spacing=1.15;pf.space_after=Pt(5)
    if name.startswith('Heading'):pf.space_before=Pt(9);pf.keep_with_next=True
    if name=='Title':pf.space_after=Pt(12)
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='A题六张图绘制要求与数据说明'
for border in list(d.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

def para(text,bold=False,size=None,style=None):
    p=d.add_paragraph(style=style);r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    p.paragraph_format.widow_control=True
    return p
def heading(text):return para(text,style='Heading 2')
def table(headers,rows,widths=None):
    tb=d.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=widths or [5.0,12.0]
    for c,w in zip(tb.columns,widths):c.width=Cm(w)
    for c,t in zip(tb.rows[0].cells,headers):c.text=t
    for row in rows:
        cells=tb.add_row().cells
        for c,t in zip(cells,row):c.text=str(t)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr();n=OxmlElement('w:cantSplit');trpr.append(n)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for key,val in [('top',70),('bottom',70),('left',90),('right',90)]:
                e=OxmlElement('w:'+key);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D5D9DE');borders.append(e)
            tcpr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF1F4');tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
    return tb
sources=[]
def source(label,path):
    p=ROOT/path;assert p.exists(),path;sources.append(path)
    para(label,bold=True,size=10)
    t=para(p.as_posix(),size=9.5);t.paragraph_format.space_after=Pt(5)
def card(n,title,loc,page):
    p=para(f'图{n} {title}',style='Heading 1');p.paragraph_format.page_break_before=True;p.paragraph_format.space_before=Pt(0)
    para(f'论文位置：{loc}　当前论文第{page}页',size=10)
def note(text):
    p=para(text,size=10);p.paragraph_format.space_before=Pt(6)
    return p

para('A题六张图绘制要求与数据说明',style='Title')
para('供论文制图与排版使用。每张图均列明绘制内容、坐标与单位、数据文件、字段选取和图注要求。按图号进入对应页，即可查到该图需要的数据。')
table(['图号','建议图名','正文位置'],[
 ['图1','问题一指定时刻的径向温湿分布','5.1.4　第9页'],
 ['图2','烘房边界观测及4小时后的平台延拓','5.2.2　第11页'],
 ['图3','问题二中心与表面的热湿响应','5.2.4　第12页'],
 ['图4','问题三全域阈值与临界时刻','5.3.4　第15页'],
 ['图5','问题四收缩半径与域内含水率剖面','5.4.4　第18页'],
 ['图6','经验边界阻力的情景比较','5.4.6　第19页']], [1.6,11.0,4.4])
heading('数据在哪里')
para('全部路径位于 D:/Document/数学建模/2026CUMCM。各图页列的是完整文件路径；可复制到文件资源管理器的地址栏打开所在文件夹。')
para('图2的观测数据可直接用 Excel/WPS 打开以下 CSV：',bold=True)
source('环境观测表','paper_output/data_cleaned/A_environment_observed.csv')
heading('怎样使用这些文件')
table(['文件类型','读取方式与用途'],[
 ['CSV','用 Excel/WPS 打开。第一行为字段名，后续为数据；使用对应的时间、半径及数值列。'],
 ['NPZ','NumPy数组文件，不能将其改扩展名当作Excel表。用支持NumPy的绘图程序按本手册字段读取，或先导出相应数组为CSV。'],
 ['JSON','结构化结果记录。按本手册给出的键名查找事件时间、报告点或六条情景值。']])
heading('六图共同约定')
para('时间由秒换小时：除以3600；半径由米换厘米：乘100；温度由K换°C：减273.15。已经是小时、厘米或摄氏度的列不再转换。')
para('图1、图3、图4采用N3200生产结果，图5采用N6400生产结果；图6全部采用独立N800情景。21个或121个绘图位置只是保存的绘图采样点数。')
para('曲线用未舍入数据，四位数值表供核对。数学变量采用Times New Roman斜体，单位和说明性下标正体；不同曲线同时用颜色与线型区分。')
note('页码对应当前31页论文版。图片插入或重新分页后，以章节及正文“【绘图位置】”定位，再核对正文20–30页要求。')

card(1,'问题一指定时刻的径向温湿分布','5.1.4 结果分析，表1和表2分析之后',9)
heading('需要画什么')
para('绘制左右双面板：左图为径向温度，右图为径向干基含水率。七条曲线分别对应100、300、600、900、1200、1500、1800 s；两个面板使用相同的时间颜色和图例次序。')
table(['坐标','要求'],[['两图横轴','距中心距离 r / cm，0–2 cm'],['左图纵轴','温度 T / °C'],['右图纵轴','干基含水率 C / (kg水/kg干物质)']])
heading('需要哪些数据')
source('主数据文件',PROD+'figures/fig_q1_profiles_data.npz')
table(['字段','形状与取法'],[['times_s','7个时刻；作为曲线标签，单位s。'],['radii_m','121个物理半径；乘100得到横轴cm。'],['temperature_C','7×121数组；第i行对应times_s[i]，无需再减273.15。'],['C','7×121数组；第i行与左图同一时刻，直接作为右图纵轴。']])
heading('作图和核对要求')
para('每个时刻取完整121点曲线，标明中心r=0和表面r=2 cm。用线图展示模型结果；正文表1、表2的五个指定半径只用于四位数值核对。')
source('核对温度表',PROD+'outputs/q1_paper_temperature.csv')
source('核对含水率表',PROD+'outputs/q1_paper_moisture.csv')
heading('图注应注明')
para('问题一、附录2物性、固定半径2 cm、N3200数值结果；表面先响应，中心含水率变化滞后。中心显示2.5500不表示其未舍入值始终不变。')

card(2,'烘房边界观测及长期平台延拓','5.2.2 变量定义与公式推导，边界延拓公式之后',11)
heading('需要画什么')
para('上图画0–4 h环境温度原始点与分段线性插值，下图画环境水分指标。再用较长时间轴或插图展示4–60 h平台；在t=4 h处画竖线，区分观测区间和假设延拓区间。')
table(['坐标','要求'],[['横轴','时间 t / h；观测细图0–4 h，长期图0–60 h。'],['上图纵轴','环境温度 / °C。'],['下图纵轴','环境水分指标 / (kg/kg)，不能标作相对湿度百分比。']])
heading('观测数据直接取这里')
source('原始观测数值的清洗表','paper_output/data_cleaned/A_environment_observed.csv')
table(['字段','具体用途'],[['time_h','横轴；241条记录覆盖0–4 h，另有1行表头。'],['temperature_C','上图观测点纵轴。'],['air_moisture_kg_per_kg','下图观测点纵轴。'],['time_s、temperature_K','辅助单位列；不与time_h、temperature_C重复转换。']])
heading('4小时后的数据怎样处理')
para('CSV只有0–4 h观测。4 h之后的平台由模型边界假设构造：对4<t≤60 h，温度恒取50°C，环境等效边界指标恒取0.05。画成两条水平虚线，不添加“实测点”标记。')
table(['位置','温度 / °C','水分指标 / (kg/kg)'],[['t=0 h观测','28','0.01963'],['t=4 h末次观测','50.165','0.04986'],['t=4 h右侧极限至60 h的假设平台','50','0.05']], [7.2,4.5,5.3])
para('相邻观测点直接连线即为线性插值。保留4 h处的实际末点；平台从4 h右侧开始，可用开端点或图注说明小跳变，不把末点改为50与0.05。',size=10)
source('平台设置记录',PROD+'Q23/summary.json')
note('读取settings中的boundary_extension=nominal、tail_temperature_C=50.0、tail_equilibrium=0.05。图注写明：0–4 h为观测与线性插值，4 h后为假设延拓；空气指标映射为材料等效边界量属于模型假设。')

card(3,'问题二中心与表面的热湿响应','5.2.4 结果分析，表3和表4分析之后',12)
heading('需要画什么')
para('分别绘制温度和含水率两个时间面板，每个面板均包含中心、表面两条曲线。温度图展示0–4 h；含水率图展示0至临界达标附近。两面板中心与表面使用一致的颜色或线型。')
table(['坐标','要求'],[['两图横轴','时间 t / h；允许采用不同横轴范围。'],['温度纵轴','T / °C；突出早期升温与中心滞后。'],['含水率纵轴','干基 C / (kg/kg)；中心r=0，表面r=2 cm。']])
heading('需要哪些数据')
source('主数据文件',PROD+'Q23/sampled_solution.npz')
table(['字段','形状与取法'],[['times_s','3452个保存时刻；除3600转换为小时。'],['material_x','21个材料位置，首项0、末项1；首列为中心，末列为表面。'],['T_K','3452×21；取首列和末列，再减273.15。'],['C','3452×21；取首列和末列，保持原精度。']])
heading('筛选时间与核对')
para('温度曲线保留times_s≤14400的样点。含水率曲线保留至临界根约57.47230195056044 h；若显示严格报告点，使用summary.json中的completion记录，不把临界等号根当作严格达标点。')
source('事件与报告时刻',PROD+'Q23/summary.json')
source('表3与表4的核对文件',PROD+'outputs/q2_paper_temperature.csv')
source('含水率核对文件',PROD+'outputs/q2_paper_moisture.csv')
heading('图注应注明')
para('问题二、附录3整组变物性、固定半径、N3200、4 h后名义平台。问题二从均匀初始状态单独起算；不要接到问题一1800 s的末态。升温接近完成不等于全域含水率已达标。')

card(4,'问题三全域阈值与临界时刻','5.3.4 结果分析，达标过程说明处',15)
heading('需要画什么')
para('主图画中心、表面及全域最大含水率，叠加C=0.15 kg/kg水平阈值线。增加临界时刻附近放大图，区分临界等号根与严格可行报告点。中心与最大值重合时直接重合，不人为错开曲线。')
table(['坐标','要求'],[['主图','横轴t / h，可设0–60 h；纵轴干基C / (kg/kg)。'],['放大图','横轴相对临界根的秒差；纵轴全域最大C减0.15。']])
heading('主图和放大图分别取数')
source('主图优先数据',PROD+'figures/fig_q3_drying_data.npz')
table(['字段','具体用途'],[['times_s','266个时刻；除3600。实际终点206902 s，约57.47278 h。'],['centre_C、surface_C','中心与真实表面曲线。'],['max_C','生产全体离散节点的最大值；不要改取21个稀疏半径的最大值。']])
source('放大图已有保存样点',PROD+'Q23/sampled_solution.npz')
para('用末尾times_s和C首列显示根附近的已有中心样点，并明确为样点连线。主图的全域最大曲线仍取max_C；严格报告点不在该NPZ中，另从下列JSON取数。',size=10)
source('临界根与严格报告点',PROD+'Q23/summary.json')
table(['标记','时间 / h','最大干基含水率'],[['临界根','57.47230195056044','0.15'],['严格报告点','57.4724','0.14999989718225315']], [3.5,6.8,6.7])
para('JSON读取completion对象：critical_event_h为临界小时，reported_drying_time_h为报告小时，reported_time_s为报告秒数，max_C_at_reported_time为报告点的全域最大含水率。',size=10)
para('两点相差约0.353 s；放大图需有足够分辨率。60 h只是可选轴上限，已计算终点之后留空，不继续外推曲线。',size=10)
heading('图注应注明')
para('N3200、4 h后长期边界假设；0.15等号对应临界根，严格报告点才满足全域最大值低于0.15。不要用四位显示浓度判断严格性，也不把该图解释为物理置信区间。')

card(5,'问题四收缩半径与域内含水率剖面','5.4.4 结果分析，表6和表7分析之后',18)
heading('需要画什么')
para('左图画半径随时间的变化，并标出右图选用时刻；右图在真实物理半径上画含水率剖面。建议选0、6、12、24、48 h及临界根51.09057478683054 h。每条曲线均在该时刻的真实表面终止。')
table(['坐标','要求'],[['左图','时间t / h，半径R / cm。'],['右图','物理半径r / cm，干基C / (kg/kg)。横轴可共用0–2 cm。']])
heading('需要哪些数据')
source('半径观测数据','paper_output/data_cleaned/A_radius_observed.csv')
para('145条数据记录，另有1行表头；time_h和radius_cm可直接作图。覆盖0–72 h、间隔0.5 h；模型按相邻点线性插值。',size=10)
source('含水率剖面与当前半径',PROD+'Q4/sampled_solution.npz')
table(['字段','形状与取法'],[['times_s','3069个时刻；选定六个时刻均已有保存样点。'],['material_x','21个0–1材料位置。'],['radius_m','与times_s逐项对应，乘100得到当前半径cm。'],['C','3069×21；第i行作为对应时刻的含水率剖面。']])
para('对第i个选定时刻：横轴 r_cm = material_x × radius_m[i] × 100；纵轴取 C[i,:]。最后一个材料位置等于1，所以自动包含真实表面端点。',bold=True,size=10)
table(['时刻 / h','0','6','12','24','48','临界根'],[['半径 / cm','2.0000','1.3740','1.2480','1.2040','1.2000','1.2000']], [3.2,2.2,2.2,2.2,2.2,2.2,2.8])
heading('域外与图注要求')
para('r>R(t)处无药材，留空或用灰区说明，不填零、不补线。叠加多时刻曲线时用各自颜色标注表面端点；需要分别标灰区时采用小面板。')
para('图注注明附录4物性、同比径向收缩、固定轴向长度、N6400。与问题三比较同时改变物性与几何，不把全部时差归因于纯收缩。')
source('正文含水率核对表',PROD+'outputs/q4_paper_moisture.csv')
source('正文半径核对表',PROD+'outputs/q4_paper_radius.csv')

card(6,'经验边界阻力的情景比较','5.4.6 边界阻力的补充修正之后',19)
heading('需要画什么')
para('横轴p=1、2、4，画两组带标记的折线或分组柱图：一组为问题二三定半径、附录3物性，另一组为问题四收缩、附录4物性。所有点统一使用N800和零潜热负荷设置。')
table(['坐标','要求'],[['横轴','经验形状参数p，无量纲，仅画1、2、4。'],['纵轴','临界时长 / h，采用等号事件时间，不采用严格上取报告时间。']])
heading('唯一取数文件与筛选条件')
source('六个情景数据源','paper_output/results/crossvalidation/isotherm_closure_v1/isotherm_closure.json')
para('读取scenarios数组，筛选intervals=800、latentFraction=0、isothermShapeP为1/2/4、status=computed_event_only。每条记录统一取顶层event_h字段，不混用嵌套输出中的显示值。')
table(['字段','作用'],[['scenario','用iso_p1_Q23等名称区分p值及问题。'],['isothermShapeP','横轴p。'],['event_h','纵轴临界时长，保留原始精度。']])
heading('可直接录入的六个数据点')
table(['p','问题二三临界时长 / h','问题四临界时长 / h'],[
 ['1','57.47232976256206','51.090599155463806'],
 ['2','59.623901353233975','52.22830295685303'],
 ['4','65.4447491379564','55.31621495305313']], [1.5,7.75,7.75])
para('情景名依次为iso_p1_Q23、iso_p2_Q23、iso_p4_Q23，以及对应iso_p1_Q4、iso_p2_Q4、iso_p4_Q4。绘图使用上表精度，标注可显示四位小数。')
heading('图注与比较范围')
para('p=1使用同一N800情景族的基线，不替换为正文N3200或N6400生产时长。不同p之间连线仅辅助阅读，不画未经计算的p值，不添加误差棒或置信带。')
para('建议图注：在相同N800网格及零潜热负荷下比较p=1、2、4的经验边界阻力情景。该族尚未由真实药材等温线标定，结果仅说明所选边界解释对临界时长的影响。')
note('本图不证明真实干燥时长的全局下界、任意p下的单调性或实测预测准确率。')

footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('第 ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
r=footer.add_run(' 页');r.font.size=Pt(9)
d.save(OUT)
manifest={'output':OUT.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),
 'sources':[{'path':p,'sha256':hashlib.sha256((ROOT/p).read_bytes()).hexdigest()} for p in dict.fromkeys(sources)],
 'figures':6,'newModelRuns':0,'includesActualImages':False,'authoring':'python-docx from bundled Python'}
(QA/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'artifact.md').write_text('用途：独立Word制图交接手册。六张图逐项说明绘制要求、坐标单位、真实数据完整路径、字段、数组形状、选点与图注。总览加六图分页，宋体正文与黑体标题，数学/西文延续Times New Roman参考方向。保留题设、观测、模拟和平台假设区别，不生成虚构数据或重算模型。不改原论文。真实渲染全部页面后交付。',encoding='utf-8')
print(json.dumps({'output':manifest['output'],'sha256':manifest['sha256'],'sources':len(manifest['sources'])},ensure_ascii=False))
