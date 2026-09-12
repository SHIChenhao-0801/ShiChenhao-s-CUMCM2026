"""Revise the user's latest DOCX with native equations and linked references.

This script edits a preserved input snapshot. It does not run or alter the model.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re
import sys

ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, str(ROOT / 'tools'))
import revise_docx_typography_20260912 as typography
from formula_omml import latex_to_omml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

QA = ROOT / 'paper_output/qa/revision_20260913'
SOURCE = QA / 'user_source.docx'
OUTPUT = ROOT / 'paper_output/paper/药材热湿耦合模型与干燥时间计算_文献公式修订版.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
W_TEXT = qn('w:t')

def visible(node):
    return ''.join(t.text or '' for t in node.iter() if t.tag in (W_TEXT, qn('m:t')))

def plain(node):
    return ''.join(t.text or '' for t in node.iter(W_TEXT))

def field(instruction, value):
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), instruction)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr'); typography.font(pr)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); pr.append(sz)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = str(value); r.append(t); f.append(r)
    return f

def bookmark(start_name, bookmark_id):
    a = OxmlElement('w:bookmarkStart'); a.set(qn('w:id'), str(bookmark_id)); a.set(qn('w:name'), start_name)
    b = OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'), str(bookmark_id))
    return a, b

def run(text, size=12):
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr'); typography.font(pr)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(round(size*2))); pr.append(sz)
    r.append(pr); t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
    return r

def inline_math_delimiters(text):
    # Balanced outer parentheses in the authored replacements denote inline math.
    # Numeric equation references remain ordinary text for Word REF conversion.
    output=[]; i=0
    while i<len(text):
        if text[i]!='(':
            output.append(text[i]); i+=1; continue
        j=i+1; depth=1
        while j<len(text) and depth:
            depth += (text[j]=='(')-(text[j]==')'); j+=1
        if depth:
            output.append(text[i:]); break
        value=text[i+1:j-1]
        if (re.search(r'[A-Za-z\\\\]|\[0,1\]',value) or re.fullmatch(r'\d+\.\d+',value)) and not any('\u4e00'<=c<='\u9fff' for c in value):
            unit = r'\ \mathrm{kg/kg}'
            if value.endswith(unit):
                output.extend(['$',value[:-len(unit)],'$ kg/kg'])
                i=j
                continue
            output.extend(['$',value,'$'])
        else:output.append(text[i:j])
        i=j
    return ''.join(output)

def replace_span(p, start, end, elements):
    """Replace ordinary text interval without flattening intervening OMML."""
    nodes = list(p.iter(W_TEXT)); cursor = 0; selected = []
    for t in nodes:
        a, b = cursor, cursor + len(t.text or ''); cursor = b
        if a < end and b > start: selected.append((t, a, b))
    assert selected, (start, end, plain(p))
    first, a, b = selected[0]; last, la, lb = selected[-1]
    first_run = first.getparent(); assert first_run.tag == qn('w:r')
    parent = first_run.getparent(); assert parent.tag == qn('w:p'), parent.tag
    prefix = (first.text or '')[:start-a]
    suffix = (last.text or '')[end-la:]
    first.text = prefix
    for t, _, _ in selected[1:]: t.text = ''
    pos = parent.index(first_run) + 1
    for e in elements: parent.insert(pos, e); pos += 1
    if first is last:
        if suffix: parent.insert(pos, run(suffix))
    else:
        last.text = suffix

def replace_exact(p, old, new):
    text = plain(p); assert text.count(old) == 1, (old, text)
    at = text.index(old)
    replace_span(p, at, at+len(old), [run(new)])

def main():
    doc = Document(SOURCE)
    original_paragraphs = list(doc.paragraphs)
    original_text = [visible(p._p) for p in original_paragraphs]
    changed = set(); changes = []
    def write(index, text):
        p = original_paragraphs[index]
        p.clear()
        text=inline_math_delimiters(text)
        typography.layout.fmt.add_inline_content(p, text, font_size=12)
        for r in p.runs:
            r.font.name = 'Times New Roman'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
            r.font.color.rgb = RGBColor(0,0,0)
        changed.add(index)
        changes.append({'paragraph': index, 'before': original_text[index], 'after': text})
        return p

    # Preserve the user's changes while completing the agreed scientific edits.
    write(10, '中药材干燥是产地初加工的重要环节，干燥方式及工艺参数会影响产品品质、能耗和生产效率。热风干燥是常用方法之一，相关技术与装备的发展现状已有系统综述[1]。数值模拟可用于分析药材干燥过程中的温度、水分分布及工艺条件的影响，为实验设计和参数选择提供参考[2]。本题给定的烘干过程分为预热平衡与恒温干燥两个阶段，本文据此研究圆柱形药材的径向热湿传递及烘干时间。')
    write(15, '问题二：将研究范围扩展至整个烘干过程，统一采用题给附录3的物性关系，建立药材温度与干基含水率的时空变化模型，计算每隔1 s、径向每隔0.1 cm的温度和含水率，并按题定时刻与半径给出结果表。')
    write(20, r'已知圆柱形药材的几何尺寸、均匀初始状态、烘房温湿度序列及附录2参数，问题一要求得到1800 s内的径向温度和含水率分布。由傅里叶导热定律与Fick扩散定律建立轴对称控制方程，并给定中心零通量和表面对流边界条件。本问热物性为常数，温度方程可独立求解，并可构造圆柱Bessel级数基准进行检验；水分扩散系数 (D_1(C)) 仍随含水率变化，采用守恒型有限体积离散和隐式时间积分求解。问题一的扩散系数不含温度项；问题二至四含Arrhenius温度因子的经验式统一以开尔文温度代入。')
    replace_exact(original_paragraphs[22]._p, '烘焙', '烘干')
    changes.append({'paragraph':22,'before':'烘焙','after':'烘干','type':'terminology'})
    write(24, r'问题三沿用问题二的变物性耦合模型及同一条数值轨迹，确定药材各处干基含水率均低于 (0.15\ \mathrm{kg/kg}) 所需的时间。以全域最大含水率构造阈值函数，通过连续数值解定位其下降穿越阈值的临界时刻，再在指定时间精度下选取并核验满足严格不等式的烘干时长。该判据同时检查全部径向位置，不预先指定中心或表面作为唯一控制点。')
    write(26, r'问题四需要处理半径随时间减小引起的移动边界。半径由附件2线性插值得到，热容量、导热系数和扩散系数采用附录4的经验关系，并随当前含水率与温度更新。')
    write(27, r'在干物质守恒及径向同比收缩的假设下，引入材料坐标 (x=r/R(t))，将随时间变化的药材区域映射到固定区间 ([0,1])。从组分守恒方程出发进行坐标变换，骨架运动产生的平流项与坐标变化项相消，收缩影响通过时变几何系数及边界条件体现。在此基础上采用有限体积法离散求解。为考察收缩本身的影响，保持附录4物性、初始条件、环境条件及数值设置一致，分别计算固定半径与给定半径变化两种情形，比较其含水率分布和烘干时长。问题三与问题四的结果之差则反映物性变化和径向收缩的综合影响。')
    write(32, r'假设药材干物质初始分布均匀，烘干过程中不存在干物质的生成、损失及相对于骨架的迁移。问题一至三忽略药材尺寸变化；问题四假设轴向长度保持不变，干物质骨架仅发生径向同比收缩，即各材料点到轴线的距离均按 (R(t)/R_0) 的比例变化。')
    write(41, r'干基含水率表示水质量与干物质质量的比值。为保持水量计算的量纲一致，引入单位当前体积内的干物质质量 (\rho_{\mathrm d})，则水质量密度为 (\rho_{\mathrm d}C)。问题一至三采用固定且空间均匀的干物质骨架；问题四允许骨架随材料收缩运动，并假设干物质相对于骨架的通量为零。下文据此建立移动区域上的组分守恒方程。')
    write(43, r'平衡含水率与材料性质、温度及环境湿度有关，通常需由吸附或解吸等温线描述[3]。附件1的烘房水分指标仅标注kg/kg，未给出其与药材干基含水率之间的转换关系。记该观测数值为 (Y_\infty)，基线取 (C_{\mathrm{eq}}=Y_\infty)，作为表面传质的等效平衡值。这一数值映射是本文的边界假设；文献中的一般平衡关系不构成该映射的标定依据。若进一步将 (Y_\infty) 解释为空气湿度比，其参考质量也需另作说明。')
    write(49, r'式(2)表示有效显热的局部变化由净导热供给；式(3)由水质量守恒结合Fick扩散定律，在干骨架密度均匀且固定时得到。导热系数和扩散系数保留在散度内，以反映状态变化引起的通量梯度。相应常系数圆柱问题的Bessel函数级数解可参考Crank的推导[4]，用于检验简化条件下的数值实现。')
    write(58, r'设归一化坐标 (x=r/R_0)，在 (0\le x\le1) 上划分 (N) 个等长区间，节点 (x_i=i/N) 包含中心和表面。每个节点对应相邻半节点之间的环形控制体，中心和表面取截断控制体，记其无量纲体积权重为 (w_i)。有限体积法先计算相邻控制体共用的面通量，再由净通量建立收支；内部面在全域求和时成对相消[5]。该处理可直接落实轴心零通量条件。')
    write(59, r'对于随含水率显著变化的扩散系数，参照浓度相关扩散方程的积分变换方法[4]，构造Kirchhoff浓度势。将题给系数写为 (D=D_0A(T)\exp(-a/C))，取')
    write(62, r'细网格缩短了局部扩散时间尺度，显式推进受到较严格的步长限制。本文将温度与含水率交错排列，利用相邻节点耦合构造解析稀疏雅可比矩阵，并采用适于刚性方程组的BDF方法进行时间积分[6]。实际使用SciPy的BDF实现及其稀疏雅可比接口[7]。状态末项累计失水，用于核对全域水量收支。题定输出时刻由同一连续数值解查询；积分器的内部步长与1 s或60 s的输出间隔分别设置。')
    write(103, '空间加密、收紧时间设置以及不同积分器的对照表明，本组数值解在所检验的采样范围内稳定。表3中1.5 h表面温度接近四位小数的舍入边界，基线与更严格时间设置可能给出不同末位；表中统一采用基线计算值47.1401°C。该差异反映计算末位对时间设置的敏感性，详细比较见附录C。')
    write(106, '问题三将给定时刻的温湿状态计算转化为烘干时间求解。沿用问题二的完整数值轨迹，以全域最大含水率构造阈值事件，定位临界时刻，并据此确定满足严格含水率约束的烘干时长。')
    write(108, r'问题三要求确定药材各处干基含水率均低于 (0.15\ \mathrm{kg/kg}) 所需的烘干时间。沿用问题二得到的含水率场 (C(r,t))，定义时刻 (t) 的全域最大含水率：')
    write(110, r'其中，(M(t)) 的单位仍为kg/kg。药材各处均达标，等价于全域最大含水率满足')
    write(112, '为定位阈值，定义辅助函数')
    write(115, r'当 (g(t)>0) 时尚有位置未达标；(g(t)=0) 对应临界边界；(g(t)<0) 表示全域达标。数值计算对全部径向节点取最大值。')
    write(116, r'记 (g(t)) 首次下降穿越零点的临界时刻为 (t_{\mathrm c})。在临界附近连续下降的条件下，根时刻对应等号。取时间间隔 (\Delta t=0.36\ \mathrm s=10^{-4}\ \mathrm h)，在该离散时间网格上选择不早于临界时刻且满足严格阈值的首个时刻：')
    write(118, r'式中，(\mathbb{N}_0) 为非负整数集合；(t_{\mathrm c})、(t_{\mathrm r}) 与 (\Delta t) 统一以秒计算。达标判定采用未舍入的数值解，最终将 (t_{\mathrm r}) 换算为小时表示。候选点的生成和逐格核验见附录C.4。')
    write(120, r'第一步：在问题二的积分过程中监测 (g(t)) 下降穿零。第二步：利用连续数值解在包围区间内定位临界根，并核对最大含水率与阈值的一致性。第三步：按式(14)选取并核验烘干时长。第四步：从同一轨迹查询每隔60 s、径向每隔0.1 cm的含水率。另在 (N=800) 的验证轨迹上采用独立二分法复核事件定位；正式结果由 (N=3200) 的计算轨迹给出。')
    write(121, r'求解阈值方程得到临界时刻 (t_{\mathrm c}\approx57.472302\ \mathrm h)。按式(14)选取并核验未舍入含水率，得到烘干时长为 (57.4724\ \mathrm h)。')
    write(123, '表5列出每隔6 h及最终烘干时刻的径向含水率，用于展示干燥过程中的空间差异；题定60 s间隔的完整结果另行导出。')
    write(127, r'注：含水率保留四位小数。(57.4724\ \mathrm h) 时的全域最大含水率约为 (0.149999897\ \mathrm{kg/kg})，表中显示为 (0.1500)；达标判定采用未舍入的计算值。')
    # Place the short explanation next to its table rather than in the analysis.
    table5 = doc.tables[5]  # symbols, table1, table2, table3, table4, table5
    assert visible(table5._tbl).startswith('时间/h'), visible(table5._tbl)[:80]
    table5._tbl.addnext(original_paragraphs[127]._p)
    original_paragraphs[127].paragraph_format.first_line_indent = Cm(0)
    original_paragraphs[127].paragraph_format.line_spacing = 1.15
    for r in original_paragraphs[127].runs:r.font.size=Pt(10.5)
    write(138, r'将水质量方程减去干质量方程的 (C) 倍，得到 (\rho_{\mathrm d}D_tC=-\nabla\cdot\mathbf j_{\mathrm w})。取 (\mathbf j_{\mathrm w}=-\rho_{\mathrm d}D\nabla C)，并利用干骨架密度的空间均匀性消去 (\rho_{\mathrm d})。令 (x=r/R(t))，定义 (\widehat C(x,t)=C(R(t)x,t))、(\widehat T(x,t)=T(R(t)x,t))。链式法则中的坐标变化项与骨架平流项相消，得到')
    write(141, r'式(17)沿用有效显热闭合，未显式包含潜热、湿分携焓及收缩功。材料坐标方程中不再出现额外网格平流项；干基比值也不另加体积浓缩项。收缩引起的干骨架密度变化满足 (\rho_{\mathrm d}(t)=\rho_{\mathrm d0}[R_0/R(t)]^2)。坐标变换和干质量守恒推导见附录A。')
    write(147, r'在相同温度与含水率下，两组扩散系数满足 (D_4/D_3=0.175\exp(0.15/C))。当 (C\approx0.08606\ \mathrm{kg/kg}) 时两者相等；高于该值时 (D_4<D_3)，低于该值时大小关系反转。因此，物性变化的综合作用由全过程数值对照评估，不能仅由扩散系数前因子判断。')
    write(149, r'在材料坐标中采用 (N=6400) 的固定网格。每次右端函数求值时，根据当前 (t) 更新 (R(t))，再更新式(16)—(20)中的几何因子、物性和边界通量。网格点随材料运动，输出时转换到题目要求的物理半径 (r)。')
    write(153, r'问题四的连续临界时刻约为 (51.090575\ \mathrm h)，按与问题三相同的时间网格和严格阈值判据，得到烘干时长为 (51.0906\ \mathrm h)，对应半径约为 (1.2000\ \mathrm{cm})。该时刻位于附件2的观测范围内。表6与表7分别给出含水率和表面半径。')
    write(173, r'达标时间由全域最大含水率事件确定，并在 (N=800) 的验证轨迹上用独立二分法复核事件定位。最终烘干时刻按式(14)确定，其可行性由未舍入的最大含水率核验；两种定位方法的比较见附录C.4。')
    write(179, r'针对扩散系数随含水率显著变化的特点，采用Kirchhoff变换构造水分扩散的界面通量；针对径向收缩引起的区域变化，采用材料坐标将移动区域映射到固定区间。空间离散后，使用隐式BDF进行时间积分，并提供解析构造的稀疏雅可比矩阵，以处理细网格离散和热湿耦合形成的刚性方程组。')
    write(182, '圆柱控制体统一处理轴心、内部界面和表面收支；磁盘连续输出降低高网格轨迹的内存占用。解析基准、网格加密、时间设置和独立实现分别检查不同求解环节，为结果一致性提供相互补充的证据。')
    write(184, '模型的主要限制来自气固平衡映射、有效显热容量和内部同比收缩等假设；这些关系仍需相应实验进一步验证。')
    write(194, 'AI参与内容包括文献检索与真实性核查、模型推导辅助、程序生成与调试、计算结果核对、公式与引用整理，以及基于已有结果数据的图表制作。本稿中的数值对照用于说明所列计算条件下的一致性。参赛队对采用的模型、源码、引用和结论承担核实责任；核心代码及最终稿的团队人工审查另行记录。')
    write(199, '本附录补充正文方程的守恒推导、数值求解步骤和核心算法。温度、时间、长度在计算中分别采用K、s、m；含水率为kg水/kg干物质。式号续正文式(1)—(21)。完整输入、源程序和结果记录作为支撑材料保留，文末仅列实际生产程序的必要片段。')
    write(261, r'离散全域最大值 (M_N(t)=\max_i C_i(t)) 定义事件 (g(t)=M_N(t)-0.15)，监测其下降穿零。节点间采用分段线性重构时，最大值在节点处取得。连续事件根 (t_{\mathrm c}) 以秒计，程序继续真实积分至 (t_{\mathrm{end}}=\delta t_{\mathrm s}[\lceil t_{\mathrm c}/\delta t_{\mathrm s}\rceil+1])，其中 (\delta t_{\mathrm s}=1\ \mathrm s)。在小时四位小数对应的时间网格上，初始候选点为')
    write(263, r'取 (\Delta t=0.36\ \mathrm s)，并依次检查 (t_0,t_0+\Delta t,\ldots) 处未舍入的 (M_N)。第一个满足 (M_N<0.15) 的候选点即正文式(14)的 (t_{\mathrm r})；若候选点超出已实际积分的区间，则报错。问题三的连续根为57.47230195056044 h，选取时长为57.4724 h；问题四分别为51.09057478683054 h和51.0906 h。表中显示精度与阈值判定精度分别处理。')

    # Identify all numbered equations before changing their displayed labels.
    equations=[]
    for i,p in enumerate(original_paragraphs):
        wt=plain(p._p).strip()
        if p._p.find(qn('m:oMath')) is not None and re.fullmatch(r'\(\d+\)',wt):
            equations.append((i,p,int(wt[1:-1])))
    assert len(equations)==37, [(i,n) for i,p,n in equations]
    eq_replacements={
      14:r't_{\mathrm r}=\min\{k\Delta t:\ k\in\mathbb{N}_0,\ k\Delta t\ge t_{\mathrm c},\ M(k\Delta t)<0.15\}',
      18:r'-\frac{k}{R(t)}\widehat T_x(1,t)=h(T_{\mathrm s}-T_\infty),\quad-\frac{D}{R(t)}\widehat C_x(1,t)=\beta(C_{\mathrm s}-C_{\mathrm{eq}})',
      37:r't_0=\Delta t\left\lceil\frac{t_{\mathrm c}}{\Delta t}\right\rceil,\qquad \Delta t=0.36\ \mathrm s.'
    }
    eq_map=[]
    for new,(i,p,old) in enumerate(equations,1):
        math=latex_to_omml(eq_replacements[new]) if new in eq_replacements else deepcopy(p._p.find(qn('m:oMath')))
        p.clear();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t');p._p.append(math);p.add_run('\t(')
        a,b=bookmark(f'eq_{new:03d}',10000+new)
        p._p.append(a);p._p.append(field(' SEQ Equation \\* ARABIC ',new));p._p.append(b);p.add_run(')')
        eq_map.append({'sourceParagraph':i,'sourceDisplayedNumber':old,'finalNumber':new,'bookmark':f'eq_{new:03d}','formulaChanged':new in eq_replacements})

    # Existing prose references still refer to the original pre-insertion numbers.
    eq_paragraph_ids={id(p._p) for _,p,_ in equations}
    ref_pattern=re.compile(r'(?:正文)?式\s*\(\d+\)(?:(?:[—–、,，]|及|和|至)\s*(?:边界式|式)?\s*\(\d+\))*')
    old_to_new={n:n if n<=11 else n+2 for n in range(1,36)}
    for i,p in enumerate(original_paragraphs):
        if id(p._p) in eq_paragraph_ids:continue
        text=plain(p._p)
        refs=[]
        for m in ref_pattern.finditer(text):
            refs.extend((m.start()+s.start()+1,m.start()+s.end()-1,int(s.group()[1:-1])) for s in re.finditer(r'\(\d+\)',m.group()))
        for a,b,n in reversed(refs):
            target=n if i in changed else old_to_new[n]
            replace_span(p._p,a,b,[field(f' REF eq_{target:03d} \\h ',target)])

    # Verified sources are inserted in first-citation order, not guessed from old labels.
    sources=json.loads((QA/'references/references_verified.json').read_text(encoding='utf-8'))
    sources=sources['references'] if isinstance(sources,dict) else sources
    assert len(sources)==7,len(sources)
    ref_anchor=original_paragraphs[195]._p
    original_paragraphs[195].paragraph_format.page_break_before=True
    for i in (196,197):original_paragraphs[i]._p.getparent().remove(original_paragraphs[i]._p)
    for n,item in enumerate(sources,1):
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.15
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
        p.add_run('[');a,b=bookmark(f'bib_{n:03d}',11000+n)
        p._p.append(a);p._p.append(field(' SEQ Bibliography \\* ARABIC ',n));p._p.append(b)
        p.add_run('] '+re.sub(r'^\[\d+\]\s*','',item['bibliography']))
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10.5);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        ref_anchor.addnext(p._p);ref_anchor=p._p

    # All citations become clickable Word REF fields with cached values.
    citation_count=0
    for p in doc.paragraphs:
        if any(x.get(qn('w:name'),'').startswith('bib_') for x in p._p.iter(qn('w:bookmarkStart'))):continue
        text=plain(p._p)
        for m in reversed(list(re.finditer(r'\[([1-7])\]',text))):
            n=int(m.group(1));replace_span(p._p,m.start(1),m.end(1),[field(f' REF bib_{n:03d} \\h ',n)]);citation_count+=1

    # All eight source comments concern items resolved by this revision.
    resolved_comments=[{'id':c._element.get(qn('w:id')),'text':c.text} for c in doc.comments]
    for c in list(doc.comments):c._element.getparent().remove(c._element)
    for e in list(doc.element.body.iter()):
        if e.tag in {qn('w:commentRangeStart'),qn('w:commentRangeEnd'),qn('w:commentReference')}:
            e.getparent().remove(e)
    # Clarify the pre-existing event symbol in the notation table.
    for row in doc.tables[0].rows:
        if '全域达标时刻' in row.cells[1].text:
            row.cells[1].text=row.cells[1].text.replace('全域达标时刻','全域含水率阈值的临界时刻')
            cell=row.cells[0]; cell.text=''; cell.paragraphs[0]._p.append(latex_to_omml(r't_{\mathrm c}'))
    # Preserve all code fragment texts and all result table cells.
    code_start=next(i for i,t in enumerate(original_text) if t=='附录 D 必要算法片段')
    assert [visible(p._p) for p in original_paragraphs[code_start:]]==original_text[code_start:]
    original_doc=Document(SOURCE)
    for before,after in zip(original_doc.tables[1:],doc.tables[1:]):
        assert [[c.text for c in row.cells] for row in before.rows]==[[c.text for c in row.cells] for row in after.rows]
    # Fix only new prose/math typography; preserve user paragraph styles.
    for i in changed:typography.ordinary_style(original_paragraphs[i]._p)
    typography.mathematical_style(doc)
    for p in doc.paragraphs:
        typography.child(p._p.get_or_add_pPr(),'w:snapToGrid').set(qn('w:val'),'0')
    for e in list(doc.settings.element.findall(qn('w:updateFields'))):doc.settings.element.remove(e)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.title='药材热湿耦合模型与干燥时间计算'
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    report={'source':str(SOURCE.relative_to(ROOT)),'sourceSha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            'output':str(OUTPUT.relative_to(ROOT)),'outputSha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
            'changes':changes,'equations':eq_map,'referenceCount':len(sources),'citationCount':citation_count,
            'resolvedSourceComments':resolved_comments,'resultTablesUnchanged':True,'codeFragmentsUnchanged':True,
            'figuresInserted':False,'productionRecomputed':False}
    (QA/'revision_build.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in {'changes','equations','resolvedSourceComments'}},ensure_ascii=False))

if __name__=='__main__':main()
