"""Check the typography/appendix revision without rerunning the frozen model."""
from pathlib import Path
from zipfile import ZipFile
import sys,json,re,hashlib
from collections import Counter
import pymupdf
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path.cwd();assert ROOT.as_posix()=='D:/Document/数学建模/2026CUMCM'
QA=ROOT/'paper_output/qa/manuscript_revision_20260912'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
version=sys.argv[1]
manifest=json.loads((QA/f'{version}_manifest.json').read_text(encoding='utf-8'))
path=ROOT/manifest['docx'];assert sha(path)==manifest['docxSha256']
pdfpath=ROOT/manifest['pdf'];assert sha(pdfpath)==manifest['pdfSha256']
doc=Document(path);pdf=pymupdf.open(pdfpath)
texts=[p.get_text() for p in pdf]
for p in manifest['pages']:assert sha(ROOT/p['image'])==p['sha256']
ap=next(i for i,t in enumerate(texts) if re.search(r'(?m)^附录\s*$',t))
bp=next(i for i,t in enumerate(texts) if re.search(r'(?m)^1\s+问题重述\s*$',t))
assert bp==1 and 20<=ap-bp<=30,(bp,ap)
figs=[]
for n in range(1,7):
    pages=[i+1 for i,t in enumerate(texts) if f'【绘图位置{n}】' in re.sub(r'\s+','',t)]
    assert len(pages)==1,pages
    figs.append({'number':n,'page':pages[0]})
ns={'w':qn('w:p').split('}')[0][1:],'m':qn('m:oMath').split('}')[0][1:]}
with ZipFile(path) as z:
    assert z.testzip() is None
    xml=etree.fromstring(z.read('word/document.xml'))
    comments=etree.fromstring(z.read('word/comments.xml'))
    assert not [n for n in z.namelist() if n.startswith('word/media/')]
comment_texts=[''.join(c.xpath('.//w:t/text()',namespaces=ns)) for c in comments]
for n in (3,4,5,6):assert sum(f'正文引用[{n}]' in t and '没有对应条目' in t for t in comment_texts)==1
body=xml.find(qn('w:body'))
visible=lambda e:''.join(e.xpath('.//w:t/text()|.//m:t/text()',namespaces=ns))
boundary=next(i for i,e in enumerate(body) if visible(e).strip()=='附录')
bodytext='\n'.join(visible(e) for e in list(body)[:boundary])
refs=bodytext.split('参考文献')[-1]
assert re.findall(r'(?m)^\[([0-9]+)\]',refs)==['1','2']
for n in (3,4,5,6):assert f'[{n}]' in bodytext
code=[p.text for p in doc.paragraphs if p._p.get(qn('w:rsidR'))=='00000002']
assert len(code)==117 and all(t.strip() and not t.lstrip().startswith(('#','import ','from ')) for t in code)
maths=xml.findall('.//'+qn('m:oMath'))
math_fonts=Counter();bad=[];ceil=[]
for m in maths:
    for r in m.findall('.//'+qn('m:r')):
        t=''.join(r.xpath('./m:t/text()',namespaces=ns))
        f=r.find(qn('w:rPr')+'/'+qn('w:rFonts'))
        name=f.get(qn('w:ascii')) if f is not None else None
        math_fonts[name]+=1
        if name!='Times New Roman':bad.append(t)
        if t in {'⌈','⌉'}:
            nor=r.find(qn('m:rPr')+'/'+qn('m:nor'))
            ceil.append({'glyph':t,'normal':nor is not None and nor.get(qn('m:val'))=='1'})
assert not bad and ceil and all(x['normal'] for x in ceil)
hour_headers=0;kp_bases=0
for p in xml.findall('.//'+qn('w:p')):
    if visible(p)=='时间/h':
        for r in p.findall(qn('w:r')):
            if ''.join(r.xpath('./w:t/text()',namespaces=ns))=='h':
                assert r.find(qn('w:rPr')+'/'+qn('w:i')).get(qn('w:val'))=='0'
                hour_headers+=1
for s in xml.findall('.//'+qn('m:sSub')):
    if visible(s.find(qn('m:e')))=='K' and visible(s.find(qn('m:sub')))=='p':
        r=s.find(qn('m:e')+'/'+qn('m:r'))
        assert r.find(qn('m:rPr')+'/'+qn('m:sty')).get(qn('m:val'))=='i'
        assert r.find(qn('m:rPr')+'/'+qn('m:nor')) is None
        kp_bases+=1
assert hour_headers==5 and kp_bases>=2,(hour_headers,kp_bases)
outside=[]
for pi,p in enumerate(pdf):
    for block in p.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            for span in line['spans']:
                x0,y0,x1,y1=span['bbox']
                if x0<0 or y0<0 or x1>p.rect.width+0.5 or y1>p.rect.height+0.5:
                    outside.append({'page':pi+1,'text':span['text'],'bbox':span['bbox']})
assert not outside,outside
numbered=[]
for p in xml.findall('.//'+qn('w:p')):
    if p.find('.//'+qn('m:oMath')) is not None:
        s=''.join(p.xpath('./w:r/w:t/text()',namespaces=ns))
        m=re.search(r'\((\d+)\)\s*$',s)
        if m:numbered.append(int(m.group(1)))
assert numbered==list(range(1,36)),numbered
fonts=sorted({f[3] for p in pdf for f in p.get_fonts()})
result={'status':'PASS','scope':'DOCX structure, source identity, typography properties, exact rendered page scope; whole-page visual review is recorded separately.',
 'docx':manifest['docx'],'docxSha256':manifest['docxSha256'],'renderManifest':str((QA/f'{version}_manifest.json').relative_to(ROOT)),
 'pageCount':len(pdf),'abstractPages':bp,'bodyPagesIncludingAIReferences':ap-bp,'appendixPages':len(pdf)-ap,'appendixStartPage':ap+1,
 'pageRequirement':'User confirmed body 20–30 pages; abstract and appendix excluded.',
 'figurePositions':figs,'numberedFormulas':len(numbered),'editableNativeFormulaObjects':len(maths),'tables':len(doc.tables),
 'codeFragments':6,'codeLines':len(code),'comments':len(comment_texts),'unmatchedCitationComments':[3,4,5,6],
 'bibliographyEntriesRemaining':[1,2],'mathRunFonts':dict(math_fonts),'ceilSymbolProperties':ceil,'actualPdfFonts':fonts,
 'outOfPageSpans':outside,'pdeRerun':False,'actualGuiOrHumanCodeReviewThisRevision':False,
 'uprightHourUnitHeaders':hour_headers,'italicKpVariableBases':kp_bases,
 'sourceSnapshotUnchanged':sha(QA/'user_edited_source.docx')=='0c8bb744b568fd751dc0ad0d883f69b5f895a74b9f8fcb3666b99975db7b05ee'}
assert result['sourceSnapshotUnchanged']
(QA/'final_structural_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
