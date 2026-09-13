from __future__ import annotations

import argparse
import csv
from dataclasses import asdict
from datetime import datetime, timezone
import gzip
import hashlib
import inspect
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import time

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
SOURCE = Path(__file__).resolve()
LOADED_EXPORTER_SHA256 = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
TEMPLATE_DIR = ROOT / 'problem_files/CUMCM2026Problems/A题/附件/附件3'
RADII_M = np.arange(21, dtype=float)*0.001
RADII_CM = [i/10 for i in range(21)]
NUMBER_FORMAT = '0.0000'
SIZE_LIMIT_BYTES = 20_000_000


def _sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as fh:
        for block in iter(lambda:fh.read(1024*1024),b''):
            h.update(block)
    return h.hexdigest()


def _record(path, role=None):
    path=Path(path).resolve()
    result={'path':path.relative_to(ROOT).as_posix(),'absolute_path':str(path),
            'bytes':path.stat().st_size,'sha256':_sha(path),'exists':True}
    if role:
        result['role']=role
    return result


def _json(path,value):
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')


def _qid(question_id):
    value=str(question_id).upper()
    if value in ['1','2','3','4']:
        value='Q'+value
    if value not in ['Q1','Q2','Q3','Q4']:
        raise ValueError('Export requires Q1, Q2, Q3 or Q4')
    return value


def _output_times(question_id,end,event):
    if question_id=='Q1':
        if end<1800-1e-7:
            raise ValueError('Q1 requires a solved interval covering 0..1800s')
        return np.arange(1801,dtype=float)
    if question_id=='Q2':
        if abs(end-round(end))>1e-7:
            raise ValueError('Q2 needs an integer post-crossing verification endpoint')
        return np.arange(int(round(end))+1,dtype=float)
    if event is None:
        raise ValueError('Q3/Q4 final export requires an actual critical drying event')
    if not 0<=event<=end:
        raise ValueError('Critical event must lie inside the solved interval')
    return np.unique(np.r_[np.arange(math.floor(end/60)+1,dtype=float)*60,event,end])


def _q4(x):
    if x is None:
        return None
    if not math.isfinite(float(x)):
        raise ValueError('A non-finite value cannot be exported as a numerical result')
    return float(round(float(x),4))


def _template(question_id):
    path=TEMPLATE_DIR/f'result{question_id[-1]}.xlsx'
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        header=next(wb.worksheets[0].iter_rows(values_only=True))
        names=wb.sheetnames
        return {'record':_record(path,'original_result_template'),'time_header':header[0],
                'sheet_names':names,'surface_header':header[-1] if question_id=='Q4' else None}
    finally:
        wb.close()


def _run_provenance(run):
    module=sys.modules[run.model.__class__.__module__]
    core=Path(inspect.getfile(run.model.__class__)).resolve()
    loaded=getattr(module,'LOADED_CODE_SHA256',None)
    if loaded is not None and loaded!=_sha(core):
        raise RuntimeError('Core file changed after this Run implementation was imported')
    if _sha(SOURCE)!=LOADED_EXPORTER_SHA256:
        raise RuntimeError('Exporter changed after import; restart for consistent provenance')
    input_records=list(getattr(run.model,'input_records',[]))
    for rec in input_records:
        if _sha(ROOT/rec['path'])!=rec['sha256']:
            raise RuntimeError('Observed input changed after solve: '+rec['path'])
    return {'core':_record(core,'solver_source'),'exporter':_record(SOURCE,'exporter_source'),
            'inputs':input_records,'settings':asdict(run.model.settings),
            'event_s':None if run.event_s is None else float(run.event_s),
            'end_s':float(run.end_s),'source_kind':'live_Run_dense_solution',
            'per_second_values_are_not_interpolated_from_60s_npz':True}


def _field_block(run,times,qid):
    T,C=run.fields(times,radii_m=RADII_M)
    radius=np.asarray(run.model.radius(times),dtype=float)
    if radius.ndim==0:
        radius=np.full(len(times),radius)
    surface_C=None
    if qid=='Q4':
        _,surface=run.fields(times,material_x=np.array([1.0]))
        surface_C=surface[:,0]
        inside=RADII_M[None,:] <= radius[:,None]+1e-12

        near_boundary=inside & (~np.isfinite(C))
        if np.any(near_boundary):
            C=np.where(near_boundary,surface_C[:,None],C)
        T=np.where(inside,T,np.nan)
        C=np.where(inside,C,np.nan)
        if not np.all(np.isfinite(C[inside])) or not np.all(np.isfinite(surface_C)):
            raise ValueError('Non-finite moisture inside Q4 domain')
    else:
        inside=np.ones_like(C,dtype=bool)
        if not np.all(np.isfinite(T)) or not np.all(np.isfinite(C)):
            raise ValueError('Non-finite output inside fixed domain')
    return T-273.15,C,radius,surface_C,inside


def _raw_headers(qid):
    h=['time_s','radius_m']
    if qid in ['Q1','Q2']:
        h += [f'T_C_r_{i/10:.1f}_cm' for i in range(21)]
    h += [f'C_r_{i/10:.1f}_cm' for i in range(21)]
    if qid=='Q4':
        h += ['C_surface']
    return h


def _raw_row(qid,t,T,C,R,surface,inside):
    row=[float(t),float(R)]
    if qid in ['Q1','Q2']:
        row+=list(map(float,T))
    row += [float(c) if valid else None for c,valid in zip(C,inside)]
    if qid=='Q4':
        row.append(float(surface))
    return row


def _rounded_sheets(qid,raw):
    t,R=raw[:2]
    if qid in ['Q1','Q2']:
        return {'温度':[_q4(t)]+[_q4(v) for v in raw[2:23]],
                '水分浓度':[_q4(t)]+[_q4(v) for v in raw[23:44]]}
    if qid=='Q3':
        return {'Sheet1':[_q4(t)]+[_q4(v) for v in raw[2:23]]}
    return {'Sheet1':[_q4(t)]+[_q4(v) for v in raw[2:24]],
            '半径':[_q4(t),_q4(R*100)]}


def _sheet(wb,name,header):
    ws=wb.create_sheet(name)
    ws.freeze_panes='B2'
    ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=30
    for j in range(2,len(header)+1):
        ws.column_dimensions[get_column_letter(j)].width=13 if j<len(header) else 16
    ws.row_dimensions[1].height=34
    cells=[]
    for value in header:
        cell=WriteOnlyCell(ws,value=value)
        cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
        cell.fill=PatternFill(fill_type='solid',fgColor='334155')
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        if isinstance(value,(int,float)):
            cell.number_format='0.0'
        cells.append(cell)


    cache=ROOT/'tmp/cache/openpyxl_exports'
    cache.mkdir(parents=True,exist_ok=True)
    previous_tempdir=tempfile.tempdir
    try:
        tempfile.tempdir=str(cache)
        ws.append(cells)
    finally:
        tempfile.tempdir=previous_tempdir
    return ws


def _append_numeric(ws,row):
    cells=[]
    for value in row:
        if value is None:
            cells.append(None)
        else:
            cell=WriteOnlyCell(ws,value=value)
            cell.number_format=NUMBER_FORMAT
            cells.append(cell)
    ws.append(cells)


def _paper_tables(run,qid,directory):
    if qid=='Q1':
        ts=np.array([100,300,600,900,1200,1500,1800],dtype=float)
    elif qid=='Q2':
        ts=np.arange(1,7,dtype=float)*1800
        if run.end_s<10800-1e-7:
            raise ValueError('Q2 paper tables require 3h of actual solution')
    else:
        from q3_model import completion
        reported_end=completion(run)['reported_time_s']
        ts=np.unique(np.r_[np.arange(21600.,reported_end+1e-8,21600.),reported_end])
    rs=np.array([0.,.005,.01,.015,.02])
    T,C=run.fields(ts,radii_m=rs)
    T=T-273.15
    R=np.asarray(run.model.radius(ts))
    headers=['时间/s' if qid=='Q1' else '时间/h']+[0,0.5,1,1.5,2]
    if qid=='Q4':
        headers += ['药材表面']
        _,Cs=run.fields(ts,material_x=[1.])
    outputs=[]
    fields=[('temperature',T),('moisture',C)] if qid in ['Q1','Q2'] else [('moisture',C)]
    for kind,values in fields:
        path=directory/f'{qid.lower()}_paper_{kind}.csv'
        with path.open('w',encoding='utf-8-sig',newline='') as fh:
            w=csv.writer(fh);w.writerow(headers)
            for i,t in enumerate(ts):
                row=[f'{t if qid=="Q1" else t/3600:.4f}']
                for j,r in enumerate(rs):
                    outside=qid=='Q4' and r>R[i]+1e-12
                    value=values[i,j]
                    if not outside and not np.isfinite(value) and qid=='Q4' and abs(r-R[i])<=1e-12:
                        value=Cs[i,0]
                    row.append('' if outside else f'{float(value):.4f}')
                if qid=='Q4':
                    row += [f'{Cs[i,0]:.4f}']
                w.writerow(row)
        outputs.append(_record(path,'paper_table'))
    if qid=='Q4':
        path=directory/'q4_paper_radius.csv'
        with path.open('w',encoding='utf-8-sig',newline='') as fh:
            w=csv.writer(fh);w.writerow(['时间/h','表面半径/cm'])
            w.writerows([[f'{t/3600:.4f}',f'{r*100:.4f}'] for t,r in zip(ts,R)])
        outputs.append(_record(path,'paper_surface_coordinates'))
    return outputs


def _validate_paper_tables(manifest,run):


    qid=manifest['question_id']
    if run is None:
        return {'status':'NOT_REQUESTED','checked_cells':0,'files':[],
                'scope':'Paper CSV hashes only; no independent live Run available.'},[]
    errors=[];records=manifest['paper_tables'];completion_info=None
    if qid=='Q1':
        times=np.asarray([100.,300.,600.,900.,1200.,1500.,1800.])
        expected_kinds=['temperature','moisture']
    elif qid=='Q2':
        times=1800.*np.arange(1,7)
        expected_kinds=['temperature','moisture']
    else:
        from q3_model import completion
        completion_info=completion(run)
        final_s=float(completion_info['reported_time_s'])


        regular=[21600.*k for k in range(1,math.floor(final_s/21600.)+1)
                 if 21600.*k<final_s-1e-8]
        times=np.asarray([*regular,final_s],dtype=float)
        expected_kinds=['moisture']+(['radius'] if qid=='Q4' else [])
        stored=manifest.get('paper_table_time_contract',{})
        for key in ['reported_time_s','reported_drying_time_h']:
            if key not in stored or not _numeric_equal(stored[key],completion_info[key]):
                errors.append('Paper CSV report-time metadata differs from independent completion: '+key)
        if not _numeric_equal(stored.get('reported_time_h'),completion_info['reported_drying_time_h']):
            errors.append('Paper CSV reported_time_h alias differs from independent completion')
    if times[-1]>run.end_s+1e-7:
        errors.append('Paper CSV contract exceeds the live solution horizon')
        return {'status':'FAIL','checked_cells':0,'files':[]},errors
    expected_names={f'{qid.lower()}_paper_{kind}.csv':kind for kind in expected_kinds}
    actual_names=[Path(r['path']).name for r in records]
    if set(actual_names)!=set(expected_names) or len(actual_names)!=len(expected_names):
        errors.append('Missing, duplicate, or unexpected paper CSV artifact')
    files=[];total_cells=0
    for rec in records:
        path=ROOT/rec['path'];kind=expected_names.get(path.name)
        if kind is None:continue
        before_errors=len(errors)
        with path.open('r',encoding='utf-8-sig',newline='') as fh:
            reader=csv.reader(fh);header=next(reader,None);rows=list(reader)
        desired_time='时间/s' if qid=='Q1' else '时间/h'
        if kind=='radius':
            correct_header=['时间/h','表面半径/cm']
            if header!=correct_header:
                errors.append(path.name+': radius/time units or header mismatch')
                continue
            radii=None
            expected=np.asarray(run.model.radius(times),dtype=float).reshape(-1,1)*100.
            inside=np.ones_like(expected,dtype=bool)
            unit='cm';surface_column=None
        else:
            wanted_columns=7 if qid=='Q4' else 6
            if header is None or len(header)!=wanted_columns or header[0]!=desired_time:
                errors.append(path.name+': time/column header mismatch');continue
            try:
                radii=np.asarray([float(v)*.01 for v in header[1:6]])
            except ValueError:
                errors.append(path.name+': nonnumeric physical-radius header');continue
            if not np.all(np.isfinite(radii)) or not np.allclose(radii,[0.,.005,.01,.015,.02],rtol=0,atol=1e-12):
                errors.append(path.name+': physical-radius grid must be 0,0.5,1,1.5,2 cm');continue
            if qid=='Q4' and header[-1]!='药材表面':
                errors.append(path.name+': distinct material surface column missing');continue
            temperature,moisture=run.fields(times,radii_m=radii)
            expected=np.asarray(temperature)-273.15 if kind=='temperature' else np.asarray(moisture)
            inside=np.ones_like(expected,dtype=bool);surface_column=None
            if qid=='Q4':
                physical_R=np.asarray(run.model.radius(times),dtype=float)
                inside=radii[None,:] <= physical_R[:,None]+1e-12
                _,surface_values=run.fields(times,material_x=np.asarray([1.]))


                coincidence=np.abs(radii[None,:]-physical_R[:,None])<=1e-12
                expected=np.where(coincidence & ~np.isfinite(expected),surface_values,expected)
                expected=np.column_stack([expected,surface_values[:,0]])
                inside=np.column_stack([inside,np.ones(len(times),dtype=bool)])
                surface_column=7
            unit='degC' if kind=='temperature' else 'kg/kg'
        if len(rows)!=len(times):
            errors.append(f'{path.name}: expected {len(times)} data rows, read {len(rows)}')
        checked=0;blanks=0;max_field_error=0.;max_time_error=0.
        for i,(row,t) in enumerate(zip(rows,times)):
            if len(row)!=len(header):
                errors.append(f'{path.name}: column count mismatch at row {i+2}');continue
            time_value=float(t if qid=='Q1' else t/3600.)
            if row[0]!=format(time_value,'.4f'):
                errors.append(f'{path.name}: time label mismatch at row {i+2}; query time is {t:.17g}s')
            else:max_time_error=max(max_time_error,abs(float(row[0])-time_value))
            checked+=1
            for j,cell in enumerate(row[1:]):
                checked+=1
                if not inside[i,j]:
                    blanks+=1
                    if cell!='':errors.append(f'{path.name}: domain outside must be blank at row {i+2}, column {j+2}')
                    continue
                value=float(expected[i,j])
                if not math.isfinite(value):
                    errors.append(f'{path.name}: live value is not finite inside material');continue
                if not re.fullmatch(r'-?\d+\.\d{4}',cell):
                    errors.append(f'{path.name}: finite four-decimal number required at row {i+2}, column {j+2}')
                    continue
                actual=float(cell)
                if not math.isfinite(actual) or abs(actual-float(format(value,'.4f')))>1e-10:
                    errors.append(f'{path.name}: independently queried value mismatch at row {i+2}, column {j+2}')
                max_field_error=max(max_field_error,abs(actual-value))
        total_cells+=checked
        files.append({'artifact':_record(path),'kind':kind,'status':'PASS' if len(errors)==before_errors else 'FAIL',
            'header':header,'data_rows':len(rows),'expected_data_rows':len(times),'checked_cells':checked,
            'source_query_times_s':times.tolist(),'display_time_unit':'s' if qid=='Q1' else 'h',
            'max_display_time_rounding_error':max_time_error,'field_unit':unit,
            'max_field_rounding_error_from_live_value':max_field_error,'outside_domain_blank_cells':blanks,
            'physical_radii_m':None if radii is None else radii.tolist(),
            'surface_column_1based':surface_column,'surface_query':'material_x=1; never fixed 2 cm' if surface_column else None})
    return {'status':'FAIL' if errors else 'PASS','checked_cells':total_cells,'files':files,
        'completion':completion_info,'completion_source':_record(SOURCE.with_name('q3_model.py')) if completion_info else None,
        'scope':'Every paper CSV cell independently re-evaluated from live Run and fixed question contracts; no generation arrays or 60s archives reused.',
        'rounding_limit':'CSV stores four decimals; matching rounded values cannot distinguish sub-rounding perturbations of the original source value.'},errors


def export_question(run,question_id,output_dir,*,chunk_rows=1000,overwrite=False):


    started=time.perf_counter()
    qid=_qid(question_id)
    directory=Path(output_dir).resolve()
    if not directory.is_relative_to(ROOT) or directory.is_relative_to(ROOT/'problem_files'):
        raise ValueError('Outputs must be inside the competition workspace, outside problem_files')
    if not 1<=int(chunk_rows)<=2000:
        raise ValueError('chunk_rows must be between 1 and 2000')
    directory.mkdir(parents=True,exist_ok=True)
    stem='result'+qid[-1]
    book_path=directory/(stem+'.xlsx')
    raw_path=directory/(stem+'_unrounded.csv.gz')
    manifest_path=directory/(stem+'.export.json')
    if not overwrite and any(p.exists() for p in [book_path,raw_path,manifest_path]):
        raise FileExistsError('Export target already exists; use a new output version')
    provenance=_run_provenance(run)
    allowed={'Q1':{'Q1'},'Q2':{'Q2','Q23'},'Q3':{'Q2','Q3','Q23'},'Q4':{'Q4'}}
    if run.model.settings.question not in allowed[qid]:
        raise ValueError('Run physics/question do not match the requested output question')
    if qid=='Q4' and not run.model.settings.shrink:
        raise ValueError('Formal Q4 output requires the shrinking-domain Run, not its fixed-radius control')
    template=_template(qid)
    times=_output_times(qid,run.end_s,run.event_s)
    if len(times)+1>1_048_576:
        raise ValueError('Required full time grid exceeds the XLSX worksheet row limit; do not truncate')
    workbook=Workbook(write_only=True)
    workbook.properties.creator=''
    workbook.properties.lastModifiedBy=''
    workbook.properties.title=f'问题{qid[-1]}结果'
    headers={name:[template['time_header']]+RADII_CM for name in template['sheet_names']}
    if qid=='Q4':
        headers['Sheet1'].append(template['surface_header'])
        headers['半径']=['时间/s','药材表面半径/cm']
    sheets={name:_sheet(workbook,name,header) for name,header in headers.items()}
    counts={name:{'data_rows':0,'total_rows':1,'columns':len(header),'outside_domain_blank_cells':0,
                  'first_time_s':None,'last_time_s':None,'header':header} for name,header in headers.items()}
    samples_idx=set(np.linspace(0,len(times)-1,min(13,len(times)),dtype=int).tolist())
    samples_idx.update([0,min(1,len(times)-1),len(times)-1])
    if run.event_s is not None:
        samples_idx.update(np.flatnonzero(times==run.event_s).tolist())
    samples=[]
    with gzip.open(raw_path,'wt',encoding='utf-8',newline='',compresslevel=6) as fh:
        raw_writer=csv.writer(fh);raw_writer.writerow(_raw_headers(qid))
        for start in range(0,len(times),int(chunk_rows)):
            ts=times[start:start+int(chunk_rows)]
            T,C,R,Cs,inside=_field_block(run,ts,qid)
            for i,t in enumerate(ts):
                raw=_raw_row(qid,t,T[i],C[i],R[i],None if Cs is None else Cs[i],inside[i])
                raw_writer.writerow(['' if v is None else format(v,'.17g') for v in raw])
                values=_rounded_sheets(qid,raw)
                for name,row in values.items():
                    _append_numeric(sheets[name],row)
                    count=counts[name]
                    count['data_rows']+=1;count['total_rows']+=1
                    count['outside_domain_blank_cells']+=sum(v is None for v in row)
                    if count['first_time_s'] is None:
                        count['first_time_s']=row[0]
                    count['last_time_s']=row[0]
                if start+i in samples_idx:
                    samples.append({'data_row_index':start+i,'source_time_s':float(t),
                                    'raw_values':raw,'rounded_sheets':values})
    workbook.save(book_path)
    paper_records=_paper_tables(run,qid,directory)
    paper_time_contract={'unit':'s' if qid=='Q1' else 'h','stored_decimals':4,
                         'query_uses_unrounded_seconds':True}
    if qid in ['Q3','Q4']:
        from q3_model import completion
        completed=completion(run)
        paper_time_contract.update({key:completed[key] for key in
                                   ['reported_time_s','reported_drying_time_h','critical_event_s',
                                    'max_C_at_reported_time','conservative_post_verification_s']})
        paper_time_contract['reported_time_h']=completed['reported_drying_time_h']
        paper_time_contract['regular_step_s']=21600.
        paper_time_contract['endpoint_convention']='6-hour samples plus upward-reported strict-drying time; distinct from workbook post-verification endpoint'
    end_provenance=_run_provenance(run)
    if end_provenance!=provenance or _sha(ROOT/template['record']['path'])!=template['record']['sha256']:
        raise RuntimeError('Source, settings or template changed during export')
    book_record=_record(book_path,'contest_result_workbook')
    archive_record=_record(raw_path,'internal_unrounded_output_grid_archive')
    manifest={
        'schema_version':'1.0','question_id':qid,'generated_by':'paper_output/code/modeling/export_outputs.py',
        'generated_at':datetime.now(timezone.utc).isoformat(),'status':'EXPORTED_PENDING_STREAM_READBACK',
        'workbook':book_record,'unrounded_archive':archive_record,'paper_tables':paper_records,
        'paper_table_time_contract':paper_time_contract,
        'provenance':provenance,'template':template['record'],'sheets':counts,
        'time_grid':{'start_s':0.,'end_s':float(times[-1]),'regular_step_s':1 if qid in ['Q1','Q2'] else 60,
                     'critical_event_s':provenance['event_s'] if qid in ['Q3','Q4'] else None,
                     'include_post_verification_endpoint':qid!='Q1',
                     'count':len(times),'rounding_duplicate_time_count':int(np.sum(np.diff(np.round(times,4))==0))},
        'radius_grid':{'fixed_radii_cm':RADII_CM,'surface_column':qid=='Q4','surface_coordinate_sheet':'半径' if qid=='Q4' else None},
        'rounding':{'decimal_places':4,'storage':'numeric rounded to four decimals','number_format':NUMBER_FORMAT,
                    'raw_archive':'17 significant digits, source times in seconds and radius in metres',
                    'threshold_judgement':'use original full precision event and max C, never rounded 0.1500'},
        'domain_rule':{'outside_values':'blank (None); never zero','comparison_tolerance_m':1e-12,
                       'Q4_surface_is_not_fixed_2cm':True},
        'samples':samples,'elapsed_s':time.perf_counter()-started,
        'size':{'workbook_bytes':book_record['bytes'],'workbook_MiB':book_record['bytes']/2**20,
                'unrounded_archive_bytes':archive_record['bytes'],'warning_limit_bytes':SIZE_LIMIT_BYTES,
                'workbook_exceeds_20_decimal_MB':book_record['bytes']>SIZE_LIMIT_BYTES,
                'archive_is_internal_not_automatically_in_submission_zip':True,
                'complete_support_archive_still_requires_actual_size_check':True,
                'no_required_time_or_space_values_removed':True},
        'validation_required':'validate_exports([record], runs={question_id: run}); full workbook/archive comparison, live Run workbook samples, and every paper CSV cell independently queried from live Run',
        'visual_studio_gui':'pending','human_review':'pending',
    }
    _json(manifest_path,manifest)
    return {'question_id':qid,'workbook_path':str(book_path),'manifest_path':str(manifest_path),
            'artifacts':[book_record,archive_record]+paper_records+[_record(manifest_path,'export_manifest')],
            'sheets':counts,'size':manifest['size'],'validation_status':'pending'}


def _manifest_path(item):
    if isinstance(item,dict):
        return Path(item['manifest_path'])
    path=Path(item)
    return path.with_suffix('.export.json') if path.suffix=='.xlsx' else path


def _numeric_equal(a,b,tolerance=1e-10):
    return a is None and b is None or (a is not None and b is not None and
        isinstance(a,(int,float)) and not isinstance(a,bool) and math.isfinite(float(a)) and abs(float(a)-float(b))<=tolerance)


def validate_exports(paths,runs=None):


    if isinstance(paths,(str,Path,dict)):
        paths=[paths]
    runs={} if runs is None else runs
    results=[]
    for item in paths:
        started=time.perf_counter()
        manifest_path=_manifest_path(item).resolve()
        manifest=json.loads(manifest_path.read_text(encoding='utf-8'))
        qid=manifest['question_id'];errors=[];checked_cells=0;checked_sheets={}
        for rec in [manifest['workbook'],manifest['unrounded_archive'],manifest['template']]+manifest['paper_tables']:
            if not (ROOT/rec['path']).exists() or _sha(ROOT/rec['path'])!=rec['sha256']:
                errors.append('Hash mismatch or missing artifact: '+rec['path'])
        if errors:
            raise RuntimeError('; '.join(errors))
        times=_output_times(qid,manifest['provenance']['end_s'],manifest['provenance']['event_s'])
        wb=load_workbook(ROOT/manifest['workbook']['path'],read_only=True,data_only=False)
        if wb.sheetnames!=list(manifest['sheets']):
            errors.append('Sheet names/order differ from export contract')
        try:
            for name,spec in manifest['sheets'].items():
                ws=wb[name]
                rows=ws.iter_rows(min_row=1,max_col=spec['columns'])
                header=[cell.value for cell in next(rows)]
                if header!=spec['header']:
                    errors.append(name+': header/physical radius grid mismatch')
                count=0;blanks=0;first=None;last=None;max_rounding_error=0.
                with gzip.open(ROOT/manifest['unrounded_archive']['path'],'rt',encoding='utf-8',newline='') as rawfh:
                    source_rows=csv.reader(rawfh)
                    if next(source_rows)!=_raw_headers(qid):
                        errors.append('Unrounded archive schema mismatch')
                    for index,source_row in enumerate(source_rows):
                        try:
                            cells=next(rows)
                        except StopIteration:
                            errors.append(name+': workbook ends before raw archive');break
                        raw=[None if v=='' else float(v) for v in source_row]
                        if index>=len(times) or abs(raw[0]-times[index])>1e-8:
                            errors.append(name+': source time schedule mismatch at '+str(index))
                        if qid=='Q4':
                            for j,r in enumerate(RADII_M):
                                outside=r>raw[1]+1e-12
                                if (raw[2+j] is None)!=outside:
                                    errors.append(name+': incorrect moving-domain mask at '+str(index))
                        elif any(v is None for v in raw):
                            errors.append(name+': unexpected blank in fixed domain')
                        expected=_rounded_sheets(qid,raw)[name]
                        for j,(cell,value) in enumerate(zip(cells,expected)):
                            if not _numeric_equal(cell.value,value):
                                errors.append(f'{name}: value mismatch at row {index+2}, column {j+1}')
                            if value is not None:
                                if cell.number_format!=NUMBER_FORMAT:
                                    errors.append(f'{name}: four-decimal format missing at row {index+2}, column {j+1}')
                                if not math.isfinite(float(cell.value)) or abs(float(cell.value)*1e4-round(float(cell.value)*1e4))>1e-5:
                                    errors.append(f'{name}: non-finite or non-four-decimal stored number')
                                max_rounding_error=max(max_rounding_error,abs(float(cell.value)-value))
                            else:
                                blanks+=1
                        checked_cells+=len(cells);count+=1
                        if first is None:first=cells[0].value
                        last=cells[0].value
                        if len(errors)>30:
                            raise RuntimeError('Export validation failed: '+'; '.join(errors[:30]))
                    if next(rows,None) is not None:
                        errors.append(name+': workbook has uncontracted extra rows')
                if count!=spec['data_rows'] or count!=len(times):
                    errors.append(name+': row count differs from time grid')
                if blanks!=spec['outside_domain_blank_cells']:
                    errors.append(name+': blank count mismatch')
                checked_sheets[name]={'data_rows':count,'total_rows':count+1,'first_time_s':first,
                    'last_time_s':last,'outside_domain_blank_cells':blanks,'value_comparison_max_abs_error':max_rounding_error}
        finally:
            wb.close()
        run=runs.get(qid)
        live={'status':'NOT_REQUESTED','samples':0,'max_abs_unrounded_difference':None}
        if run is not None:
            current=_run_provenance(run)
            if current!=manifest['provenance']:
                errors.append('Live Run provenance differs from the exported Run')
            samples=manifest['samples'];ts=np.array([s['source_time_s'] for s in samples])
            T,C,R,Cs,inside=_field_block(run,ts,qid)
            difference=0.
            for i,sample in enumerate(samples):
                actual=_raw_row(qid,ts[i],T[i],C[i],R[i],None if Cs is None else Cs[i],inside[i])
                for a,b in zip(actual,sample['raw_values']):
                    if a is None or b is None:
                        if not (a is None and b is None):errors.append('Live Run sample domain mismatch')
                    else:
                        difference=max(difference,abs(a-b))
            if difference>1e-10:
                errors.append('Independent live Run sample mismatch exceeds 1e-10')
            live={'status':'PASS' if difference<=1e-10 else 'FAIL','samples':len(samples),
                  'max_abs_unrounded_difference':difference}
        paper_validation,paper_errors=_validate_paper_tables(manifest,run)
        errors.extend(paper_errors)
        report={'schema_version':'1.0','question_id':qid,'generated_by':'paper_output/code/modeling/export_outputs.py',
                'generated_at':datetime.now(timezone.utc).isoformat(),
                'status':'FAIL' if errors else 'PASS','fully_verified_with_live_Run':not errors and run is not None,
                'checked_cells':checked_cells,'sheets':checked_sheets,'live_Run_sample_validation':live,
                'paper_CSV_live_Run_validation':paper_validation,
                'errors':errors,'elapsed_s':time.perf_counter()-started,
                'workbook':manifest['workbook'],'export_manifest':_record(manifest_path),
                'scope':'Full workbook/raw-archive validation; independent workbook live Run samples and every paper CSV time/position/value/surface/radius cell when live Run is provided.',
                'four_decimals_are_storage_and_display_not_physical_accuracy_claim':True,
                'visual_render':'pending','visual_studio_gui':'pending','human_review':'pending'}
        report_path=manifest_path.with_name(manifest_path.name.replace('.export.json','.validation.json'))
        _json(report_path,report);report['report_path']=str(report_path);results.append(report)
        if errors:
            raise RuntimeError('Export validation failed: '+'; '.join(errors[:30]))
    return {'status':'PASS' if all(r['status']=='PASS' for r in results) else 'FAIL',
            'fully_verified_with_live_Run':all(r['fully_verified_with_live_Run'] for r in results),'exports':results}


def _paper_csv_selfcheck(directory):

    from drying_core import Settings,solve_case
    import copy
    directory=Path(directory).resolve()
    directory.mkdir(parents=True,exist_ok=True)
    summaries=[];negative_checks=[]
    for qid in ['Q1','Q2','Q3','Q4']:
        settings={'question':'Q23' if qid in ['Q2','Q3'] else qid,
                  'intervals':40,'face_scheme':'kirchhoff','shrink':qid=='Q4',
                  'rtol':1e-8,'atol_temperature':1e-8,'atol_moisture':1e-10,
                  'early_max_step_s':5.}
        if qid=='Q2':
            settings.update(constant_D=2e-8,beta=8e-6,horizon_h=24.)
        print('PAPER_CSV_SELFCHECK '+qid+': actual N40 solve',flush=True)
        run=solve_case(Settings(**settings))
        try:
            if qid=='Q2' and (run.event_s is None or run.end_s<10800.):
                raise RuntimeError('Accelerated export test must still cover 3h and a real drying event')
            exported=export_question(run,qid,directory/qid)
            checked=validate_exports([exported],runs={qid:run})
            manifest=json.loads(Path(exported['manifest_path']).read_text(encoding='utf-8'))
            defects={'Q1':['temperature_value'],'Q2':['moisture_value'],
                     'Q3':['last_time_label'],'Q4':['outside_nonblank','surface_missing','radius_value']}[qid]
            for defect in defects:
                changed=copy.deepcopy(manifest)
                kind='temperature' if defect=='temperature_value' else 'radius' if defect=='radius_value' else 'moisture'
                selected=next(r for r in changed['paper_tables'] if Path(r['path']).name==f'{qid.lower()}_paper_{kind}.csv')
                original_path=ROOT/selected['path']
                with original_path.open('r',encoding='utf-8-sig',newline='') as fh:rows=list(csv.reader(fh))
                if defect in ['temperature_value','moisture_value']:
                    rows[1][1]=format(float(rows[1][1])+.01,'.4f')
                elif defect=='last_time_label':
                    rows[-1][0]=format(float(rows[-1][0])-.001,'.4f')
                elif defect=='outside_nonblank':
                    ri,ci=next((i,j) for i in range(1,len(rows)) for j in range(1,6) if rows[i][j]=='')
                    rows[ri][ci]='0.0000'
                elif defect=='surface_missing':rows[-1][-1]=''
                elif defect=='radius_value':rows[-1][1]=format(float(rows[-1][1])+.01,'.4f')
                target=directory/'deliberate_defects'/defect/original_path.name
                target.parent.mkdir(parents=True,exist_ok=True)
                with target.open('w',encoding='utf-8-sig',newline='') as fh:csv.writer(fh).writerows(rows)
                selected.update(_record(target,selected.get('role')))
                result,errors=_validate_paper_tables(changed,run)
                if not errors or result['status']!='FAIL':
                    raise AssertionError('Independent paper validator missed deliberate defect: '+defect)
                negative_checks.append({'question_id':qid,'defect':defect,'correctly_rejected':True,
                                        'errors':errors,'artifact':_record(target)})
            summaries.append({'question_id':qid,'settings':asdict(run.model.settings),'export':exported,
                              'validation':checked,'event_s':run.event_s,'end_s':run.end_s,
                              'Q2_is_accelerated_constant_D_export_test_only':qid=='Q2'})
            print(json.dumps({'question_id':qid,'status':'PASS','end_s':run.end_s,
                'paper_cells':checked['exports'][0]['paper_CSV_live_Run_validation']['checked_cells']},ensure_ascii=False),flush=True)
        finally:
            run.close()
    report={'status':'PASS','generated_at':datetime.now(timezone.utc).isoformat(),
            'scope':'N40 exporter fidelity only. Q2 uses constant_D=2e-8 and beta=8e-6 solely to shorten the real full-event export test; these are not formal physical results. Q1/Q3/Q4 use their nominal physics.',
            'source':_record(SOURCE),'exports':summaries,'negative_checks':negative_checks,
            'negative_test_scope':'Copied paper CSVs only; six deliberate defects must be rejected even after their recorded hashes are updated.',
            'visual_studio_gui':'pending','human_review':'pending'}
    _json(directory/'paper_csv_selfcheck_report.json',report)
    print(json.dumps({'status':'PASS','questions':4,'deliberate_defects_rejected':len(negative_checks),
                      'report':str(directory/'paper_csv_selfcheck_report.json')},ensure_ascii=False),flush=True)
    return 0


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--selfcheck',action='store_true')
    parser.add_argument('--paper-selfcheck',action='store_true')
    parser.add_argument('--output-dir',default='paper_output/results/export_selfcheck')
    args=parser.parse_args()
    if not args.selfcheck and not args.paper_selfcheck:
        parser.error('Use the Python API with final Runs, or --selfcheck for a real Q1 validation export')
    if Path.cwd().resolve()!=ROOT:
        raise RuntimeError('Use the competition workspace as cwd')
    if args.paper_selfcheck:
        return _paper_csv_selfcheck(ROOT/args.output_dir)
    from drying_core import Settings,solve_case


    run=solve_case(Settings(question='Q1',intervals=40,face_scheme='kirchhoff',
                           rtol=1e-8,atol_temperature=1e-8,atol_moisture=1e-10,early_max_step_s=5.))
    record=export_question(run,'Q1',ROOT/args.output_dir)
    checks=validate_exports([record],runs={'Q1':run})
    output=Path(args.output_dir).resolve()
    note={'status':checks['status'],'scope':'Actual Q1 export selfcheck only, not final production or numerical-accuracy acceptance',
          'settings':asdict(run.model.settings),'export':record,'validation':checks}
    _json(output/'selfcheck_summary.json',note)
    print(json.dumps({'status':checks['status'],'fully_verified_with_live_Run':checks['fully_verified_with_live_Run'],
                      'output':str(output),'size':record['size'],'sheets':record['sheets']},ensure_ascii=False,indent=2))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
