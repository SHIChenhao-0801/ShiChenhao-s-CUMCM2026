"""Independent data/portability audit for the six-figure handoff (no model runs)."""
from pathlib import Path
import csv, hashlib, json, re, sys, zipfile
import numpy as np
from openpyxl import load_workbook
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
PKG = ROOT / 'paper_output/handoff/six_figures_20260912'
PROD = ROOT / 'paper_output/results/production/final_v6a'
QA = ROOT / 'paper_output/qa/figure_package_20260912'
manifest = json.loads((PKG / '数据清单.json').read_text(encoding='utf-8'))
checks = []
data = {}


def record(name, **details):
    checks.append({'name': name, 'status': 'PASS', **details})


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def read_csv(fig, filename):
    path = f'CSV/图{fig}/{filename}.csv'
    b = (PKG / path).read_bytes()
    assert b.startswith(b'\xef\xbb\xbf'), path
    with (PKG / path).open(encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        columns = next(reader)
        rows = list(reader)
    data[path] = (columns, rows)
    return rows


def compare(fig, filename, expected, tolerance=0):
    actual = np.array(read_csv(fig, filename), dtype=float)
    expected = np.array(expected, dtype=float)
    assert actual.shape == expected.shape, (filename, actual.shape, expected.shape)
    assert np.isfinite(actual).all(), filename
    error = float(np.max(np.abs(actual - expected))) if actual.size else 0
    assert error <= tolerance, (filename, error, tolerance)
    record(f'figure_{fig}_{filename}', rows=len(actual), columns=actual.shape[1], maximum_abs_error=error)


assert len(manifest['sources']) == 15
for s in manifest['sources']:
    rel = Path(s['file'])
    assert not rel.is_absolute() and '..' not in rel.parts
    assert sha(PKG / rel) == sha(ROOT / s['source']) == s['sha256']
record('all_original_source_copies', files=15)

with np.load(PROD / 'figures/fig_q1_profiles_data.npz') as z:
    for field, file in [('temperature_C', 'temperature'), ('C', 'moisture')]:
        compare(1, file, np.column_stack((z['radii_m'] * 100, z[field].T)))

with (ROOT / 'paper_output/data_cleaned/A_environment_observed.csv').open(encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
compare(2, 'observed_0_4h', [[float(r[k]) for k in ('time_h', 'temperature_C', 'air_moisture_kg_per_kg')] for r in rows])
assert len(rows) == 241
assert [float(rows[-1][k]) for k in ('time_h', 'temperature_C', 'air_moisture_kg_per_kg')] == [4, 50.165, 0.04986]
summary = json.loads((PROD / 'Q23/summary.json').read_text())
assert summary['settings']['boundary_extension'] == 'nominal'
assert summary['settings']['tail_temperature_C'] == 50
assert summary['settings']['tail_equilibrium'] == .05
compare(2, 'assumed_platform_4_60h', [[4, 50, .05, 0], [60, 50, .05, 1]])
completion = summary['completion']
event_s = completion['critical_event_s']

with np.load(PROD / 'Q23/sampled_solution.npz') as z:
    assert z['material_x'][0] == 0 and z['material_x'][-1] == 1
    select = z['times_s'] <= 14400
    compare(3, 'temperature_0_4h', np.column_stack((z['times_s'][select]/3600, z['T_K'][select, 0]-273.15, z['T_K'][select, -1]-273.15)))
    select = z['times_s'] <= event_s
    compare(3, 'moisture_to_event', np.column_stack((z['times_s'][select]/3600, z['C'][select, 0], z['C'][select, -1])))
    select = (z['times_s'] >= event_s - 120) & (z['times_s'] <= event_s + 2)
    compare(4, 'centre_near_event', np.column_stack((z['times_s'][select], z['times_s'][select]-event_s, z['C'][select, 0], z['C'][select, 0]-.15)))
    assert not np.any(z['times_s'] == completion['reported_time_s'])

with np.load(PROD / 'figures/fig_q3_drying_data.npz') as z:
    compare(4, 'drying_main', np.column_stack((z['times_s']/3600, z['centre_C'], z['surface_C'], z['max_C'], np.full(len(z['times_s']), .15))))
    select = z['times_s'] >= event_s-120
    compare(4, 'maximum_near_event', np.column_stack((z['times_s'][select], z['times_s'][select]-event_s, z['max_C'][select], z['max_C'][select]-.15)))
    assert z['times_s'][-1] == 206902

rows = read_csv(4, 'event_markers')
assert len(rows) == 2
assert [r[0] for r in rows] == ['critical_equality', 'strict_report']
expected = [[event_s, completion['critical_event_h'], 0, .15, 0, 0],
            [completion['reported_time_s'], completion['reported_drying_time_h'], completion['reported_time_s']-event_s,
             completion['max_C_at_reported_time'], completion['max_C_at_reported_time']-.15, 1]]
assert np.array_equal(np.array([r[1:] for r in rows], dtype=float), np.array(expected))
record('figure_4_event_markers', rows=2, strictly_feasible_report_point=True)

with (ROOT / 'paper_output/data_cleaned/A_radius_observed.csv').open(encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
compare(5, 'radius_observed', [[float(r['time_h']), float(r['radius_cm'])] for r in rows])
assert len(rows) == 145

q4event = json.loads((PROD / 'Q4/summary.json').read_text())['completion']['critical_event_s']
selected_rows = read_csv(5, 'selected_times_and_radii')
assert [r[0] for r in selected_rows] == ['0h', '6h', '12h', '24h', '48h', 'event']
with np.load(PROD / 'Q4/sampled_solution.npz') as z:
    requested = [0, 21600, 43200, 86400, 172800, q4event]
    pairs, selected = [], []
    for t in requested:
        indices = np.flatnonzero(z['times_s'] == t)
        assert len(indices) == 1
        i = indices[0]
        pairs += [z['material_x'] * z['radius_m'][i] * 100, z['C'][i]]
        selected.append([t, t/3600, z['radius_m'][i]*100])
    compare(5, 'moisture_profiles_xy_pairs', np.column_stack(pairs), tolerance=1e-14)
    assert np.array_equal(np.array([r[1:] for r in selected_rows], dtype=float), np.array(selected))
    assert np.array_equal(np.array(selected)[:, 2], [2, 1.374, 1.248, 1.204, 1.2, 1.2])
record('figure_5_selected_times_and_radii', rows=6, exact_existing_times=True, endpoint_reaches_surface=True)

scenarios = json.loads((ROOT / 'paper_output/results/crossvalidation/isotherm_closure_v1/isotherm_closure.json').read_text())['scenarios']
expected = []
for p in (1, 2, 4):
    row = [p]
    for question in ('Q23', 'Q4'):
        choices = [s for s in scenarios if s['question'] == question and s['isothermShapeP'] == p]
        assert len(choices) == 1
        s = choices[0]
        assert s['intervals'] == 800 and s['latentFraction'] == 0 and s['status'] == 'computed_event_only'
        row.append(s['event_h'])
    expected.append(row)
compare(6, 'scenario_comparison', expected)

assert len(data) == 14
assert set(data) == {t['csv'] for t in manifest['tables']}
for t in manifest['tables']:
    assert t['columns'] == data[t['csv']][0]
    assert t['data_rows'] == len(data[t['csv']][1])
    assert (PKG / t['source'].split('；')[0]).exists()
assert manifest['total_data_rows'] == sum(len(v[1]) for v in data.values())
record('manifest_completeness', tables=14, total_data_rows=manifest['total_data_rows'])

workbook_path = PKG / '六图绘图数据.xlsx'
if workbook_path.exists():
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    largest_error, cells = 0., 0
    for t in manifest['tables']:
        ws = wb[t['name']]
        columns, rows = data[t['csv']]
        observed = list(ws.iter_rows(min_row=6, max_row=6+len(rows), max_col=len(columns), values_only=True))
        assert list(observed[0]) == columns
        for got, want in zip(observed[1:], rows):
            for gv, wv in zip(got, want):
                try:
                    number = float(wv)
                except ValueError:
                    assert gv == wv
                else:
                    assert isinstance(gv, (int, float))
                    err = abs(gv-number)
                    assert err <= max(1e-12, abs(number)*1e-14), (t['name'], gv, wv)
                    largest_error = max(largest_error, err)
                    cells += 1
    wb.close()
    with zipfile.ZipFile(workbook_path) as z:
        assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
    record('xlsx_vs_independently_verified_csv', numeric_cells=cells, maximum_abs_error=largest_error, external_links=0)

docx_path = PKG / '六图绘图说明.docx'
if docx_path.exists():
    with zipfile.ZipFile(docx_path) as z:
        assert z.testzip() is None
        x = etree.fromstring(z.read('word/document.xml'))
    texts = x.xpath('//w:t/text()', namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
    text = '\n'.join(texts)
    assert not re.search(r'[A-Za-z]:[/\\]', text), 'Word contains local drive path'
    for t in manifest['tables']:
        assert t['name'] in text
        # The Word provides a figure-level CSV directory, then individual filenames.
        assert Path(t['csv']).parent.as_posix()+'/' in text
        assert Path(t['csv']).name in text
    record('docx_portable_table_references', all_14_csv_and_sheet_names=True, local_drive_paths=0)

status = 'PASS' if workbook_path.exists() and docx_path.exists() else 'PASS_DATA_PENDING_FINAL_ASSETS'
result = {'status':status, 'model_runs':0, 'checks':checks,
          'scope':'Frozen-source-to-CSV independent comparison; copied input hashes; manifest paths and data counts. XLSX/Word checks run when final assets exist. ZIP audit belongs to root final check.',
          'package':PKG.relative_to(ROOT).as_posix(),
          'final_assets':[{'file':p.relative_to(PKG).as_posix(),'sha256':sha(p)} for p in [workbook_path,docx_path] if p.exists()]}
(QA / 'agent_audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
(QA / 'agent_audit.md').write_text('# 六图数据包独立核验\n\n状态：'+status+'\n\n'+ '\n'.join('- '+c['name']+'：PASS' for c in checks)+'\n\n未重跑模型。导出数据逐值比对冻结来源；图5半径只允许乘法顺序导致的浮点误差。Excel采用数值精度容差核对。\n',encoding='utf-8')
print(json.dumps({'status':status, 'checks':len(checks), 'data_rows':manifest['total_data_rows']}, ensure_ascii=False))
