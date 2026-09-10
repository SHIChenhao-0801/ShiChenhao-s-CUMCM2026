"""Independent raw-input/scalar audit. Does not import or run modeling code."""
from __future__ import annotations
import ast
import csv
import hashlib
import json
import math
import pathlib
import platform
import sys
import time
from datetime import datetime, timezone
import openpyxl
import pymupdf

ROOT = pathlib.Path(__file__).resolve().parents[3]
OUT = pathlib.Path(__file__).resolve().parent
assert ROOT.name == '2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')

def stamp():
    return datetime.now(timezone.utc).isoformat()

def record(p):
    b = p.read_bytes()
    return {'path':p.relative_to(ROOT).as_posix(),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()}

def read_raw(path):
    w = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for s in w.worksheets:
        cells = list(s.iter_rows())
        vals = [[c.value for c in row] for row in cells]
        sheets.append({'name':s.title,'max_row':s.max_row,'max_column':s.max_column,
            'headers':vals[0],'data_rows':len(vals)-1,'rows':vals[1:],
            'missing_data_cells':sum(v is None for row in vals[1:] for v in row),
            'nonnumeric_data_cells':sum(not isinstance(v,(int,float)) for row in vals[1:] for v in row),
            'formula_cells':sum(c.data_type=='f' for row in cells for c in row),
            'merged_ranges': 'not exposed by read_only; separately checked worksheet XML',
            'time_unique_count':len(set(row[0] for row in vals[1:])),
            'duplicate_entire_rows':len(vals[1:])-len(set(tuple(row) for row in vals[1:])),
            'time_gap_s':sorted(set(vals[i+1][0]-vals[i][0] for i in range(1,len(vals)-1)))})
    w.close()
    import zipfile, xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        for i,s in enumerate(sheets,1):
            root=ET.fromstring(z.read(f'xl/worksheets/sheet{i}.xml'))
            s['merged_ranges']=[e.attrib['ref'] for e in root.findall('m:mergeCells/m:mergeCell',ns)]
    return sheets

def compare(raw, path, expected, tolerance):
    with path.open(encoding='utf-8-sig',newline='') as f:
        clean=list(csv.DictReader(f))
    if len(raw)!=len(clean):
        raise AssertionError(f'row mismatch {path}')
    details={name:{'count':0,'max_abs_error':0.,'max_error_raw_excel_row':None,'nonzero_float_differences':0,'failures':0,'tolerance':tolerance[name]} for name in expected(raw[0])}
    for i,(r,c) in enumerate(zip(raw,clean),2):
        for name,v in expected(r).items():
            delta=abs(float(c[name])-v)
            d=details[name]
            d['count']+=1
            d['nonzero_float_differences']+=delta!=0
            if delta>d['max_abs_error']:
                d['max_abs_error']=delta;d['max_error_raw_excel_row']=i
            d['failures']+=delta>tolerance[name]
    return {'rows':len(raw),'numeric_values_compared':sum(d['count'] for d in details.values()),'fields':details,'pass':not any(d['failures'] for d in details.values())}

def prop(q,C,T):
    # Independently transcribed from raw PDF pp.3-4, not imported from solver.
    if q=='Q1': return 820.,2600.,.36,7e-9*math.exp(-.89/C)
    if q=='Q23': return 650+128*C,1450+2736*C/(1+C),.21+.38*C/(1+C),2.4e-3*math.exp(-.45/C)*math.exp(-3850/T)
    if q=='Q4': return 760+90*C,1850+2150*C/(1+C),.12+.20*C/(1+C),4.2e-4*math.exp(-.30/C)*math.exp(-3850/T)
    raise ValueError(q)

def scalars(q,C,T,R,t):
    rho,cp,k,D=prop(q,C,T); B=rho*cp; alpha=k/B
    return {'question':q,'C_kgwater_per_kgdry':C,'T_K':T,'R_m':R,'reference_time_s':t,
        'rho_kg_m3':rho,'cp_J_kgK':cp,'k_W_mK':k,'B_J_m3K':B,'D_m2_s':D,'alpha_m2_s':alpha,
        'Bi_heat_radius':25*R/k,'Bi_mass_radius':8e-7*R/D,'alpha_over_D_effective_Le':alpha/D,
        'Fo_heat_frozen_coefficient':alpha*t/R**2,'Fo_mass_frozen_coefficient':D*t/R**2,
        'tau_heat_radial_s':R**2/alpha,'tau_mass_radial_s':R**2/D,
        'tau_heat_radial_h':R**2/alpha/3600,'tau_mass_radial_h':R**2/D/3600,
        'tau_heat_surface_lumped_s':B*R/(2*25),'tau_mass_surface_lumped_s':R/(2*8e-7)}

def alpha_candidates(q,lo,hi):
    # alpha=(n0+n1*C)/(d0+d1*C+d2*C²), stationarity has quadratic numerator.
    if q=='Q1': return [lo,hi]
    a,b,p,c,k,v = (650,128,1450,2736,.21,.38) if q=='Q23' else (760,90,1850,2150,.12,.20)
    n0,n1=k,k+v;d0,d1,d2=a*p,a*(p+c)+b*p,b*(p+c)
    aa,bb,cc=-n1*d2,-2*n0*d2,n1*d0-n0*d1
    discriminant=bb*bb-4*aa*cc
    roots=[] if discriminant<0 else [(-bb-math.sqrt(discriminant))/(2*aa),(-bb+math.sqrt(discriminant))/(2*aa)]
    return [lo,hi]+[r for r in roots if lo<r<hi]

def main():
    started=stamp();tic=time.perf_counter()
    base=ROOT/'problem_files/CUMCM2026Problems/A题'
    pdf=base/'A题.pdf'; raw1=base/'附件/附件1.xlsx';raw2=base/'附件/附件2.xlsx'
    clean1=ROOT/'paper_output/data_cleaned/A_environment_observed.csv';clean2=ROOT/'paper_output/data_cleaned/A_radius_observed.csv'
    core=ROOT/'paper_output/code/modeling/drying_core.py'
    summaries={q:ROOT/f'paper_output/results/production/final_v6a/{q}/summary.json' for q in ('Q1','Q23','Q4')}
    inputs=[pdf,raw1,raw2,clean1,clean2,core,*summaries.values()]
    before=[record(p) for p in inputs]
    with pymupdf.open(pdf) as doc:
        pages=[p.get_text() for p in doc]
        (OUT/'raw_problem_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{p}' for i,p in enumerate(pages)),encoding='utf-8')
    env_s=read_raw(raw1);rad_s=read_raw(raw2)
    assert len(env_s)==len(rad_s)==1
    env=env_s[0]['rows'];rad=rad_s[0]['rows']
    assert env_s[0]['headers']==['时间','温度','水分浓度']
    assert rad_s[0]['headers']==['时间','半径']
    assert all(s['missing_data_cells']==s['nonnumeric_data_cells']==s['formula_cells']==0 for s in env_s+rad_s)
    c1=compare(env,clean1,lambda r:{'time_s':r[0],'time_h':r[0]/3600,'temperature_C':r[1],'temperature_K':r[1]+273.15,'air_moisture_kg_per_kg':r[2]},dict(time_s=0,time_h=1e-13,temperature_C=0,temperature_K=1e-12,air_moisture_kg_per_kg=0))
    c2=compare(rad,clean2,lambda r:{'time_s':r[0],'time_h':r[0]/3600,'radius_cm':r[1],'radius_m':r[1]/100},dict(time_s=0,time_h=1e-13,radius_cm=0,radius_m=1e-16))
    assert c1['pass'] and c2['pass']
    tail=[r for r in env if r[0]>=10800]
    env_stats={'count':len(env),'support_s':[env[0][0],env[-1][0]],'temperature_C_min_max':[min(r[1] for r in env),max(r[1] for r in env)],'air_moisture_min_max':[min(r[2] for r in env),max(r[2] for r in env)],'first':env[0],'last':env[-1],
        'last_hour_inclusive_samples':len(tail),'last_hour_temperature_mean_C':sum(r[1] for r in tail)/len(tail),'last_hour_air_mean':sum(r[2] for r in tail)/len(tail),
        'nominal_extension_right_minus_observed_left_at_4h':{'temperature_C':50-env[-1][1],'equivalent_Ceq':.05-env[-1][2]},
        'observations_after_4h':sum(r[0]>14400 for r in env),'physical_response_values_excluding_time':2*len(env)}
    rad_stats={'count':len(rad),'support_s':[rad[0][0],rad[-1][0]],'radius_cm_min_max':[min(r[1] for r in rad),max(r[1] for r in rad)],'first':rad[0],'last':rad[-1],
        'increasing_intervals':sum(rad[i+1][1]>rad[i][1] for i in range(len(rad)-1)),
        'equal_radius_intervals':sum(rad[i+1][1]==rad[i][1] for i in range(len(rad)-1)),
        'min_dR_dt_m_s':min((rad[i+1][1]-rad[i][1])/100/(rad[i+1][0]-rad[i][0]) for i in range(len(rad)-1))}
    # Static equivalence against the function AST, never execute the solver.
    tree=ast.parse(core.read_text(encoding='utf-8'))
    properties=next(n for c in tree.body if isinstance(c,ast.ClassDef) and c.name=='RadialModel' for n in c.body if isinstance(n,ast.FunctionDef) and n.name=='properties')
    branches=[n for n in ast.walk(properties) if isinstance(n,ast.If) and isinstance(n.test,ast.Compare) and ast.unparse(n.test.left)=='s.question']
    expected={'Q1':['rho = np.full_like(C, 820.)','cp = np.full_like(C, 2600.)','k = np.full_like(C, .36)','D = 7e-9 * np.exp(-.89 / positive_C)'],
        'Q23':['rho, cp, k = 650 + 128 * positive_C, 1450 + 2736 * wet, .21 + .38 * wet','D = 2.4e-3 * np.exp(-.45 / positive_C - 3850 / T)'],
        'Q4':['rho, cp, k = 760 + 90 * positive_C, 1850 + 2150 * wet, .12 + .20 * wet','D = 4.2e-4 * np.exp(-.30 / positive_C - 3850 / T)']}
    static=[]
    for q,branch in zip(('Q1','Q23','Q4'),branches):
        actual=[n for n in branch.body if isinstance(n,ast.Assign)]
        checks=[ast.dump(a,include_attributes=False)==ast.dump(ast.parse(e).body[0],include_attributes=False) for a,e in zip(actual,expected[q])]
        static.append({'question':q,'lineno':branch.lineno,'all_assignments_match':all(checks) and len(actual)==len(expected[q]),'actual_assignments':[ast.unparse(n) for n in actual]})
    assert len(static)==3 and all(r['all_assignments_match'] for r in static)
    init_assignment=next(n for n in tree.body if isinstance(n,ast.Assign) and ast.unparse(n.targets[0])=='(C0, T0, R0, LENGTH)')
    assert ast.literal_eval(init_assignment.value)==(2.55,301.15,.02,.25)
    sums={q:json.loads(p.read_text(encoding='utf-8')) for q,p in summaries.items()}
    C0,T0,R0,L=2.55,301.15,.02,.25;V0=math.pi*R0**2*L
    states=[];mass=[];cover={}
    for q,s in sums.items():
        t=1800 if q=='Q1' else s['completion']['reported_time_s']
        Rend=R0 if q!='Q4' else s['diagnostics']['radius_end_m']
        states.append(dict(label='initial',scope='initial empirical coefficient',**scalars(q,C0,T0,R0,t)))
        states.append(dict(label='threshold_reference',scope='C=.15 and T=50C illustrative state, not a uniform solved field',**scalars(q,.15,323.15,Rend,t)))
        if q!='Q1':
            Cs=s['diagnostics']['final_surface_C']; Ts=s['reproduction_samples']['T_K'][-1][-1]
            states.append(dict(label='production_end_surface_reference',scope='saved production state used only as scalar reference, not raw observation',**scalars(q,Cs,Ts,Rend,t)))
            cover[q]={'reported_total_h':t/3600,'observed_environment_h':4.,'extension_h':t/3600-4.,'fraction_environment_unobserved':(t-14400)/t,'nominal_extension_not_new_observations':True,
                'radius_observed_through_h':72.,'radius_extrapolation_required':q=='Q4' and t>rad[-1][0]}
        rho,cp,k,D=prop(q,C0,T0); Md=V0*rho/(1+C0); Mwet=V0*rho; Mw=Md*C0; deltaC=2.4; Lv=2.4e6
        mass.append({'question':q,'initial_rho_kg_m3':rho,'initial_cp_J_kgK':cp,'V0_m3':V0,'initial_wet_mass_reference_kg':Mwet,'initial_dry_mass_calibration_kg':Md,'initial_water_mass_kg':Mw,
            'removed_water_lower_bound_if_everywhere_C_le_0p15_kg':Md*deltaC,'Lv_audit_assumption_J_kg':Lv,'latent_energy_lower_reference_J':Md*deltaC*Lv,
            'constant_initial_capacity_sensible_reference_J':V0*rho*cp*22,'latent_to_constant_capacity_sensible_ratio':deltaC*Lv/((1+C0)*cp*22),
            'energy_not_measured_not_total_consumption_forecast':True})
    states.append(dict(label='3h_initial_frozen_coefficient_reference',scope='illustrative Fourier numbers for Q2 three-hour table, not a solved field',**scalars('Q23',C0,T0,R0,10800)))
    ranges=[];Tlo=T0;Thi=max(r[1] for r in env)+273.15;Clo,Chi=.05,C0
    for q in ('Q1','Q23','Q4'):
        cs=alpha_candidates(q,Clo,Chi)
        avals=[(prop(q,C,Tlo)[2]/(prop(q,C,Tlo)[0]*prop(q,C,Tlo)[1]),C) for C in cs]
        alo,ahi=min(avals),max(avals);dlo=prop(q,Clo,Tlo)[3];dhi=prop(q,Chi,Thi)[3]
        rlo=min(r[1] for r in rad)/100 if q=='Q4' else R0
        # Le decreases with C here; verify derivative sign across this range analytically sampled.
        # Report endpoint/sweep envelope, not rigorous joint extremum if no proof.
        sampled=[]
        for i in range(2001):
            C=Clo+(Chi-Clo)*i/2000
            for T in (Tlo,Thi):
                rr,cc,kk,dd=prop(q,C,T);sampled.append((kk/(rr*cc)/dd,C,T))
        ranges.append({'question':q,'range_definition':'independent reference box; not set of jointly observed/solved states','C_range':[Clo,Chi],'T_K_range':[Tlo,Thi],'R_m_range':[rlo,R0],
            'alpha_min_value_and_C':alo,'alpha_max_value_and_C':ahi,'D_min_max':[dlo,dhi],
            'Bi_heat_radius_box_min_max':[25*rlo/prop(q,Chi,Tlo)[2],25*R0/prop(q,Clo,Tlo)[2]],
            'Bi_mass_radius_box_min_max':[8e-7*rlo/dhi,8e-7*R0/dlo],
            'tau_heat_h_box_min_max':[rlo*rlo/ahi[0]/3600,R0*R0/alo[0]/3600],
            'tau_mass_h_box_min_max':[rlo*rlo/dhi/3600,R0*R0/dlo/3600],
            'effective_Le_sampled_min':min(sampled),'effective_Le_sampled_max':max(sampled),'Le_scalar_samples':len(sampled)})
    R72=rad[-1][1]/100;rhoD0=(760+90*C0)/(1+C0);Rcrit=R0*math.sqrt(rhoD0/760)
    first_below=next(r for r in rad if r[1]/100<Rcrit)
    adjacent=next((rad[i],rad[i+1]) for i in range(len(rad)-1) if rad[i][1]/100>=Rcrit>rad[i+1][1]/100)
    a,b=adjacent;tcrit=a[0]+(Rcrit-a[1]/100)/(b[1]/100-a[1]/100)*(b[0]-a[0])
    geometry={'L0_m':L,'R0_m':R0,'V0_m3':V0,'side_area_m2':2*math.pi*R0*L,'two_end_area_m2':2*math.pi*R0**2,'end_to_side_area_ratio':R0/L,
        'lumped_side_characteristic_length_m':R0/2,'lumped_full_cylinder_V_over_A_m':R0*L/(2*(L+R0)),
        'R72_m':R72,'R72_over_R0':R72/R0,'volume72_over_volume0_fixed_L':(R72/R0)**2,'diffusion_prefactor72_over_initial':(R0/R72)**2,'boundary_beta_over_R_prefactor72_over_initial':R0/R72,
        'production_Q4_Rend_m':sums['Q4']['diagnostics']['radius_end_m'],'production_Q4_diffusion_prefactor_over_initial':(R0/sums['Q4']['diagnostics']['radius_end_m'])**2,
        'production_Q4_conserved_dry_density_kg_m3':rhoD0*(R0/sums['Q4']['diagnostics']['radius_end_m'])**2,
        'production_Q4_max_drymass_fraction_if_literal_rho':760*(sums['Q4']['diagnostics']['radius_end_m']/R0)**2/rhoD0,
        'literal_wet_rho_d0_kg_m3':rhoD0,'literal_wet_rho_d_upper_bound_for_C_ge_0':760.,'R_necessary_min_m_for_fixed_L_conserved_Md':Rcrit,
        'first_observed_below_necessary_min_t_s':first_below[0],'linearly_interpolated_crossing_h':tcrit/3600,'interpolation_is_assumption':True,
        'max_remaining_drymass_fraction_at_72h':760*(R72/R0)**2/rhoD0,'necessary_L72_over_L0_min_for_literal_rho':rhoD0/(760*(R72/R0)**2)}
    # Units/meaning were visually checked against raw PDF renders; coefficients have empirical status.
    formulas=[
        {'id':'E01','question':'Q1','source_pdf_page':3,'expression':'D=7e-9 exp(-0.89/C)','dependent':'D [m²/s]','independent':'C [kgwater/kgdry]'},
        {'id':'E02','question':'Q23','source_pdf_page':4,'expression':'rho=650+128 C','dependent':'rho [kg/m³]','independent':'C [kgwater/kgdry]'},
        {'id':'E03','question':'Q23','source_pdf_page':4,'expression':'cp=1450+2736 C/(1+C)','dependent':'cp [J/(kg K)]','independent':'C [kgwater/kgdry]'},
        {'id':'E04','question':'Q23','source_pdf_page':4,'expression':'k=0.21+0.38 C/(1+C)','dependent':'k [W/(m K)]','independent':'C [kgwater/kgdry]'},
        {'id':'E05','question':'Q23','source_pdf_page':4,'expression':'D=2.4e-3 exp(-0.45/C) exp(-3850/T)','dependent':'D [m²/s]','independent':'C [kgwater/kgdry], T [K]'},
        {'id':'E06','question':'Q4','source_pdf_page':4,'expression':'rho=760+90 C','dependent':'rho [kg/m³]','independent':'C [kgwater/kgdry]'},
        {'id':'E07','question':'Q4','source_pdf_page':4,'expression':'cp=1850+2150 C/(1+C)','dependent':'cp [J/(kg K)]','independent':'C [kgwater/kgdry]'},
        {'id':'E08','question':'Q4','source_pdf_page':4,'expression':'k=0.12+0.20 C/(1+C)','dependent':'k [W/(m K)]','independent':'C [kgwater/kgdry]'},
        {'id':'E09','question':'Q4','source_pdf_page':4,'expression':'D=4.2e-4 exp(-0.30/C) exp(-3850/T)','dependent':'D [m²/s]','independent':'C [kgwater/kgdry], T [K]'}]
    after=[record(p) for p in inputs];assert before==after
    result={'status':'INDEPENDENT_RAW_DATA_AND_SCALAR_RECHECK_COMPLETE','started_utc':started,'ended_utc':stamp(),'elapsed_s':time.perf_counter()-tic,
        'runtime':{'python':sys.version,'executable':sys.executable,'platform':platform.platform(),'openpyxl':openpyxl.__version__,'pymupdf':pymupdf.__version__},
        'scope':'fresh raw XLSX/PDF read, cleaned scalar comparison, formula static check and independent scalar calculations; no PDE run, no GUI, no Git',
        'source_records_before':before,'source_records_after':after,'all_inputs_hash_unchanged':before==after,'helper':record(pathlib.Path(__file__)),
        'raw_sheets':{'environment':env_s,'radius':rad_s},'cleaned_comparison':{'environment':c1,'radius':c2},'environment':env_stats,'radius':rad_stats,
        'empirical_formulas':formulas,'static_core_formula_checks':static,'initial_constants_static_match':True,
        'constant_units':{'C0':'2.55 kg water/kg dry material','T0':'28 C = 301.15 K','R0':'2 cm = 0.02 m','L0':'25 cm = 0.25 m','h':'25 W/(m² K)','beta':'8e-7 m/s'},
        'scalar_states':states,'independent_reference_boxes':ranges,'mass_and_energy_references':mass,'geometry_and_literal_density_risk':geometry,'environment_prediction_coverage':cover,
        'air_basis_risk':{'source_wording':'烘房 水分浓度 kg/kg; does not specify dry-air denominator or sorption law','direct_Ceq_use_assumption_starts_at_s':0,'initial_equivalent_driving_difference':C0-env[0][2],'no_isotherm_or_gas_film_partition_measurements':True},
        'interpretation_limits':['nine empirical fits are supplied, not re-estimated from 241/145 data points','heat/mass Biot use radius; V/A side convention gives half','effective alpha/D is a solid thermal-to-moisture ratio, not measured gas Lewis number','all Fourier/diffusion-time references freeze local coefficients and are not nonlinear drying-time predictions','reference boxes allow independent C,T,R combinations and are not observed trajectories','latent energy is only an audit assumption with initial rho dry-mass calibration; sensible reference is constant initial B times 22K, not actual energy consumption','allwhere C<.15 yields strictly more than Md*2.4 water loss; displayed lower reference uses non-strict <= for conservative bound','there are zero observed environment samples after4h; numerical outputs and interpolated values are not additional observations']}
    (OUT/'data_scale_recheck.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'status':result['status'],'cleaned_numeric_checks':c1['numeric_values_compared']+c2['numeric_values_compared'],'raw1_data_rows':len(env),'raw2_data_rows':len(rad),'mass_and_energy_references':mass,'scalar_states':states,'geometry':geometry,'coverage':cover},ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
