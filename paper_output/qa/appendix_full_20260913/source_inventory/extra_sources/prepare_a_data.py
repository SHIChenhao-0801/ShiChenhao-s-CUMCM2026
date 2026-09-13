from __future__ import annotations

import csv
import hashlib
import json
import os
import platform
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
if Path.cwd().resolve() != ROOT:
    raise RuntimeError(f"Run with workdir {ROOT}")
os.environ["MPLCONFIGDIR"] = str(ROOT / "tmp/cache/matplotlib")

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as font_manager
import numpy as np
import openpyxl

OUT = ROOT / "paper_output"
DATA = OUT / "data_cleaned"
PLAN = OUT / "plan"
FIG = OUT / "figures"
MANIFEST = OUT / "input_manifest.json"
SCRIPT = Path(__file__).resolve()
EXPECTED = {
    "environment": "problem_files/CUMCM2026Problems/A题/附件/附件1.xlsx",
    "radius": "problem_files/CUMCM2026Problems/A题/附件/附件2.xlsx",
}


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path: Path, header: list[str], data: np.ndarray) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows([[format(float(v), ".15g") for v in row] for row in data])


def base_record() -> dict:
    return {"schema_version": "1.0", "generated_by": rel(SCRIPT),
            "generated_at": now(), "source_code_sha256": sha(SCRIPT)}


def read_and_diagnose() -> tuple[dict, dict, dict]:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8-sig"))
    entries = {entry["path"]: entry for entry in manifest["entries"]}
    report = {**base_record(), "status": "PASS", "input_manifest_used": True,
              "input_manifest": rel(MANIFEST), "input_manifest_sha256": sha(MANIFEST),
              "active_problem": "A", "input_dirs": ["problem_files/CUMCM2026Problems/A题/附件"],
              "data_files": [], "skipped_files": [], "pdf_diagnostics": [],
              "warnings": [], "errors": []}
    for path, entry in entries.items():
        if path not in EXPECTED.values():
            report["skipped_files"].append({"path": path, "role": entry.get("role"),
                "reason": "outside_active_A_raw_inputs"})
    raw_arrays, diagnostics = {}, {}
    for key, relative in EXPECTED.items():
        entry = entries.get(relative, {})
        if entry.get("role") != "raw_data" or not entry.get("usable_for_modeling"):
            raise RuntimeError(f"Input is not authorized raw data in manifest: {relative}")
        path = ROOT / relative
        before = sha(path)
        if entry.get("sha256") != before:
            raise RuntimeError(f"Manifest hash differs from raw input: {relative}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if wb.sheetnames != ["Sheet1"]:
                raise ValueError(f"Unexpected worksheet schema: {relative}")
            values = list(wb["Sheet1"].iter_rows(values_only=True))
        finally:
            wb.close()
        header, body = list(values[0]), values[1:]
        expected_header = ["时间", "温度", "水分浓度"] if key == "environment" else ["时间", "半径"]
        if header != expected_header:
            raise ValueError(f"Unexpected header: {relative}: {header}")
        if any(len(row) != len(header) for row in body):
            raise ValueError(f"Unequal row length: {relative}")
        if any(not isinstance(v, (int, float)) or isinstance(v, bool) for row in body for v in row):
            raise ValueError(f"Missing or nonnumeric input found, review required: {relative}")
        data = np.asarray(body, dtype=float)
        if not np.isfinite(data).all():
            raise ValueError(f"Nonfinite values found: {relative}")
        step = 60 if key == "environment" else 1800
        end = 14400 if key == "environment" else 259200
        if data[0, 0] != 0 or data[-1, 0] != end or not np.all(np.diff(data[:, 0]) == step):
            raise ValueError(f"Unexpected time coverage or duplicate time: {relative}")
        if key == "radius" and (np.any(np.diff(data[:, 1]) > 0) or np.any(data[:, 1] <= 0)):
            raise ValueError("Radius must be positive and nonincreasing; review raw data")
        diag = {"path": relative, "kind": "spreadsheet", "ext": ".xlsx", "readable": True,
                "sha256": before, "unchanged_after_read": sha(path) == before,
                "sheets": [{"name": "Sheet1", "rows": len(values), "cols": len(header), "sample_cols": header}],
                "numeric_record_count": len(data), "missing_numeric_cells": 0,
                "nonfinite_values": 0, "duplicate_times": 0, "time_start_s": 0,
                "time_end_s": end, "time_step_s": step, "warnings": [], "errors": [],
                "minima": data.min(axis=0).tolist(), "maxima": data.max(axis=0).tolist()}
        report["data_files"].append(diag)
        raw_arrays[key], diagnostics[key] = data, diag
    report["summary"] = {"file_count_scanned": 2, "data_file_count": 2,
                         "readable_data_file_count": 2, "pdf_file_count": 0}
    report["warnings"] = ["Air kg/kg is not automatically solid dry-basis equilibrium moisture.",
                          "Room observations end at 4 h; radius observations end at 72 h.",
                          "No internal temperature/moisture validation measurements are supplied."]
    save_json(DATA / "load_report.json", report)
    return raw_arrays, diagnostics, report


def make_figures(environment: np.ndarray, radius: np.ndarray) -> list[dict]:
    font_file = Path("C:/Windows/Fonts/msyh.ttc")
    if font_file.exists():
        font_manager.fontManager.addfont(str(font_file))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font_file)).get_name()
    plt.rcParams.update({"font.size": 10, "axes.titlesize": 11, "axes.labelsize": 10,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.unicode_minus": False, "savefig.dpi": 220,
                         "svg.fonttype": "path", "axes.grid": True, "grid.alpha": 0.18})
    records = []
    configs = [
        {"figure_id": "fig_A_environment", "title": "烘房温度与水分指标原始观测",
         "question_id": "Q1", "question_ids": ["Q1", "Q2", "Q3", "Q4"],
         "data_source": "paper_output/data_cleaned/A_environment_observed.csv",
         "candidate_x": "time_h", "candidate_y": ["temperature_C", "air_moisture_kg_per_kg"],
         "purpose": "展示0–4小时实测边界及末期平台，标明长期延拓的观测边界",
         "paper_usage": "数据审计与环境边界条件", "caption": "241条原始观测；连线仅连接采样点，不表示4小时以后实测。空气kg/kg指标未直接等同药材平衡含水率。"},
        {"figure_id": "fig_A_radius", "title": "药材半径收缩原始观测",
         "question_id": "Q4", "question_ids": ["Q4"],
         "data_source": "paper_output/data_cleaned/A_radius_observed.csv",
         "candidate_x": "time_h", "candidate_y": ["radius_cm"],
         "purpose": "展示0–72小时给定半径和48–72小时尾段，限定收缩几何的数据支持范围",
         "paper_usage": "收缩数据与移动边界", "caption": "145条原始观测；半径保持非增。右图放大48–72小时，未对72小时以后外推。"},
    ]
    for item in configs:
        fig, axes = plt.subplots(1, 2, figsize=(10.8, 3.7), layout="constrained")
        if item["figure_id"] == "fig_A_environment":
            for ax, y, color, ylabel, title in [
                (axes[0], environment[:, 2], "#1b6b93", "温度 / °C", "烘房温度"),
                (axes[1], environment[:, 4], "#b86625", "空气水分指标 / (kg/kg)", "烘房水分指标"),
            ]:
                ax.plot(environment[:, 1], y, color=color, lw=1.35, marker=".", ms=2)
                ax.set(xlabel="时间 / h", ylabel=ylabel, title=title, xlim=(0, 4))
                ax.axvspan(3, 4, color=color, alpha=0.065)
        else:
            axes[0].plot(radius[:, 1], radius[:, 2], color="#336950", lw=1.35, marker=".", ms=2)
            axes[0].set(xlabel="时间 / h", ylabel="半径 / cm", title="完整观测区间", xlim=(0, 72))
            tail = radius[:, 1] >= 48
            axes[1].plot(radius[tail, 1], radius[tail, 2], color="#336950", lw=1.35, marker="o", ms=2.2)
            axes[1].set(xlabel="时间 / h", ylabel="半径 / cm", title="48–72 h 尾段放大", xlim=(48, 72))
            axes[1].ticklabel_format(axis="y", style="plain", useOffset=False)
        paths = {}
        for ext in ["png", "svg", "pdf"]:
            path = FIG / f"{item['figure_id']}.{ext}"
            fig.savefig(path)
            paths[ext] = {"path": rel(path), "sha256": sha(path), "bytes": path.stat().st_size}
        plt.close(fig)
        records.append({**item, "path": paths["png"]["path"], "output_path": paths["png"]["path"],
                        "chart_type": "observed_time_series", "artifacts": paths,
                        "status": "generated_from_observations", "ok": True, "placeholder": False,
                        "exists": True, "planned": False, "used_in": item["paper_usage"],
                        "source_sha256": sha(ROOT / item["data_source"]),
                        "visual_review": "pending", "human_review": "pending"})
    return records


def main() -> None:
    for directory in [DATA, PLAN, FIG]:
        directory.mkdir(parents=True, exist_ok=True)
    raw_arrays, diagnostics, report = read_and_diagnose()
    air, rad = raw_arrays["environment"], raw_arrays["radius"]
    environment = np.column_stack([air[:, 0], air[:, 0]/3600, air[:, 1], air[:, 1]+273.15, air[:, 2]])
    radius = np.column_stack([rad[:, 0], rad[:, 0]/3600, rad[:, 1], rad[:, 1]*0.01])
    env_header = ["time_s", "time_h", "temperature_C", "temperature_K", "air_moisture_kg_per_kg"]
    rad_header = ["time_s", "time_h", "radius_cm", "radius_m"]
    write_csv(DATA / "A_environment_observed.csv", env_header, environment)
    write_csv(DATA / "A_radius_observed.csv", rad_header, radius)

    env_back = np.loadtxt(DATA / "A_environment_observed.csv", delimiter=",", skiprows=1, encoding="utf-8-sig")
    rad_back = np.loadtxt(DATA / "A_radius_observed.csv", delimiter=",", skiprows=1, encoding="utf-8-sig")
    if not np.allclose(env_back, environment, rtol=0, atol=1e-10):
        raise ValueError("Environment CSV round-trip failed")
    if not np.allclose(rad_back, radius, rtol=0, atol=1e-10):
        raise ValueError("Radius CSV round-trip failed")
    data_files = []
    for key, header, name in [("environment", env_header, "A_environment_observed.csv"),
                              ("radius", rad_header, "A_radius_observed.csv")]:
        data_files.append({**diagnostics[key], "role": "raw_data", "type": "xlsx",
            "columns": header, "numeric_columns": header, "categorical_columns": [],
            "cleaned_output": rel(DATA / name), "output_sha256": sha(DATA / name),
            "cleaning_tasks": ["numeric/schema/finite/duplicate/time-step checks", "unit conversion only"],
            "operations": {"rows_removed": 0, "values_filled": 0, "values_changed": 0,
                           "smoothing": False, "extrapolation": False}})
    data_plan = {**base_record(), "source_contracts": [rel(MANIFEST),
        "paper_output/step1/problem_analysis.json", "paper_output/plan/model_route.json"],
        "data_files": data_files, "question_links": [
        {"question_id": q, "required_fields": env_header + (rad_header if q == "Q4" else []),
         "input_sources": [EXPECTED["environment"]] + ([EXPECTED["radius"]] if q == "Q4" else []),
         "expected_outputs": ["temperature/moisture fields and validation" if q in ["Q1", "Q2"] else "moisture field, maximum criterion and drying time"],
         "status": "model_outputs_pending"} for q in ["Q1", "Q2", "Q3", "Q4"]],
        "assumption_boundaries": {"continuous_interpolation": "between observed samples",
            "air_to_solid_boundary_mapping": "t=0; requires explicit closure",
            "environment_extrapolation_start_s": 14400,
            "radius_extrapolation_start_s": 259200,
            "material_velocity_and_axial_length": "t=0; unobserved assumptions"},
        "limitations": report["warnings"],
        "validation": {"csv_round_trip": "PASS", "originals_unchanged": True,
                       "visual_studio_gui": "pending", "human_review": "pending"}}
    save_json(PLAN / "data_plan.json", data_plan)
    figures = make_figures(environment, radius)
    save_json(PLAN / "visualization_plan.json", {**base_record(), "source_contracts": [
        "paper_output/plan/model_route.json", "paper_output/plan/data_plan.json"],
        "figures": figures, "scope": "observed inputs only; model figures are a later stage"})
    save_json(OUT / "figure_index.json", {**base_record(), "figures": figures})
    run_record = {**base_record(), "status": "PASS", "operation": "S3 observed input processing",
        "runtime": {"executable": sys.executable, "python": sys.version, "platform": platform.platform(),
                    "numpy": np.__version__, "openpyxl": openpyxl.__version__, "matplotlib": matplotlib.__version__},
        "input_manifest_sha256": sha(MANIFEST), "record_counts": {"environment": len(air), "radius": len(rad)},
        "observations_changed": 0, "csv_round_trip_max_absolute_error": {
            "environment": float(np.max(np.abs(env_back-environment))),
            "radius": float(np.max(np.abs(rad_back-radius)))},
        "figures_generated": [item["figure_id"] for item in figures],
        "model_solve": False, "visual_studio_gui_reproduction": False, "user_human_review": False}
    save_json(DATA / "A_data_run_record.json", run_record)
    print(json.dumps(run_record, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
