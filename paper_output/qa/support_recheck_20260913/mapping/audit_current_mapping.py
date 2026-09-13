from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
from docx import Document
import csv
import hashlib
import io
import json
import re
import sys
import datetime
import pymupdf

ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')
QA = ROOT / 'paper_output/qa/support_recheck_20260913/mapping'
QA.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT / '药材热湿耦合模型与干燥时间计算_附录修订版.docx'
SUPPORT = ROOT / '支撑材料'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'

def sha(data):
    return hashlib.sha256(data).hexdigest()

def txt(element):
    out = []
    for child in element.iter():
        if child.tag == W + 't': out.append(child.text or '')
        elif child.tag == W + 'tab': out.append('\t')
        elif child.tag == W + 'br': out.append('\n')
    return ''.join(out)

raw = SOURCE.read_bytes()
snapshot = QA / 'current_paper_readonly_snapshot.docx'
snapshot.write_bytes(raw)
doc = Document(io.BytesIO(raw))
paragraphs = [p.text for p in doc.paragraphs]
with ZipFile(io.BytesIO(raw)) as z:
    body = ET.fromstring(z.read('word/document.xml')).find('w:body', NS)
    media = [{'part': n, 'bytes': len(z.read(n)), 'sha256': sha(z.read(n))}
             for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')

sections = []
for i, element in enumerate(body):
    sect = element if element.tag == W + 'sectPr' else element.find('w:pPr/w:sectPr', NS)
    if sect is None: continue
    sections.append({'body_element_index': i, 'section_properties_xml': ET.tostring(sect, encoding='unicode')})

table_audits = []
tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
for p in sorted((SUPPORT / '04_结果表格/Referrence Table Files').glob('q*_paper_*.csv')):
    rows = list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig'))))
    exact = [i for i, t in enumerate(tables) if t == rows]
    data_matches = [i for i, t in enumerate(tables) if t[1:] == rows[1:]]
    table_audits.append({'csv': str(p.relative_to(ROOT)), 'csv_sha256': sha(p.read_bytes()),
                         'exact_table_indices': exact, 'data_only_table_indices': data_matches,
                         'data_cells': sum(len(r) for r in rows[1:]),
                         'headers_csv': rows[0], 'headers_docx': tables[data_matches[0]][0] if data_matches else None})

build = json.loads((ROOT / 'paper_output/qa/appendix_revision_20260913/build_audit.json').read_text(encoding='utf-8-sig'))
headings = [(i, txt(el)) for i, el in enumerate(body) if re.match(r'^9\.2\.\d+\s', txt(el))]
source_checks = []
for old, (i, heading) in zip(build['sources'], headings, strict=True):
    p = ROOT / old['relative_path']
    lines = p.read_text(encoding='utf-8-sig').splitlines()
    extracted = [txt(el) for el in body[i + 1:i + 1 + len(lines)]]
    different = [{'line': j + 1, 'source': a, 'docx': b}
                 for j, (a, b) in enumerate(zip(lines, extracted)) if a != b]
    source_checks.append({'number': old['number'], 'source': str(p.relative_to(ROOT)),
                          'heading': heading, 'source_lines': len(lines),
                          'exact_lines_equal': lines == extracted,
                          'current_source_sha256': sha(p.read_bytes()),
                          'matches_old_manifest_hash': sha(p.read_bytes()) == old['sha256'],
                          'first_differences': different[:5]})

pdfs = []
for p in sorted((SUPPORT / '02_参考文献与网络资料').rglob('*.pdf')):
    d = pymupdf.open(p)
    pdfs.append({'file': str(p.relative_to(ROOT)), 'sha256': sha(p.read_bytes()), 'pages': len(d),
                 'first_page_text': d[0].get_text(), 'last_page_text': d[-1].get_text()})

figroot = ROOT / 'paper_output/figures/review_20260913'
figure_assets = []
for p in sorted(figroot.glob('图*')):
    if p.suffix.lower() not in ('.png', '.svg', '.pdf'): continue
    digest = sha(p.read_bytes())
    figure_assets.append({'file': str(p.relative_to(ROOT)), 'sha256': digest,
                          'docx_media_exact_matches': [m['part'] for m in media if m['sha256'] == digest],
                          'support_file_same_hash': [str(s.relative_to(ROOT)) for s in SUPPORT.rglob('*')
                                                    if s.is_file() and s.stat().st_size == p.stat().st_size and sha(s.read_bytes()) == digest]})
figure_csvs = []
for p in sorted((figroot / 'input_data/CSV').rglob('*.csv')):
    rows = list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig'))))
    figure_csvs.append({'file': str(p.relative_to(ROOT)), 'data_rows': len(rows)-1,
                        'sha256': sha(p.read_bytes())})

frozen_checks = []
for p in sorted((SUPPORT / '04_结果表格').glob('result*.xlsx')):
    frozen = ROOT / 'paper_output/results/production/final_v6a/outputs' / p.name
    frozen_checks.append({'file': str(p.relative_to(ROOT)), 'bytes': p.stat().st_size,
                           'sha256': sha(p.read_bytes()), 'frozen_exists': frozen.exists(),
                           'frozen_sha256': sha(frozen.read_bytes()) if frozen.exists() else None,
                           'equals_frozen': frozen.exists() and p.read_bytes() == frozen.read_bytes()})

allfiles = sorted(p for p in SUPPORT.rglob('*') if p.is_file())
ai_files = [str(p.relative_to(ROOT)) for p in allfiles if re.search(r'AI.*(详情|报告|使用)', p.name, re.I)]
report = {
    'audit_time_local': datetime.datetime.now().isoformat(timespec='seconds'),
    'scope': 'Current saved DOCX and support files only; no document edits, no solver run, no layout approval, no independent AI file restoration.',
    'paper': str(SOURCE), 'snapshot': str(snapshot), 'paper_sha256': sha(raw), 'paper_bytes': len(raw),
    'paper_last_write_local': datetime.datetime.fromtimestamp(SOURCE.stat().st_mtime).isoformat(timespec='seconds'),
    'unchanged_during_audit': SOURCE.read_bytes() == raw,
    'prior_appendix_output_sha256': build['output_sha256'],
    'equals_prior_appendix_output': sha(raw) == build['output_sha256'],
    'seven_result_tables': table_audits,
    'all_297_result_cells_equal': len(table_audits) == 7 and sum(r['data_cells'] for r in table_audits) == 297 and all(len(r['data_only_table_indices']) == 1 for r in table_audits),
    'four_workbooks': frozen_checks,
    'appendix_sources': source_checks, 'source_total_lines': sum(s['source_lines'] for s in source_checks),
    'all_36_source_texts_equal': len(source_checks) == 36 and all(s['exact_lines_equal'] for s in source_checks),
    'docx_media': media, 'figure_assets_in_workspace': figure_assets,
    'figure_csvs_in_workspace': figure_csvs, 'figure_csv_data_rows': sum(r['data_rows'] for r in figure_csvs),
    'support_image_files': [str(p.relative_to(ROOT)) for p in allfiles if p.suffix.lower() in ('.png', '.svg', '.jpg', '.jpeg')],
    'support_r_scripts': [str(p.relative_to(ROOT)) for p in allfiles if p.suffix.lower() == '.r'],
    'pdf_inventory': pdfs,
    'reference_paragraphs': [p for p in paragraphs if re.match(r'^\[[1-7]\]', p)],
    'support_html': [str(p.relative_to(ROOT)) for p in allfiles if p.suffix.lower() == '.html'],
    'ai_files_in_support': ai_files,
    'first_fifteen_paragraphs': paragraphs[:15],
    'section_properties': sections,
    'support_catalog_evidence_before_source': [{'body_element_index': i, 'text': txt(el)}
        for i, el in enumerate(body[:headings[0][0]])
        if re.search(r'支撑|文件列表|文件清单|01_|02_|04_|表9|表10|表11|续表10', txt(el), re.I)],
    'final_ai_report_paragraphs': paragraphs[next(i for i,p in enumerate(paragraphs) if p == '9.3 AI 使用报告'):],
    'appendix_cross_references_in_main_text': [p for p in paragraphs[:226] if re.search(r'附录[A-D]', p)],
}
(QA / 'current_mapping_audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
(QA / 'current_paper_body_text.txt').write_text('\n'.join(paragraphs[:470]), encoding='utf-8')
print(json.dumps({k: report[k] for k in ['paper_sha256','paper_bytes','unchanged_during_audit','equals_prior_appendix_output','all_297_result_cells_equal','all_36_source_texts_equal','source_total_lines','figure_csv_data_rows','ai_files_in_support']},ensure_ascii=False,indent=2))
print('WORKBOOKS',json.dumps(frozen_checks,ensure_ascii=False))
print('TABLES',json.dumps(table_audits,ensure_ascii=False))
print('MEDIA_EXACT',json.dumps([f for f in figure_assets if f['docx_media_exact_matches']],ensure_ascii=False))
