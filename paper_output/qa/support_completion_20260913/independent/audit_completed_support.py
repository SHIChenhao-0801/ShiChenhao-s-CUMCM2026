from pathlib import Path
from collections import Counter
import ast
import csv
import datetime
import hashlib
import io
import json
import re
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path.cwd()
assert ROOT.as_posix() == 'D:/Document/数学建模/2026CUMCM'
sys.stdout.reconfigure(encoding='utf-8')
QA = ROOT / 'paper_output/qa/support_completion_20260913/independent'
QA.mkdir(parents=True, exist_ok=True)
SUPPORT = ROOT / '支撑材料'

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def read_json(p):
    return json.loads(p.read_text(encoding='utf-8-sig'))

prior = read_json(ROOT / 'paper_output/qa/support_recheck_20260913/directory_final_audit.json')
baseline = {p['path'].replace('\\', '/'): p for p in prior['files']}
now = {p.relative_to(SUPPORT).as_posix(): {'sha256': digest(p), 'bytes': p.stat().st_size}
       for p in SUPPORT.rglob('*') if p.is_file()}
changed = [p for p, old in baseline.items() if p in now and old['sha256'] != now[p]['sha256']]
missing = sorted(set(baseline) - set(now))
added = sorted(set(now) - set(baseline))
core30 = [p for p in baseline if p.startswith('03_程序代码/')]
frozen = []
for i in range(1, 5):
    p = SUPPORT / f'04_结果表格/result{i}.xlsx'
    f = ROOT / f'paper_output/results/production/final_v6a/outputs/result{i}.xlsx'
    template = SUPPORT / f'01_赛题与原始数据/附件/附件3/result{i}.xlsx'
    frozen.append({'file': str(p.relative_to(SUPPORT)), 'sha256': digest(p),
                   'frozen_equal': digest(p) == digest(f), 'distinct_from_blank_template': digest(p) != digest(template)})

old05sources = [p for p in baseline if p.startswith('05_数值检验与实验/') and Path(p).suffix.lower() in {'.py', '.m', '.mjs'}]
method_path = '05_数值检验与实验/Python检验源码/paper_output/code/verification/method_comparison.py'
before_method = ROOT / 'paper_output/qa/support_completion_20260913/verification/method_comparison.before.py'
oldtree = ast.parse(before_method.read_text(encoding='utf-8-sig'))
newtree = ast.parse((SUPPORT / method_path).read_text(encoding='utf-8-sig'))
def exclude_root(tree):
    tree.body = [node for node in tree.body if not (isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'ROOT' for t in node.targets))]
    return ast.dump(tree, include_attributes=False)
method_only_root = exclude_root(oldtree) == exclude_root(newtree)
requirements = (SUPPORT / '05_数值检验与实验/Python检验源码/requirements.txt').read_text(encoding='utf-8-sig').splitlines()
data05 = []
for n in ['A_environment_observed.csv', 'A_radius_observed.csv']:
    p = SUPPORT / '05_数值检验与实验/Python检验源码/paper_output/data_cleaned' / n
    original = SUPPORT / '03_程序代码/inputs/cleaned' / n
    data05.append({'name': n, 'sha256': digest(p), 'equals_production_input': digest(p) == digest(original),
                   'rows': len(list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig')))))-1})

fig = SUPPORT / '06_绘图程序与数据'
figsource = ROOT / 'paper_output/figures/review_20260913'
csv06 = []
for p in sorted((fig / 'CSV').rglob('*.csv')):
    src = figsource / 'input_data/CSV' / p.relative_to(fig / 'CSV')
    csv06.append({'file': p.relative_to(fig).as_posix(), 'sha256': digest(p), 'source_equal': digest(p) == digest(src),
                  'rows': len(list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig')))))-1})
png06 = [{'file': p.relative_to(fig).as_posix(), 'sha256': digest(p), 'source_equal': digest(p) == digest(figsource / p.name)}
         for p in sorted((fig / 'PNG').glob('*.png'))]
figfiles = [p for p in fig.rglob('*') if p.is_file()]
figforbidden = [p.relative_to(fig).as_posix() for p in figfiles if p.suffix.lower() in {'.json', '.jsonl', '.md', '.markdown', '.pdf', '.svg', '.zip', '.rar'}]
figure_runtime = ROOT / 'paper_output/qa/support_completion_20260913/figures'
figure_run_rows = list(csv.DictReader(io.StringIO((figure_runtime / 'run_status.csv').read_text(encoding='utf-8-sig'))))
figure_r = {'sha256': digest(fig / 'draw_figures.R'),
            'equals_executed_isolated_R': digest(fig / 'draw_figures.R') == digest(figure_runtime / 'isolated_support/draw_figures.R'),
            'actual_process_exit_zero': len(figure_run_rows) == 1 and figure_run_rows[0]['exit_code'] == '0',
            'manual_source_review': 'R reads --file from commandArgs to resolve its own directory, uses adjacent CSV, and emit writes only PNG. Data assignments and plot functions remain in the delivered source. Independent source parser evidence records 32 unchanged assignments, unchanged data checks and six emit calls; the executed isolated copy matches current delivery.'}

ai_path = SUPPORT / 'AI工具使用详情.docx'
ai = Document(ai_path)
ai_text = '\n'.join([p.text for p in ai.paragraphs] + ['\n'.join(c.text for row in t.rows for c in row.cells) for t in ai.tables])
paper_path = ROOT / '药材热湿耦合模型与干燥时间计算_附录修订版.docx'
paper_hash = digest(paper_path)
ai_report = {'docx_sha256': digest(ai_path), 'table_count': len(ai.tables), 'contains_requested_tools': all(t in ai_text for t in ['Deepseek Harness', 'Deepseek v4.1 flash', 'Codex', 'GPT6-Astra']),
             'has_pending_fields': all(t in ai_text for t in ['待填写', '待本队确认', '待粘贴真实提示词', '待粘贴真实 AI 回复']),
             'no_checked_confirmation_symbol': not any(t in ai_text for t in ['☑', '✅', '✔']),
             'manual_semantic_review': 'Known tool uses remain within root manuscript section 9.3. Prompt and answer records are explicitly unfilled. Human confirmation is conditional and unchecked; no completed sign-off or invented dated interaction was found.',
             'remaining': 'Team must complete exact version/date/interaction/acceptance/review fields from actual evidence before submitting; editable DOCX is not final PDF.'}
(QA / 'ai_details_extracted_text.txt').write_text(ai_text, encoding='utf-8')

run_evidence = []
for p in sorted((ROOT / 'paper_output/qa/support_completion_20260913/verification').glob('isolated_*/run_result.json')):
    r = read_json(p)
    input_changed = [x['path'] for x in r['input_files'] if digest(SUPPORT / '05_数值检验与实验/Python检验源码' / x['path']) != x['sha256']]
    run_evidence.append({'record': str(p.relative_to(ROOT)), 'check': r['check'], 'returncode': r['returncode'],
                          'source_input_unchanged': r['source_input_unchanged'], 'wall_seconds': r['wall_seconds'],
                          'current_delivery_diff_from_executed': input_changed})

inventory_path = ROOT / '支撑材料文件列表_附录用.docx'
inventory = {'exists': inventory_path.exists()}
if inventory_path.exists():
    doc = Document(inventory_path)
    group = None
    entries = []
    summary_rows = []
    for el in doc.element.body:
        if el.tag.endswith('}p'):
            text = Paragraph(el, doc).text
            m = re.match(r'^表(\d+)\s+(.+)$', text)
            if m: group = m[2]
        elif el.tag.endswith('}tbl'):
            t = Table(el, doc)
            rows = [[c.text for c in r.cells] for r in t.rows]
            if rows[0] == ['目录', '文件数']:
                summary_rows = rows[1:]
            elif rows[0] == ['序号', '文件名', '用途']:
                for number, name, purpose in rows[1:]:
                    path = name if group == '根目录' else group + '/' + name
                    entries.append({'number': int(number), 'path': path, 'purpose': purpose})
            else: raise AssertionError('Unknown inventory table header: ' + str(rows[0]))
    paths = [e['path'] for e in entries]
    counts = Counter(paths)
    categories = Counter(p.split('/')[0] if '/' in p else '根目录' for p in now)
    parsed_summary = {name: int(count) for name, count in summary_rows}
    template_purposes = []
    for e in entries:
        if re.fullmatch(r'result[1-4]\.xlsx', Path(e['path']).name):
            template_purposes.append({'path': e['path'], 'purpose': e['purpose'], 'correct': ('完整计算结果' in e['purpose']) if e['path'].startswith('04_结果表格/') else ('空白结果模板' in e['purpose'])})
    txt_path = SUPPORT / '00_文件清单.txt'
    txt_paths = [line.strip() for line in txt_path.read_text(encoding='utf-8-sig').splitlines() if line.strip() in now]
    appendix_txt = SUPPORT / '00_可用于论文附录的支撑清单.txt'
    appendix_paths = [m[1] for line in appendix_txt.read_text(encoding='utf-8-sig').splitlines() if (m := re.match(r'^\d+\. (.+)$', line))]
    inventory.update({'sha256': digest(inventory_path), 'entry_count': len(entries), 'actual_file_count': len(now),
                       'entries': entries, 'each_actual_file_exactly_once': len(paths) == len(now) and set(paths) == set(now) and all(n == 1 for n in counts.values()),
                       'missing': sorted(set(now) - set(paths)), 'extra': sorted(set(paths) - set(now)),
                       'duplicate_paths': [p for p,n in counts.items() if n > 1],
                       'continuous_indices': [e['number'] for e in entries] == list(range(1,len(entries)+1)),
                       'summary_counts_equal_actual': parsed_summary == dict(categories), 'template_purpose_checks': template_purposes,
                       'txt_complete_once': Counter(txt_paths) == Counter(now.keys()),
                       'appendix_txt_complete_once': Counter(appendix_paths) == Counter(now.keys())})

checks = {
    'root_manuscript_unchanged': paper_hash == 'b6fbac8fefd87c10497a7c63f3638a7ead1f7e82c9bb0af7e773031a51f6960c',
    'all_30_production_files_unchanged': len(core30)==30 and all(p not in missing and baseline[p]['sha256'] == now[p]['sha256'] for p in core30),
    'all_four_frozen_results_preserved': all(r['frozen_equal'] and r['distinct_from_blank_template'] for r in frozen),
    'all_12_result_directory_files_preserved': all(p not in changed and p not in missing for p in baseline if p.startswith('04_结果表格/')),
    'original_05_sources_only_method_changed': len(old05sources) == 33 and [p for p in old05sources if p in changed] == [method_path] and not any(p in missing for p in old05sources),
    'method_before_matches_preceding_audit': digest(before_method) == baseline[method_path]['sha256'],
    'method_non_ROOT_ast_unchanged': method_only_root,
    'dependency_matplotlib_included': 'matplotlib==3.11.2' in requirements,
    'two_verification_inputs_exact': len(data05)==2 and all(d['equals_production_input'] for d in data05),
    'figures_14_CSV_4627_rows_exact': len(csv06)==14 and sum(c['rows'] for c in csv06)==4627 and all(c['source_equal'] for c in csv06),
    'six_delivered_PNGs_frozen_exact': len(png06)==6 and all(p['source_equal'] for p in png06),
    'figure_delivery_24_files_and_no_forbidden_formats': len(figfiles)==24 and not figforbidden,
    'current_R_matches_actual_independent_execution': figure_r['equals_executed_isolated_R'] and figure_r['actual_process_exit_zero'],
    'ai_detail_contains_known_uses_and_explicit_pending_fields': ai_report['contains_requested_tools'] and ai_report['has_pending_fields'] and ai_report['no_checked_confirmation_symbol'],
    'three_actual_small_runs_bind_current_source': len(run_evidence)==3 and all(r['returncode']==0 and r['source_input_unchanged'] and not r['current_delivery_diff_from_executed'] for r in run_evidence),
}
if inventory['exists']:
    checks.update({'inventory_DOCX_lists_every_current_file_once': inventory['each_actual_file_exactly_once'] and inventory['continuous_indices'] and inventory['summary_counts_equal_actual'],
                   'both_TXT_lists_complete_once': inventory['txt_complete_once'] and inventory['appendix_txt_complete_once'],
                   'blank_templates_and_results_distinguished': all(r['correct'] for r in inventory['template_purpose_checks'])})

report = {'at_local': datetime.datetime.now().isoformat(timespec='seconds'), 'status': 'PASS_STATIC_AUDIT' if all(checks.values()) and inventory['exists'] else 'WAIT_INVENTORY' if not inventory['exists'] and all(checks.values()) else 'FAIL',
          'checks': checks, 'paper_sha256': paper_hash, 'actual_support_files': len(now), 'changed_from_preceding_106': changed, 'missing_from_preceding_106': missing, 'added_from_preceding_106': added,
          'frozen_workbooks': frozen, 'verification_inputs': data05, 'figure_csv': csv06, 'figure_png': png06, 'figure_forbidden': figforbidden, 'figure_R': figure_r, 'ai_details': ai_report,
          'actual_run_records_reviewed': run_evidence, 'inventory': inventory,
          'scope': 'Independent read-only content/hash/path audit; no solver or document generation, no team sign-off, no final PDF approval. The new method ROOT line intentionally differs from the unchanged manuscript appendix code.'}
(QA / 'independent_audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: report[k] for k in ['status','actual_support_files','checks','changed_from_preceding_106','missing_from_preceding_106']},ensure_ascii=False,indent=2))
