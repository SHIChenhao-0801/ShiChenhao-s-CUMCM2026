from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('D:/Document/数学建模/2026CUMCM')
QA = ROOT / 'paper_output/qa/support_completion_20260913/ai'
OUT = ROOT / '支撑材料/AI工具使用详情.docx'
SOURCE = ROOT / '药材热湿耦合模型与干燥时间计算_附录修订版.docx'
QA.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)
sec.footer_distance = Cm(0.85)

def font(run, size=10.5, bold=False, color='000000'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rp = run._element.get_or_add_rPr()
    rf = rp.rFonts
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rp.insert(0, rf)
    for attr, val in [('eastAsia', '宋体'), ('ascii', 'Times New Roman'), ('hAnsi', 'Times New Roman')]:
        rf.set(qn('w:' + attr), val)

for name, size, bold in [('Normal', 10.5, False), ('Title', 21, True), ('Heading 1', 14, True), ('Heading 2', 11.5, True)]:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.space_before = Pt(0 if name in ('Normal', 'Title') else 8)

for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

def para(text='', style=None, size=10.5, bold=False, after=6):
    p = doc.add_paragraph(style=style)
    font(p.add_run(text), size=size, bold=bold)
    p.paragraph_format.space_after = Pt(after)
    return p

def heading(text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    return p

def newpage():
    doc.add_page_break()

def table(headers, rows, widths, fs=10.5, padding=95, keep=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        v = OxmlElement('w:' + side)
        v.set(qn('w:val'), 'single')
        v.set(qn('w:sz'), '6')
        v.set(qn('w:color'), 'D9D9D9')
        borders.append(v)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side in ['top', 'left', 'bottom', 'right']:
        v = OxmlElement('w:' + side)
        v.set(qn('w:w'), str(padding))
        v.set(qn('w:type'), 'dxa')
        margins.append(v)
    props.append(margins)
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    def fill(row, values, head=False):
        trpr = row._tr.get_or_add_trPr()
        if keep:
            trpr.append(OxmlElement('w:cantSplit'))
        if head:
            trpr.append(OxmlElement('w:tblHeader'))
        for idx, (c, text, width) in enumerate(zip(row.cells, values, widths)):
            c.width = Cm(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if head or (idx == 0 and width < 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for n, line in enumerate(str(text).split('\n')):
                if n:
                    p.add_run().add_break()
                font(p.add_run(line), size=fs, bold=head)
            if head:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'E8EEF3')
                c._tc.get_or_add_tcPr().append(shd)
    fill(t.rows[0], headers, True)
    for row in rows:
        fill(t.add_row(), row)
    return t

# Page numbers help cross-reference the three interaction records after editing.
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run('第 '), size=9)
for instr in ['PAGE', 'NUMPAGES']:
    if instr == 'NUMPAGES':
        font(footer.add_run(' 页  共 '), size=9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instr)
    rr = OxmlElement('w:r')
    tt = OxmlElement('w:t')
    tt.text = '1' if instr == 'PAGE' else '6'
    rr.append(tt)
    field.append(rr)
    footer._p.append(field)
font(footer.add_run(' 页'), size=9)

title = doc.add_paragraph('AI工具使用详情', style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(13)
para('对应论文  药材热湿耦合模型与干燥时间计算', size=11, after=10).alignment = WD_ALIGN_PARAGRAPH.CENTER
para('本队使用 Deepseek Harness 与 Deepseek v4.1 flash 组合辅助检查源代码语法并进行部分程序调试；使用 Codex 与 GPT6-Astra 组合辅助理解题目、检索参考文献、检查建模结果和整理公式。具体使用记录、采纳修改情况及人工核验信息列于本文件。')

heading('一 AI工具清单')
table(['工具及模型', '精确版本记录', '主要用途'], [
    ['Deepseek Harness\n与 Deepseek v4.1 flash', '模型：Deepseek v4.1 flash\n客户端版本：待填写', '检查已编写源代码的语法问题；结合报错信息和代码片段辅助部分程序调试。'],
    ['Codex\n与 GPT6-Astra', '模型：GPT6-Astra\n客户端版本：待填写', '理解题目；检索参考文献；检查建模结果；整理公式与符号说明。'],
], [4.2, 5.1, 7.7], fs=11, padding=125)
para('', after=2)
para('使用日期或时段：待填写。模型名称与版本须按实际使用界面或记录核对；如使用过不同版本，逐项增补。', size=10.5)

heading('交互方式及记录方法', level=2)
para('交互方式：待填写实际使用形式，例如文字对话、提交报错信息或附代码片段。涉及上传文件时，补充文件名称及其用途。')
para('典型交互记录保留完整提示词、AI 回复、队伍处理方式和人工核验过程。后附三张记录表可直接填写或粘贴聊天截图；记录较长时可续页，并在表内注明续页位置。')
para('填写说明：本文件中的“待填写”“待本队确认”栏目，以及未勾选的确认项，须由本队按真实记录补齐或核实后再定稿。', size=10.5)

newpage()
heading('二 分环节使用记录')
para('逐项登记七个环节的实际使用情况。标为“已使用”的内容对应已说明的用途；需要进一步核实的环节由本队补充使用情况与记录位置。')
table(['论文环节', '是否使用', '工具及具体工作', '记录或采用位置'], [
    ['赛题理解', '已使用', 'Codex 与 GPT6-Astra：梳理各问研究目标、已知条件、输入数据和输出要求；解释温度、干基含水率、物性参数及半径变化等量的含义和联系。', '待填写对应问题、论文位置或记录编号。'],
    ['模型假设', '待本队确认', '请核实是否使用 AI 辅助提出或检查假设；如使用，填写工具、假设内容与队伍处理方式。', '待填写；如未使用，经确认后注明。'],
    ['建模设计', '待本队确认', '请核实公式整理是否涉及模型设计、模型选择或推导建议；如涉及，填写 AI 的实际作用与本队完成的内容。', '待填写；如未使用，经确认后注明。'],
    ['代码求解', '已使用', 'Deepseek Harness 与 Deepseek v4.1 flash：检查括号、引号、缩进、语句结构及函数调用；解释报错并梳理变量、参数与函数间的传递关系，辅助部分调试。', '待填写源码文件、报错与修改位置，以及记录编号。'],
    ['结果检验', '已使用', 'Codex 与 GPT6-Astra：检查数值结果与题设条件、模型方程和边界条件的对应关系；核对温度、含水率和干燥时间的单位、趋势与表达。', '待填写结果表、图或检验记录编号。'],
    ['论文撰写', '已使用', 'Codex 与 GPT6-Astra：整理公式，梳理推导步骤，统一符号写法，检查公式衔接及其与变量说明的对应关系。', '待填写公式或章节位置；其他写作辅助情况待本队确认。'],
    ['辅助工作', '已使用', 'Codex 与 GPT6-Astra：围绕药材干燥、热湿耦合传递及相关数值方法检索文献线索和理论资料。', '待填写实际查阅文献、采用位置及记录编号。'],
], [2.1, 2.6, 8.65, 3.65], fs=10.5, padding=110)

records = [
    ('1', '源代码语法检查或程序调试', 'Deepseek Harness 与 Deepseek v4.1 flash', '代码求解'),
    ('2', '题意理解或参考文献检索', 'Codex 与 GPT6-Astra', '赛题理解或辅助工作'),
    ('3', '建模结果检查或公式整理', 'Codex 与 GPT6-Astra', '结果检验或论文撰写'),
]
for n, topic, tool, stage in records:
    newpage()
    heading('三 典型交互记录' if n == '1' else '三 典型交互记录续表')
    heading(f'记录{n} {topic}', level=2)
    para('请从对应用途的真实交互中选取一例并完整保留原文。截图中的姓名、学校、账号等身份信息按匿名要求处理，保留任务相关内容。', size=10.5)
    table(['记录字段', '填写内容'], [
        ['工具及环节', f'工具：{tool}\n环节：{stage}；实际日期与时间：待填写'],
        ['具体任务及材料', '待填写实际问题、输入文件或代码片段，以及本次交互需要解决的事项。\n'],
        ['完整提示词', '待粘贴真实提示词全文，或插入完整且清晰的截图。\n\n\n'],
        ['AI回复全文', '待粘贴真实 AI 回复全文，或插入完整且清晰的截图。内容较长可续页并标明位置。\n\n\n\n'],
        ['队伍处理方式', '待填写：采纳、修改后采纳、未采纳或仅作参考；说明具体修改及理由，并标明最终采用位置。\n'],
        ['人工核验过程与结果', '待填写实际采用的核验方法、核验内容和结论，并给出相应计算、运行、原文核对或其他记录的位置。\n'],
        ['记录核对', '核对状态：□ 已核对  □ 待核对\n核对日期：待填写；记录编号或续页位置：待填写'],
    ], [3.2, 13.8], fs=10.5, padding=100, keep=False)

newpage()
heading('四 采纳修改及人工主导确认')
heading('采纳与修改记录', level=2)
table(['内容类别', 'AI相关工作', '队伍采纳及修改情况'], [
    ['建模思路', '题意理解与文献线索；模型设计是否涉及 AI 建议待本队确认。', '待填写采纳范围、调整内容及理由；未采纳亦须据实注明。'],
    ['公式', '整理推导步骤、统一符号、检查公式衔接及变量说明。', '待填写修改前后内容及对应公式位置。'],
    ['代码', '源代码语法检查与部分程序调试。', '待填写报错定位、实际修改位置及修改后核验记录。'],
    ['分析内容', '检查结果与题设及模型条件的对应关系，核对单位、趋势和表达。', '待填写采纳意见、修订结果与依据，以及保留或未采纳的内容。'],
], [2.4, 6.6, 8.0], fs=10.5, padding=85)

heading('本队主导情况确认', level=2)
para('由本队逐项核实并勾选。若仍需补充，填写实际情况；本表不预先代替本队作出确认。', size=10, after=5)
table(['核心环节', '本队确认事项', '核实状态'], [
    ['模型创新与选型', '模型方案的选择、适用性判断及创新内容由本队主导完成。', '□ 已确认  □ 待确认'],
    ['公式推导', '关键公式及推导由本队理解、核查并主导完成。', '□ 已确认  □ 待确认'],
    ['代码逻辑', '核心代码逻辑由本队掌握并完成必要人工审查；运行结果经本队核验。', '□ 已确认  □ 待确认'],
    ['核心论证', '结论、解释及局限性由本队审查并主导完成。', '□ 已确认  □ 待确认'],
], [2.9, 10.0, 4.1], fs=10.5, padding=78)
para('需补充或更正的事项：待填写；全部核实且无补充时填写“无”。', size=10, after=4)

heading('真实性声明', level=2)
para('供本队核实后确认：本队确认，本文件所填 AI 工具、版本、使用环节、交互记录和采纳修改情况与实际过程一致，所附提示词与 AI 回复为真实记录。本队已对采用的内容进行必要的人工核验，并对论文及支撑材料的真实性、准确性和完整性负责。', size=10.5, after=5)
para('声明确认：□ 已确认  □ 待确认    确认日期：待填写', size=10.5, after=0)

# Neutral metadata; no team, institution, names or machine paths in the deliverable.
cp = doc.core_properties
cp.title = 'AI工具使用详情'
cp.subject = '药材热湿耦合模型与干燥时间计算 AI工具使用记录'
cp.author = ''
cp.last_modified_by = ''
cp.comments = ''
cp.keywords = ''
cp.category = ''
doc.save(OUT)

source_doc = Document(SOURCE)
texts = [p.text for p in source_doc.paragraphs]
idx = next(i for i, t in enumerate(texts) if t.startswith('9.3 AI'))
(QA / 'source_ai_report.txt').write_text('\n'.join(texts[idx:]) + '\n', encoding='utf-8')
audit = {
    'source': SOURCE.relative_to(ROOT).as_posix(),
    'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'output': OUT.relative_to(ROOT).as_posix(),
    'sha256': hashlib.sha256(OUT.read_bytes()).hexdigest(),
    'output_size': OUT.stat().st_size,
    'expected_pages': 6,
    'tables': len(doc.tables),
    'source_report_scope': ['syntax review', 'partial debugging', 'problem interpretation', 'literature search', 'result checking', 'formula organization'],
    'unfilled_fields': ['Harness client version', 'Codex client version', 'actual use dates', 'model assumption/design usage confirmation', 'source and paper location references', 'three actual prompt/reply records', 'adoption details', 'human verification and main-author confirmations', 'truthfulness confirmation'],
    'no_invented_prompt_or_response': True,
    'human_confirmation_not_preselected': True,
    'render_status': 'pending',
}
(QA / 'build_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(audit, ensure_ascii=False, indent=2))
