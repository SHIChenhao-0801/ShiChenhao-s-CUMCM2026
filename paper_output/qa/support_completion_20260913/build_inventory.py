from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
SUPPORT=ROOT/'支撑材料'
QA=Path(__file__).resolve().parent
OUT=ROOT/'支撑材料文件列表_附录用.docx'

USES={
 '00_支撑材料总说明.txt':'材料组织与复现入口说明',
 '00_文件清单.txt':'全部文件的相对路径清单',
 '00_可用于论文附录的支撑清单.txt':'含简短用途的完整文件清单',
 '00_整理验收说明.txt':'材料整理与真实核查范围',
 '00_材料来源说明.txt':'数据 程序 文献及图表来源',
 '源码去注释与沙盒核查.txt':'既有去注释与独立沙盒核查范围',
 'AI工具使用详情.docx':'AI工具与使用环节记录 可编辑详情',
 'A题.pdf':'A题题面及参数和结果要求',
 '附件1.xlsx':'原始环境温湿度观测数据',
 '附件2.xlsx':'原始药材半径观测数据',
 '数据说明.txt':'原始数据与清洗字段说明',
 'A_environment_observed.csv':'环境观测数据 时间 温度和湿度',
 'A_radius_observed.csv':'药材半径观测数据',
 '参考文献文件说明.txt':'参考文献版本及实际阅读范围',
 '补充资料说明.txt':'补充文献和网页资料来源',
 '中药材干燥技术与装备研究现状.pdf':'王乐意等 2024年综述全文28页',
 '数值模拟仿真研究现状及其在中药干燥领域应用展望_王晓辉.pdf':'王晓辉等 2023年综述全文8页',
 'Handbook of Industrial Drying.pdf':'工业干燥手册第4版5页节选',
 'Handbook of numerical analysis.pdf':'有限体积法作者2019年更新稿254页',
 'Shampine1997_BYU原始扫描_35页.pdf':'常微分方程求解补充资料35页扫描',
 'FAO_干燥原理_官方网页.html':'FAO干燥原理网页存档',
 'FAO_干燥原理_官方网页_离线文本.html':'FAO干燥原理可离线阅读文本',
 'SciPy_solve_ivp_官方API网页.html':'SciPy积分接口网页存档',
 'SciPy_solve_ivp_官方API网页_离线文本.html':'SciPy积分接口可离线阅读文本',
 'A_CodeReview.sln':'Visual Studio方案入口',
 'A_CodeReview.pyproj':'Visual Studio Python工程配置',
 'analyticJacobian.py':'正式模型解析稀疏Jacobian',
 'diskDense.py':'BDF连续输出多项式存储',
 'dryingCore.py':'热湿耦合方程离散与BDF求解',
 'exportOutputs.py':'四份结果工作簿导出与回读',
 'q1Model.py':'问题一模型配置',
 'q2Model.py':'问题二与三共同轨迹配置',
 'q3Model.py':'问题三严格达标时刻判定',
 'q4Model.py':'问题四收缩模型配置',
 'input_manifest.csv':'正式输入 模板和参照校验清单',
 'numerical_design.txt':'离散方程与阈值事件说明',
 'README.txt':'目录运行与依赖说明',
 'requirements.txt':'Python依赖及固定版本',
 'runDelivery.py':'四问生产计算与输入预检入口',
 'runLogged.ps1':'生产运行监督与退出码记录',
 '运行验收说明.txt':'生产运行证据与适用范围',
 '源码变更说明.txt':'生产源码来源和变更说明',
 'reference_data.py':'冻结数值与精度参照常量',
 '结果表说明.txt':'完整工作簿与正文表格说明',
 'q1_paper_moisture.csv':'正文问题一含水率结果表',
 'q1_paper_temperature.csv':'正文问题一温度结果表',
 'q2_paper_moisture.csv':'正文问题二含水率结果表',
 'q2_paper_temperature.csv':'正文问题二温度结果表',
 'q3_paper_moisture.csv':'正文问题三含水率结果表',
 'q4_paper_moisture.csv':'正文问题四含水率结果表',
 'q4_paper_radius.csv':'正文问题四半径结果表',
 '独立沙盒运行矩阵.csv':'既有沙盒逐任务运行状态',
 '独立沙盒逐源码核查.txt':'既有逐源码实算及失败范围',
 '数值检验说明.txt':'检验程序用途及运行前提',
 '源码文件与校验值.txt':'检验源码用途 来源与SHA256',
 'runCrossCheck.m':'MATLAB独立热湿模型对照实现',
 'compareMatlab.mjs':'MATLAB与Python结果比较程序',
 'analytic_jacobian.py':'检验模型解析Jacobian',
 'axisymmetric_check.py':'轴对称假设与轴向影响检验',
 'compare_analytic.py':'数值解与圆柱解析基准比较',
 'disk_dense.py':'检验模型BDF连续输出存储',
 'drying_core.py':'检验所用热湿模型核心',
 'export_outputs.py':'检验所用工作簿导出模块',
 'production_provenance.py':'运行监督 输入源码版本记录',
 'publication_plots.py':'既有绘图辅助程序',
 'q1_model.py':'检验所用问题一配置',
 'q2_model.py':'检验所用问题二配置',
 'q3_model.py':'检验所用问题三达标判定',
 'q4_model.py':'检验所用问题四配置',
 'run_experiments.py':'既有实验任务入口',
 'run_modeling.py':'旧模型流程入口 保留原运行前提',
 'validate_bessel.py':'圆柱Robin与Bessel解析基准',
 'verify_convergence.py':'空间网格加密误差检验',
 'verify_dense_storage.py':'BDF连续输出存储一致性检验',
 'verify_time_accuracy.py':'时间容差与步长敏感性检验',
 'analytic_metric_check.py':'解析误差指标历史诊断',
 'crossvalidate_series.py':'描述序列对照及历史趋势诊断',
 'crossvalidate_solver.py':'不同积分器的数值对照',
 'energy_balance_check.py':'模型能量收支检查',
 'energy_balance_diagnose.py':'能量闭合历史诊断',
 'energy_balance_diagnose2.py':'能量闭合进一步历史诊断',
 'isotherm_activity_closure.py':'历史等温线闭合试验',
 'isotherm_diagnose.py':'历史等温线接口诊断',
 'isotherm_jacobian_probe.py':'历史Jacobian注入探测',
 'latent_heat_scenarios.py':'潜热情景计算',
 'method_comparison.py':'两种守恒通量格式对照',
 'sensitivity_analysis.py':'参数扰动与敏感性计算',
 'threshold_and_scaling_checks.py':'阈值事件和标度检验',
 'draw_figures.R':'六幅论文图的PNG绘制程序',
 'run_checks.py':'网格 时间和方法检验统一入口',
 'run_verification.py':'网格 时间和方法检验统一入口',
 'requirements_verification.txt':'数值检验独立依赖及版本',
 'temperature.csv':'图1径向温度曲线数据',
 'moisture.csv':'图1径向含水率曲线数据',
 'observed_0_4h.csv':'图2前4小时环境实测点',
 'assumed_platform_4_60h.csv':'图2后续环境平台假设数据',
 'temperature_0_4h.csv':'图3中心和表面温度数据',
 'moisture_to_event.csv':'图3中心和表面含水率数据',
 'drying_main.csv':'图4全域含水率与达标曲线',
 'centre_near_event.csv':'图4临界附近中心含水率数据',
 'maximum_near_event.csv':'图4临界附近全域最大含水率',
 'event_markers.csv':'图4连续事件与严格报告点',
 'moisture_profiles_xy_pairs.csv':'图5收缩域内含水率XY数据',
 'radius_observed.csv':'图5原始半径观测数据',
 'selected_times_and_radii.csv':'图5所选时刻及对应半径',
 'scenario_comparison.csv':'图6经验边界阻力情景数据',
}

def purpose(rel):
    name=rel.name
    if re.fullmatch(r'result[1-4]\.xlsx',name):
        n=name[6]
        return f'问题{n}完整计算结果' if rel.parts[0]=='04_结果表格' else f'问题{n}空白结果模板'
    if name=='sampled_solution.npz':
        return f'{rel.parent.name}冻结采样数组'
    if rel.suffix.lower()=='.png':return name.removesuffix('.png')+'成图'
    if name in USES:return USES[name]
    if rel.suffix.lower()=='.txt':
        if any(x in name for x in ['依赖','环境']):return '运行环境与依赖说明'
        if any(x in name for x in ['复现','运行','适配']):return '本轮路径适配与实际复现说明'
        if any(x in name for x in ['来源','说明']):return '目录材料来源与使用说明'
    if rel.suffix.lower()=='.csv' and any(x in name for x in ['manifest','来源','校验','sha256']):return '文件来源与校验清单'
    raise ValueError('Missing explicit purpose: '+str(rel))

def font(run,size=10.5,bold=False):
    run.font.name='Times New Roman'
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)

def table(doc,head,rows,widths):
    t=doc.add_table(rows=1, cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for column,w in zip(t.columns,widths):column.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','bottom','left','right','insideH','insideV']:
        el=OxmlElement('w:'+name);el.set(qn('w:val'),'single' if name in {'top','bottom'} else 'nil');el.set(qn('w:sz'),'10');el.set(qn('w:color'),'000000');borders.append(el)
    pr.append(borders)
    for cell,txt,w in zip(t.rows[0].cells,head,widths):cell.width=Cm(w);font(cell.paragraphs[0].add_run(txt),10.5,True)
    hr=t.rows[0]._tr.get_or_add_trPr();el=OxmlElement('w:tblHeader');hr.append(el)
    for c in t.rows[0].cells:
        b=OxmlElement('w:tcBorders');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'5');b.append(bottom);c._tc.get_or_add_tcPr().append(b)
    for data in rows:
        cells=t.add_row().cells
        for cell,text,w in zip(cells,data,widths):cell.width=Cm(w);font(cell.paragraphs[0].add_run(str(text)))
    for r in t.rows:
        no=OxmlElement('w:cantSplit');r._tr.get_or_add_trPr().append(no)
        for i,c in enumerate(r.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,sz in [('top',65),('bottom',65),('left',70),('right',70)]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),str(sz));m.set(qn('w:type'),'dxa');mar.append(m)
            cpr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
                if i==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def main():
    files=sorted(p for p in SUPPORT.rglob('*') if p.is_file())
    forbidden=[p for p in files if p.suffix.lower() in {'.json','.jsonl','.md','.markdown','.pyc','.zip','.rar','.7z'} or '.vs' in p.parts]
    assert not forbidden,forbidden
    rels=[p.relative_to(SUPPORT) for p in files]
    assert Path('AI工具使用详情.docx') in rels
    assert len([p for p in rels if p.parts[0]=='06_绘图程序与数据' and p.suffix.lower()=='.png'])==6
    descriptions={p.as_posix():purpose(p) for p in rels}
    cats=Counter(p.parts[0] if len(p.parts)>1 else '根目录' for p in rels)
    groups=defaultdict(list)
    for rel in rels:groups[rel.parent.as_posix()].append(rel)
    listtext=['支撑材料完整文件列表','',f'共{len(rels)}个文件，含本目录内的TXT清单与AI详情DOCX。图像仅保留PNG，文件夹未压缩。','', '相对路径以本目录为根。']
    listtext.extend(['']+[p.as_posix() for p in rels])
    (SUPPORT/'00_文件清单.txt').write_text('\n'.join(listtext)+'\n',encoding='utf-8-sig')
    tabtext=['支撑材料文件列表','',f'共{len(rels)}个文件。以下逐项列出相对路径和用途；可用于论文附录。','']
    for i,p in enumerate(rels,1):tabtext.append(f'{i}. {p.as_posix()}\n   {descriptions[p.as_posix()]}')
    (SUPPORT/'00_可用于论文附录的支撑清单.txt').write_text('\n'.join(tabtext)+'\n',encoding='utf-8-sig')
    d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    for sty in ['Normal','Title','Heading 1','Heading 2']:
        s=d.styles[sty];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(5)
    p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('支撑材料文件列表'),16,True)
    p=d.add_paragraph(f'支撑材料共{len(rels)}个文件。下表覆盖题面与数据、参考资料、生产程序、结果表格、数值检验、绘图资料和AI工具使用详情。')
    p=d.add_paragraph('分组标题给出相对目录，表内文件名与该目录共同构成完整路径；所有路径均以“支撑材料”文件夹为基准。')
    p=d.add_paragraph('目录中的AI工具使用详情当前为可编辑DOCX。定稿转为PDF后，应同步更新本列表。')
    for p in d.paragraphs[1:]:
        for r in p.runs:font(r,10.5)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run('表1 支撑材料分类汇总'),11,True)
    table(d,['目录','文件数'],[(k,v) for k,v in sorted(cats.items(),key=lambda x:(x[0]!='根目录',x[0]))],[13.2,2.8])
    idx=0
    for gi,(group,paths) in enumerate(sorted(groups.items(),key=lambda x:(x[0]!='.',x[0])),2):
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(4)
        font(p.add_run(f'表{gi} '+('根目录' if group=='.' else group)),10.5,True)
        rows=[]
        for rel in paths:
            idx+=1;rows.append((idx,rel.name,descriptions[rel.as_posix()]))
        table(d,['序号','文件名','用途'],rows,[1.0,8.2,6.8])
    core=d.core_properties;core.author='';core.last_modified_by='';core.title='支撑材料文件列表';core.subject='A题支撑材料附录文件列表';core.comments='';core.keywords=''
    for root in [d.styles.element,d.element]:
        for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
    d.save(OUT)
    expected_paths={p.as_posix() for p in rels}
    assert expected_paths=={p.relative_to(SUPPORT).as_posix() for p in SUPPORT.rglob('*') if p.is_file()}
    report={'at_utc':datetime.now(timezone.utc).isoformat(),'docx':str(OUT),'docx_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'file_count':len(rels),'support_bytes':sum(p.stat().st_size for p in files),'categories':dict(cats),'groups':{k:[p.name for p in v] for k,v in groups.items()},'entries':[{'path':p.as_posix(),'purpose':descriptions[p.as_posix()],'sha256':hashlib.sha256((SUPPORT/p).read_bytes()).hexdigest()} for p in rels]}
    (QA/'inventory_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in {'groups','entries'}},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
